from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from backend.app.db import managed_connection  # noqa: E402
from backend.app.services.app_service import AppService  # noqa: E402
from backend.app.services.app_service_provider import create_app_service  # noqa: E402


DEFAULT_COMPANY_CODE = "COMPANY-MAIN"
TARGET_WORKSHOP_CODE = "1车间"
TARGET_WORKSHOP_NAME = "1车间"
TARGET_LINE_CODE = "球囊产线"
TARGET_LINE_NAME = "球囊产线"
PROCESS_CODE_PREFIX = "B"
PROCESS_CODE_START = 10
PROCESS_CODE_STEP = 10


@dataclass(frozen=True)
class BalloonProcessRow:
    process_name_cn: str
    capacity_per_shift: float
    required_workers: int
    required_machines: int


def _find_balloon_xlsx(resource_dir: Path) -> Path:
    candidates = sorted(resource_dir.glob("*.xlsx"))
    if not candidates:
        raise RuntimeError(f"No .xlsx file found under {resource_dir}")
    for path in candidates:
        workbook = load_workbook(path, data_only=True)
        if "Sheet2" in workbook.sheetnames and "Sheet1" in workbook.sheetnames:
            sheet2 = workbook["Sheet2"]
            header = [str(sheet2.cell(1, c).value or "").strip() for c in range(1, sheet2.max_column + 1)]
            if "工序名称" in header and "日设备标准产能" in header:
                return path
    raise RuntimeError("Cannot locate target balloon workbook by required headers.")


def _header_index(sheet, header_name: str, row_index: int = 1) -> int:
    for col in range(1, sheet.max_column + 1):
        value = str(sheet.cell(row_index, col).value or "").strip()
        if value == header_name:
            return col
    raise RuntimeError(f"Missing required header '{header_name}' in sheet '{sheet.title}' row {row_index}.")


def _to_float_or_none(value: Any) -> float | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number


def _to_int_or_none(value: Any) -> int | None:
    number = _to_float_or_none(value)
    if number is None:
        return None
    return int(round(number))


def _normalize_process_name(text: Any) -> str:
    return str(text or "").strip()


def _parse_required_machines(device_type: Any, device_code: Any) -> int:
    normalized_type = str(device_type or "").strip()
    code_text = str(device_code or "").strip()
    if normalized_type not in {"设备", "工装"}:
        return 0
    if not code_text:
        return 0
    parts = [part.strip() for part in re.split(r"[\\/／]+", code_text) if part.strip()]
    return len(parts) if parts else 0


def _read_sheet1_worker_map(sheet) -> dict[str, int]:
    process_col = _header_index(sheet, "工序")
    worker_col = _header_index(sheet, "人力")
    out: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        process_name = _normalize_process_name(sheet.cell(row, process_col).value)
        if not process_name:
            continue
        worker = _to_int_or_none(sheet.cell(row, worker_col).value)
        if worker is None or worker <= 0:
            continue
        if process_name not in out:
            out[process_name] = worker
    return out


def _read_sheet2_process_rows(sheet, worker_map: dict[str, int]) -> list[BalloonProcessRow]:
    process_col = _header_index(sheet, "工序名称")
    capacity_col = _header_index(sheet, "日设备标准产能")
    device_type_col = _header_index(sheet, "设备类型")
    device_code_col = _header_index(sheet, "设备编码")

    process_map: dict[str, BalloonProcessRow] = {}
    ordered_names: list[str] = []
    for row in range(2, sheet.max_row + 1):
        process_name = _normalize_process_name(sheet.cell(row, process_col).value)
        if not process_name:
            continue
        capacity = _to_float_or_none(sheet.cell(row, capacity_col).value)
        if capacity is None or capacity <= 0:
            continue
        required_machines = _parse_required_machines(
            sheet.cell(row, device_type_col).value,
            sheet.cell(row, device_code_col).value,
        )
        required_workers = worker_map.get(process_name, 1)
        if required_workers <= 0:
            required_workers = 1
        if process_name not in process_map:
            ordered_names.append(process_name)
        process_map[process_name] = BalloonProcessRow(
            process_name_cn=process_name,
            capacity_per_shift=capacity,
            required_workers=required_workers,
            required_machines=required_machines,
        )
    return [process_map[name] for name in ordered_names]


def _existing_process_codes(service: AppService) -> set[str]:
    codes = set()
    for row in service._list_line_topology_rows():
        code = str(row.get("process_code") or "").strip().upper()
        if code:
            codes.add(code)
    for row in service._list_route_rows():
        code = str(row.get("process_code") or "").strip().upper()
        if code:
            codes.add(code)
    return codes


def _build_mapping_payload(service: AppService, process_rows: list[BalloonProcessRow]) -> tuple[dict[str, Any], dict[str, str]]:
    config = service.get_masterdata_config()["data"]
    line_skeletons = list(config.get("line_skeletons") or [])
    line_topology = list(config.get("line_topology") or [])
    scopes = list(config.get("workshop_manager_line_scopes") or [])

    skeleton_exists = any(
        str(row.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() == DEFAULT_COMPANY_CODE
        and str(row.get("workshop_code") or "").strip() == TARGET_WORKSHOP_CODE
        and str(row.get("line_code") or "").strip() == TARGET_LINE_CODE
        for row in line_skeletons
    )
    if not skeleton_exists:
        line_skeletons.append(
            {
                "company_code": DEFAULT_COMPANY_CODE,
                "workshop_code": TARGET_WORKSHOP_CODE,
                "workshop_name": TARGET_WORKSHOP_NAME,
                "line_code": TARGET_LINE_CODE,
                "line_name": TARGET_LINE_NAME,
                "enabled_flag": 1,
            }
        )

    existing_balloon_rows = [
        row
        for row in line_topology
        if (
            str(row.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() == DEFAULT_COMPANY_CODE
            and str(row.get("workshop_code") or "").strip() == TARGET_WORKSHOP_CODE
            and str(row.get("line_code") or "").strip() == TARGET_LINE_CODE
        )
    ]
    existing_balloon_codes = {
        str(row.get("process_code") or "").strip().upper()
        for row in existing_balloon_rows
        if str(row.get("process_code") or "").strip()
    }

    # Remove previous balloon rows to keep import idempotent.
    line_topology = [
        row
        for row in line_topology
        if not (
            str(row.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() == DEFAULT_COMPANY_CODE
            and str(row.get("workshop_code") or "").strip() == TARGET_WORKSHOP_CODE
            and str(row.get("line_code") or "").strip() == TARGET_LINE_CODE
        )
    ]

    used_codes = _existing_process_codes(service)
    for code in existing_balloon_codes:
        used_codes.discard(code)

    name_to_code: dict[str, str] = {}
    for index, item in enumerate(process_rows):
        code_number = PROCESS_CODE_START + index * PROCESS_CODE_STEP
        process_code = f"{PROCESS_CODE_PREFIX}{code_number:03d}"
        if process_code in used_codes:
            raise RuntimeError(
                f"Process code conflict for balloon import: {process_code} already used by non-balloon data."
            )
        used_codes.add(process_code)
        name_to_code[item.process_name_cn] = process_code
        line_topology.append(
            {
                "company_code": DEFAULT_COMPANY_CODE,
                "workshop_code": TARGET_WORKSHOP_CODE,
                "workshop_name": TARGET_WORKSHOP_NAME,
                "line_code": TARGET_LINE_CODE,
                "line_name": TARGET_LINE_NAME,
                "process_code": process_code,
                "capacity_per_shift": item.capacity_per_shift,
                "required_workers": item.required_workers,
                "required_machines": item.required_machines,
                "enabled_flag": 1,
            }
        )

    payload = {
        "line_skeletons": line_skeletons,
        "line_topology": line_topology,
        "workshop_manager_line_scopes": scopes,
    }
    return payload, name_to_code


def main() -> None:
    project_root = Path(__file__).resolve().parents[2]
    resource_dir = project_root / "resource"
    workbook_path = _find_balloon_xlsx(resource_dir)
    workbook = load_workbook(workbook_path, data_only=True)
    worker_map = _read_sheet1_worker_map(workbook["Sheet1"])
    process_rows = _read_sheet2_process_rows(workbook["Sheet2"], worker_map)
    if not process_rows:
        raise RuntimeError("No valid process rows found in Sheet2.")

    with managed_connection() as connection:
        service = create_app_service(connection)
        payload, name_to_code = _build_mapping_payload(service, process_rows)
        service.save_masterdata_config(payload)

    print(f"Imported workbook: {workbook_path.name}")
    print(f"Target line: {TARGET_WORKSHOP_CODE} / {TARGET_LINE_CODE}")
    print(f"Imported process count: {len(process_rows)}")
    print("Process mapping (name -> code):")
    print(json.dumps(name_to_code, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
