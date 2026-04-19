from __future__ import annotations

from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from ..db import fetch_all, fetch_one, transaction, utc_now
from ..errors import bad_request, forbidden, not_found, server_error
from ..json_utils import dumps
from .reporting_import_support import (
    REPORTING_IMPORT_OPTIONAL_HEADER_MAP,
    REPORTING_IMPORT_REQUIRED_HEADER_MAP,
    WORK_REPORT_COLUMNS,
    compute_file_sha256,
    normalize_sha256_text,
    parse_optional_number,
    parse_optional_text,
    parse_report_datetime,
    parse_required_positive_number,
    to_utc_iso_text,
)
from .user_scope_support import assert_actor_can_access_line, is_workshop_manager


def _to_number(value: object, fallback: float = 0.0) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return fallback
    return number if number == number else fallback


def _local_date_from_iso_datetime(value: object) -> str | None:
    from datetime import timezone, timedelta, datetime

    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None
    local_timezone = timezone(timedelta(hours=8))
    return parsed.astimezone(local_timezone).date().isoformat()


def _normalize_date_text(value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        from datetime import date

        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def _parse_enabled_flag(value: object) -> int:
    if isinstance(value, bool):
        return 1 if value else 0
    if isinstance(value, int):
        if value in (0, 1):
            return value
        raise ValueError("enabled_flag must be 0 or 1.")
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "on"}:
        return 1
    if text in {"0", "false", "no", "off"}:
        return 0
    raise ValueError("enabled_flag must be 0/1 or a boolean-like value.")


class ReportingCommandService:
    def __init__(self, host: Any) -> None:
        self.host = host
        self.connection = host.connection

    def import_mes_reportings_from_xlsx(self, payload: dict[str, Any]) -> dict[str, Any]:
        file_path = str(payload.get("file_path") or "").strip()
        if not file_path:
            raise bad_request(
                code="REPORTING_IMPORT_FILE_PATH_REQUIRED",
                message="file_path is required.",
            )
        if not file_path.lower().endswith(".xlsx"):
            raise bad_request(
                code="REPORTING_IMPORT_FILE_TYPE_INVALID",
                message="Only .xlsx files are supported.",
                details={"file_path": file_path},
            )
        workbook_path = Path(file_path)
        if not workbook_path.exists():
            raise bad_request(
                code="REPORTING_IMPORT_FILE_NOT_FOUND",
                message="xlsx file does not exist.",
                details={"file_path": file_path},
            )

        file_sha256 = normalize_sha256_text(payload.get("file_sha256"))
        if file_sha256 is None:
            file_sha256 = normalize_sha256_text(payload.get("source_file_name"))
        if file_sha256 is None:
            file_sha256 = compute_file_sha256(workbook_path)

        sha_source_file_name = f"sha256:{file_sha256}"
        explicit_source_file_name = str(payload.get("source_file_name") or "").strip()
        source_file_names: list[str] = []
        for candidate in (sha_source_file_name, explicit_source_file_name, file_path):
            if candidate and candidate not in source_file_names:
                source_file_names.append(candidate)

        original_file_name = (
            str(payload.get("original_file_name") or "").strip() or workbook_path.name
        )
        file_size_bytes = int(workbook_path.stat().st_size or 0)
        size_input = str(payload.get("file_size_bytes") or "").strip()
        if size_input:
            try:
                file_size_bytes = max(0, int(size_input))
            except ValueError:
                file_size_bytes = int(workbook_path.stat().st_size or 0)

        company_code = str(payload.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
        create_missing_orders = False
        if payload.get("create_missing_orders") is not None:
            try:
                create_missing_orders = _parse_enabled_flag(payload.get("create_missing_orders")) == 1
            except ValueError as exc:
                raise bad_request(
                    code="REPORTING_IMPORT_CREATE_MISSING_ORDERS_INVALID",
                    message=str(exc),
                ) from exc
        sheet_names_payload = payload.get("sheet_names")
        sheet_names: list[str] | None = None
        if sheet_names_payload is not None:
            if not isinstance(sheet_names_payload, list):
                raise bad_request(
                    code="REPORTING_IMPORT_SHEET_NAMES_INVALID",
                    message="sheet_names must be an array of strings.",
                )
            normalized_sheet_names: list[str] = []
            for item in sheet_names_payload:
                name = str(item or "").strip()
                if not name:
                    raise bad_request(
                        code="REPORTING_IMPORT_SHEET_NAMES_INVALID",
                        message="sheet_names must not contain empty names.",
                    )
                if name not in normalized_sheet_names:
                    normalized_sheet_names.append(name)
            sheet_names = normalized_sheet_names

        required_headers = ("报工日期", "生产订单号", "工序编码", "报工数量")
        header_to_field: dict[str, str] = {
            **REPORTING_IMPORT_REQUIRED_HEADER_MAP,
            **REPORTING_IMPORT_OPTIONAL_HEADER_MAP,
            "工序编码": "process_code",
            "工序名称": "process_name",
            "宸ュ簭缂栫爜": "process_code",
            "宸ュ簭鍚嶇О": "process_name",
        }

        try:
            workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        except Exception as exc:
            raise bad_request(
                code="REPORTING_IMPORT_WORKBOOK_LOAD_FAILED",
                message="Failed to read xlsx workbook.",
                details={"file_path": file_path, "error": str(exc)},
            ) from exc

        try:
            available_sheet_names = list(workbook.sheetnames)
            selected_sheet_names = sheet_names or available_sheet_names
            if not selected_sheet_names:
                raise bad_request(
                    code="REPORTING_IMPORT_SHEETS_EMPTY",
                    message="xlsx workbook has no sheets.",
                    details={"file_path": file_path},
                )
            missing_sheets = [
                name
                for name in selected_sheet_names
                if name not in available_sheet_names
            ]
            if missing_sheets:
                raise bad_request(
                    code="REPORTING_IMPORT_SHEET_NOT_FOUND",
                    message="Some sheets do not exist in workbook.",
                    details={
                        "file_path": file_path,
                        "missing_sheet_names": missing_sheets,
                        "available_sheet_names": available_sheet_names,
                    },
                )

            imported_count = 0
            skipped_existing_count = 0
            failed_count = 0
            total_row_count = 0
            failures: list[dict[str, Any]] = []
            order_exists_cache: dict[str, bool] = {}
            created_missing_order_count = 0
            created_missing_order_nos: list[str] = []

            for sheet_name in selected_sheet_names:
                sheet = workbook[sheet_name]
                header_cells = [
                    str(sheet.cell(1, c).value or "").strip()
                    for c in range(1, sheet.max_column + 1)
                ]
                header_index: dict[str, int] = {}
                for col_index, header in enumerate(header_cells, start=1):
                    if header and header not in header_index:
                        header_index[header] = col_index

                missing_headers = [
                    header for header in required_headers if header not in header_index
                ]
                if missing_headers:
                    raise bad_request(
                        code="REPORTING_IMPORT_REQUIRED_HEADERS_MISSING",
                        message=f"Missing required header(s) in sheet '{sheet_name}'.",
                        details={
                            "sheet_name": sheet_name,
                            "missing_headers": missing_headers,
                        },
                    )

                sheet_columns: list[tuple[str, int]] = []
                for header, field_name in header_to_field.items():
                    col = header_index.get(header)
                    if col is None:
                        continue
                    sheet_columns.append((field_name, col))

                for row_no in range(2, sheet.max_row + 1):
                    raw: dict[str, Any] = {
                        field_name: sheet.cell(row_no, col).value
                        for field_name, col in sheet_columns
                    }
                    report_datetime_raw = raw.get("report_datetime")
                    order_no_raw = raw.get("production_order_no")
                    process_code_raw = raw.get("process_code")
                    report_qty_raw = raw.get("report_qty")
                    if (
                        report_datetime_raw is None
                        and not parse_optional_text(order_no_raw)
                        and not parse_optional_text(process_code_raw)
                        and parse_optional_text(report_qty_raw) is None
                    ):
                        continue

                    total_row_count += 1
                    source_placeholder_sql = ",".join("?" for _ in source_file_names)
                    existing = fetch_one(
                        self.connection,
                        f"""
                        SELECT report_id
                        FROM work_reports
                        WHERE source_file_name IN ({source_placeholder_sql})
                          AND source_sheet_name = ?
                          AND source_row_no = ?
                        LIMIT 1
                        """,
                        tuple(source_file_names + [sheet_name, row_no]),
                    )
                    if existing is not None:
                        skipped_existing_count += 1
                        continue

                    report_dt = parse_report_datetime(report_datetime_raw)
                    if report_dt is None:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "error": "报工日期无法解析。",
                            }
                        )
                        continue
                    report_time = to_utc_iso_text(report_dt)

                    order_no_text = parse_optional_text(order_no_raw)
                    if not order_no_text:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "error": "生产订单号为空。",
                            }
                        )
                        continue
                    base_order_no = str(order_no_text.split("-")[0]).strip()
                    if not base_order_no:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "order_no": order_no_text,
                                "error": "生产订单号格式无效。",
                            }
                        )
                        continue

                    process_code_text = parse_optional_text(process_code_raw)
                    if not process_code_text:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "order_no": order_no_text,
                                "error": "工序编码为空。",
                            }
                        )
                        continue
                    process_code = process_code_text.upper()

                    report_qty = parse_required_positive_number(report_qty_raw)
                    if report_qty is None:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "order_no": order_no_text,
                                "process_code": process_code,
                                "error": "报工数量必须大于 0。",
                            }
                        )
                        continue

                    if base_order_no not in order_exists_cache:
                        order_exists_cache[base_order_no] = (
                            fetch_one(
                                self.connection,
                                """
                                SELECT production_order_no
                                FROM production_orders
                                WHERE production_order_no = ?
                                LIMIT 1
                                """,
                                (base_order_no,),
                            )
                            is not None
                        )

                    if not order_exists_cache[base_order_no]:
                        if not create_missing_orders:
                            failed_count += 1
                            failures.append(
                                {
                                    "sheet_name": sheet_name,
                                    "row_no": row_no,
                                    "order_no": order_no_text,
                                    "base_order_no": base_order_no,
                                    "process_code": process_code,
                                    "error": "生产订单不存在，请先同步或导入生产订单后再导入报工。",
                                }
                            )
                            continue
                        material_code = parse_optional_text(raw.get("product_code")) or "UNKNOWN"
                        material_name = (
                            parse_optional_text(raw.get("product_name"))
                            or material_code
                            or base_order_no
                        )
                        material_specification = parse_optional_text(
                            raw.get("product_specification")
                        )
                        with transaction(self.connection):
                            self.connection.execute(
                                """
                                INSERT INTO production_orders (
                                    production_order_no,
                                    material_code,
                                    material_name,
                                    material_specification,
                                    production_qty,
                                    status,
                                    planned_start_date,
                                    planned_end_date,
                                    source_bill_no,
                                    material_list_no,
                                    updated_at
                                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                ON CONFLICT(production_order_no) DO NOTHING
                                """,
                                (
                                    base_order_no,
                                    material_code,
                                    material_name,
                                    material_specification,
                                    0.0,
                                    "OPEN",
                                    None,
                                    None,
                                    None,
                                    None,
                                    report_time,
                                ),
                            )
                        order_exists_cache[base_order_no] = True
                        created_missing_order_count += 1
                        if len(created_missing_order_nos) < 50:
                            created_missing_order_nos.append(base_order_no)

                    process_name = parse_optional_text(raw.get("process_name"))
                    report_id = f"RPT-{uuid4().hex[:10].upper()}"
                    row_payload: dict[str, Any] = {
                        "report_id": report_id,
                        "production_order_no": base_order_no,
                        "process_code": process_code,
                        "process_name": process_name,
                        "company_code": company_code,
                        "workshop_code": None,
                        "workshop_name": None,
                        "line_code": None,
                        "line_name": None,
                        "report_qty": float(report_qty),
                        "report_time": report_time,
                        "operator_code": parse_optional_text(raw.get("operator_code")),
                        "operator_name": parse_optional_text(raw.get("operator_name")),
                        "section_leader_name": parse_optional_text(
                            raw.get("section_leader_name")
                        ),
                        "dispatch_no": parse_optional_text(raw.get("dispatch_no")),
                        "product_code": parse_optional_text(raw.get("product_code")),
                        "product_name": parse_optional_text(raw.get("product_name")),
                        "product_specification": parse_optional_text(
                            raw.get("product_specification")
                        ),
                        "resource_group_name": parse_optional_text(
                            raw.get("resource_group_name")
                        ),
                        "resource_name": parse_optional_text(raw.get("resource_name")),
                        "department_name": parse_optional_text(
                            raw.get("department_name")
                        ),
                        "source_process_code": process_code,
                        "source_process_name": process_name,
                        "mold_code": parse_optional_text(raw.get("mold_code")),
                        "support_count": parse_optional_number(raw.get("support_count")),
                        "weight_kg": parse_optional_number(raw.get("weight_kg")),
                        "cavity_count": parse_optional_number(raw.get("cavity_count")),
                        "total_cycle_time": parse_optional_number(
                            raw.get("total_cycle_time")
                        ),
                        "production_quota": parse_optional_number(
                            raw.get("production_quota")
                        ),
                        "work_duration": parse_optional_number(raw.get("work_duration")),
                        "clamp_or_assembly_weight": parse_optional_number(
                            raw.get("clamp_or_assembly_weight")
                        ),
                        "unit_weight": parse_optional_number(raw.get("unit_weight")),
                        "source_sheet_name": sheet_name,
                        "source_row_no": row_no,
                        "source_file_name": sha_source_file_name,
                        "updated_at": report_time,
                    }
                    columns_sql = ", ".join(WORK_REPORT_COLUMNS)
                    placeholders = ", ".join("?" for _ in WORK_REPORT_COLUMNS)
                    values = tuple(row_payload.get(column) for column in WORK_REPORT_COLUMNS)
                    with transaction(self.connection):
                        self.connection.execute(
                            f"INSERT INTO work_reports ({columns_sql}) VALUES ({placeholders})",
                            values,
                        )
                    imported_count += 1

            actor = payload.get("actor") if isinstance(payload.get("actor"), dict) else {}
            imported_by_user_id = str(actor.get("user_id") or "").strip() or None
            imported_by_username = str(actor.get("username") or "").strip() or None
            imported_by_display_name = (
                str(actor.get("display_name") or "").strip() or None
            )
            imported_at = utc_now()
            previous_imported_at: str | None = None
            existing_file_record = fetch_one(
                self.connection,
                """
                SELECT last_imported_at
                FROM reporting_import_files
                WHERE file_sha256 = ?
                """,
                (file_sha256,),
            )
            if existing_file_record is not None:
                previous_imported_at = str(
                    existing_file_record.get("last_imported_at") or ""
                ).strip() or None

            with transaction(self.connection):
                self.connection.execute(
                    """
                    INSERT INTO reporting_import_files (
                        file_sha256,
                        source_file_name,
                        file_path,
                        original_file_name,
                        file_size_bytes,
                        sheet_names_json,
                        imported_by_user_id,
                        imported_by_username,
                        imported_by_display_name,
                        total_row_count,
                        imported_count,
                        skipped_existing_count,
                        failed_count,
                        created_missing_order_count,
                        created_at,
                        last_imported_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(file_sha256) DO UPDATE SET
                        source_file_name = excluded.source_file_name,
                        file_path = excluded.file_path,
                        original_file_name = excluded.original_file_name,
                        file_size_bytes = excluded.file_size_bytes,
                        sheet_names_json = excluded.sheet_names_json,
                        imported_by_user_id = excluded.imported_by_user_id,
                        imported_by_username = excluded.imported_by_username,
                        imported_by_display_name = excluded.imported_by_display_name,
                        total_row_count = excluded.total_row_count,
                        imported_count = excluded.imported_count,
                        skipped_existing_count = excluded.skipped_existing_count,
                        failed_count = excluded.failed_count,
                        created_missing_order_count = excluded.created_missing_order_count,
                        last_imported_at = excluded.last_imported_at
                    """,
                    (
                        file_sha256,
                        sha_source_file_name,
                        file_path,
                        original_file_name,
                        int(file_size_bytes),
                        dumps(selected_sheet_names),
                        imported_by_user_id,
                        imported_by_username,
                        imported_by_display_name,
                        int(total_row_count),
                        int(imported_count),
                        int(skipped_existing_count),
                        int(failed_count),
                        int(created_missing_order_count),
                        imported_at,
                        imported_at,
                    ),
                )

            return {
                "file_path": file_path,
                "source_file_name": sha_source_file_name,
                "file_sha256": file_sha256,
                "original_file_name": original_file_name,
                "file_size_bytes": int(file_size_bytes),
                "previous_imported_at": previous_imported_at,
                "sheet_names": selected_sheet_names,
                "total_row_count": total_row_count,
                "imported_count": imported_count,
                "skipped_existing_count": skipped_existing_count,
                "failed_count": failed_count,
                "created_missing_order_count": created_missing_order_count,
                "created_missing_order_nos": created_missing_order_nos,
                "failures": failures,
            }
        finally:
            workbook.close()

    def select_reporting_capacity_compare(self, payload: dict[str, Any]) -> dict[str, Any]:
        report_id = str(payload.get("report_id") or "").strip()
        if not report_id:
            raise bad_request(
                code="REPORT_ID_REQUIRED",
                message="report_id is required.",
            )
        audit_id = str(payload.get("audit_id") or "").strip()
        if not audit_id:
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_AUDIT_ID_REQUIRED",
                message="audit_id is required.",
            )
        actor = payload.get("actor") if isinstance(payload.get("actor"), dict) else {}
        existing = fetch_one(
            self.connection,
            """
            SELECT
                report_id,
                process_code,
                workshop_code,
                line_code,
                report_time,
                daily_capacity_compare_audit_id,
                daily_capacity_compare_qty
            FROM work_reports
            WHERE report_id = ?
            """,
            (report_id,),
        )
        if existing is None:
            raise not_found(
                code="REPORT_NOT_FOUND",
                message="Report does not exist.",
                details={"report_id": report_id},
            )

        workshop_code = str(existing.get("workshop_code") or "").strip().upper()
        line_code = str(existing.get("line_code") or "").strip().upper()
        process_code = str(existing.get("process_code") or "").strip().upper()
        report_local_date = _local_date_from_iso_datetime(existing.get("report_time"))
        if not report_local_date or not workshop_code or not line_code or not process_code:
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_REPORT_KEY_INVALID",
                message=(
                    "The report record is missing report_time, workshop_code, line_code, "
                    "or process_code required for capacity comparison."
                ),
                details={"report_id": report_id},
            )

        if is_workshop_manager(actor):
            assert_actor_can_access_line(
                self.connection,
                actor,
                company_code="COMPANY-MAIN",
                workshop_code=workshop_code,
                line_code=line_code,
                missing_user_error_code="REPORT_ACTOR_USER_REQUIRED",
                forbidden_error_code="REPORT_LINE_SCOPE_FORBIDDEN",
                forbidden_message=(
                    "Current workshop manager is not allowed to select a capacity comparison for this reporting record."
                ),
            )

        candidate_rows = fetch_all(
            self.connection,
            """
            SELECT
                audit_id,
                new_planned_capacity_qty
            FROM daily_line_capacity_plan_audit
            WHERE calendar_date = ?
              AND company_code = ?
              AND workshop_code = ?
              AND line_code = ?
              AND process_code = ?
            ORDER BY changed_at DESC, audit_id DESC
            """,
            (
                report_local_date,
                "COMPANY-MAIN",
                workshop_code,
                line_code,
                process_code,
            ),
        )
        if len(candidate_rows) == 0:
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_AUDITS_EMPTY",
                message="No same-day daily capacity audit records exist for this report.",
                details={
                    "report_id": report_id,
                    "report_local_date": report_local_date,
                    "workshop_code": workshop_code,
                    "line_code": line_code,
                    "process_code": process_code,
                },
            )

        selected_audit = next(
            (
                row
                for row in candidate_rows
                if str(row.get("audit_id") or "").strip() == audit_id
            ),
            None,
        )
        if selected_audit is None:
            any_audit = fetch_one(
                self.connection,
                """
                SELECT audit_id
                FROM daily_line_capacity_plan_audit
                WHERE audit_id = ?
                LIMIT 1
                """,
                (audit_id,),
            )
            if any_audit is None:
                raise not_found(
                    code="DAILY_CAPACITY_AUDIT_NOT_FOUND",
                    message="Daily capacity audit does not exist.",
                    details={"audit_id": audit_id},
                )
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_AUDIT_SCOPE_MISMATCH",
                message=(
                    "The selected daily capacity audit does not match the report date, "
                    "workshop, line, and process."
                ),
                details={
                    "report_id": report_id,
                    "audit_id": audit_id,
                    "report_local_date": report_local_date,
                    "workshop_code": workshop_code,
                    "line_code": line_code,
                    "process_code": process_code,
                },
            )

        compare_qty = _to_number(selected_audit.get("new_planned_capacity_qty"), 0)
        selected_at = utc_now()
        with transaction(self.connection):
            self.connection.execute(
                """
                UPDATE work_reports
                SET daily_capacity_compare_audit_id = ?,
                    daily_capacity_compare_qty = ?,
                    daily_capacity_compare_selected_at = ?
                WHERE report_id = ?
                """,
                (audit_id, compare_qty, selected_at, report_id),
            )
        return {
            "report_id": report_id,
            "report_local_date": report_local_date,
            "daily_capacity_compare_audit_id": audit_id,
            "daily_capacity_compare_qty": compare_qty,
            "daily_capacity_compare_selected_at": selected_at,
        }

    def create_reporting(self, payload: dict[str, Any]) -> dict[str, Any]:
        from datetime import datetime, timezone, timedelta
        from uuid import uuid4

        order_no = str(payload.get("order_no") or "").strip() or None
        process_code = str(payload.get("process_code") or "").strip().upper()
        report_qty = _to_number(payload.get("report_qty"), -1)
        report_scope = str(payload.get("report_scope") or "").strip().upper()
        if not process_code:
            raise bad_request(
                code="PROCESS_CODE_REQUIRED",
                message="process_code is required.",
            )
        if report_qty <= 0:
            raise bad_request(
                code="REPORT_QTY_INVALID",
                message="report_qty must be greater than 0.",
            )
        if report_scope and report_scope not in {"ORDER", "LINE_OUTPUT"}:
            raise bad_request(
                code="REPORT_SCOPE_INVALID",
                message="report_scope must be ORDER or LINE_OUTPUT.",
            )
        if order_no:
            self.host._require_order(order_no)
        elif report_scope == "ORDER":
            raise bad_request(
                code="REPORT_ORDER_REQUIRED",
                message="订单报工必须指定 order_no。",
            )
        normalized_report_scope = report_scope or ("ORDER" if order_no else "LINE_OUTPUT")
        process_name_by_code = self.host._process_name_by_code()
        report_id = f"RPT-{uuid4().hex[:10].upper()}"
        report_time_text = str(payload.get("report_time") or "").strip()
        calendar_date_input = payload.get("calendar_date")
        calendar_date_text = str(calendar_date_input or "").strip()
        normalized_calendar_date = _normalize_date_text(calendar_date_input)
        if calendar_date_text and normalized_calendar_date is None:
            raise bad_request(
                code="CALENDAR_DATE_INVALID",
                message="calendar_date must be a valid YYYY-MM-DD date.",
            )
        local_timezone = timezone(timedelta(hours=8))
        if report_time_text:
            try:
                report_time_dt = datetime.fromisoformat(report_time_text)
            except ValueError as exc:
                raise bad_request(
                    code="REPORT_TIME_INVALID",
                    message="report_time must be a valid ISO datetime.",
                ) from exc
            if report_time_dt.tzinfo is None or report_time_dt.utcoffset() is None:
                raise bad_request(
                    code="REPORT_TIME_TIMEZONE_REQUIRED",
                    message="report_time must include timezone offset.",
                )
            report_time = report_time_dt.astimezone(timezone.utc).replace(microsecond=0).isoformat()
            reporting_local_date = report_time_dt.astimezone(local_timezone).date().isoformat()
        elif normalized_calendar_date:
            reporting_local_date = normalized_calendar_date
            report_time = (
                datetime.fromisoformat(f"{normalized_calendar_date}T10:00:00+08:00")
                .astimezone(timezone.utc)
                .replace(microsecond=0)
                .isoformat()
            )
        else:
            simulation_state = self.host._get_simulation_state()
            reporting_local_date = _normalize_date_text(simulation_state.get("current_date"))
            if reporting_local_date is None:
                raise server_error(
                    code="SIMULATION_CURRENT_DATE_INVALID",
                    message="Current simulation date is invalid.",
                    details={"current_date": simulation_state.get("current_date")},
                )
            report_time = (
                datetime.fromisoformat(f"{reporting_local_date}T10:00:00+08:00")
                .astimezone(timezone.utc)
                .replace(microsecond=0)
                .isoformat()
            )
        operator_name = (
            str(
                payload.get("operator_name_cn")
                or payload.get("operator_name")
                or "绯荤粺濉姤"
            ).strip()
            or "绯荤粺濉姤"
        )
        workshop_code = str(payload.get("workshop_code") or "").strip().upper() or None
        workshop_name = str(payload.get("workshop_name") or workshop_code or "").strip() or workshop_code
        line_code = str(payload.get("line_code") or "").strip().upper() or None
        line_name = str(payload.get("line_name") or line_code or "").strip() or line_code
        company_code = str(payload.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
        actor = payload.get("actor") if isinstance(payload.get("actor"), dict) else {}
        if is_workshop_manager(actor):
            if not workshop_code or not line_code:
                raise bad_request(
                    code="REPORT_LINE_REQUIRED",
                    message="workshop_code and line_code are required for workshop manager reporting.",
                )
            assert_actor_can_access_line(
                self.connection,
                actor,
                company_code=company_code,
                workshop_code=workshop_code,
                line_code=line_code,
                missing_user_error_code="REPORT_ACTOR_USER_REQUIRED",
                forbidden_error_code="REPORT_LINE_SCOPE_FORBIDDEN",
                forbidden_message="Current workshop manager is not allowed to report on this line.",
            )
        with transaction(self.connection):
            self.connection.execute(
                """
                INSERT INTO work_reports (
                    report_id,
                    production_order_no,
                    report_scope,
                    process_code,
                    process_name,
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    report_id,
                    order_no,
                    normalized_report_scope,
                    process_code,
                    process_name_by_code.get(process_code, process_code),
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    report_time,
                ),
            )
            if order_no:
                self._sync_order_state_from_reporting(
                    order_no=order_no,
                    process_code=process_code,
                )
        if workshop_code and line_code:
            self.host.rebuild_line_daily_actual_capacity({"calendar_date": reporting_local_date})
        return {
            "report_id": report_id,
            "order_no": order_no,
            "report_scope": normalized_report_scope,
            "process_code": process_code,
            "process_name_cn": process_name_by_code.get(process_code, process_code),
            "report_qty": report_qty,
            "report_time": report_time,
            "operator_name_cn": operator_name,
        }

    def _sync_order_state_from_reporting(self, *, order_no: str, process_code: str) -> None:
        base_row = self._require_order(order_no)
        state_row = self._get_order_state(order_no) or {}
        product_code = str(base_row.get("material_code") or "").strip().upper()
        final_meta = self._resolve_final_process_meta_by_product({product_code}).get(product_code) or {}
        final_process_code = str(final_meta.get("process_code") or "").strip().upper()
        if not final_process_code or final_process_code != process_code:
            return

        final_report_row = fetch_one(
            self.connection,
            """
            SELECT COALESCE(SUM(report_qty), 0) AS total_qty
            FROM work_reports
            WHERE production_order_no = ?
              AND UPPER(TRIM(COALESCE(process_code, ''))) = ?
            """,
            (order_no, final_process_code),
        )
        completed_qty = _to_number((final_report_row or {}).get("total_qty"), 0)
        production_qty = _to_number(base_row.get("production_qty"), 0)
        remaining_qty = max(0.0, production_qty - completed_qty)
        progress_rate = (completed_qty / production_qty * 100) if production_qty > 0 else 0
        normalized_status = (
            "DONE"
            if production_qty > 0 and completed_qty + 1e-9 >= production_qty
            else "IN_PROGRESS"
            if completed_qty > 0
            else "OPEN"
        )
        next_row = {
            "production_order_no": order_no,
            "promised_due_date": state_row.get("promised_due_date") or base_row.get("planned_end_date"),
            "expected_start_date": state_row.get("expected_start_date") or base_row.get("planned_start_date"),
            "expected_start_time": state_row.get("expected_start_time")
            or self._iso_at(state_row.get("expected_start_date") or base_row.get("planned_start_date"), "08:00:00"),
            "expected_finish_time": state_row.get("expected_finish_time")
            or self._iso_at(state_row.get("promised_due_date") or base_row.get("planned_end_date"), "18:00:00"),
            "priority_level": self._normalize_priority_level(state_row.get("priority_level"), 5),
            "urgent_flag": 1 if self._normalize_priority_level(state_row.get("priority_level"), 5) == 1 else 0,
            "lock_flag": int(_to_number(state_row.get("lock_flag"), 0)),
            "frozen_flag": int(_to_number(state_row.get("frozen_flag"), 0)),
            "status": normalized_status,
            "order_status": normalized_status,
            "completed_qty": completed_qty,
            "remaining_qty": remaining_qty,
            "progress_rate": min(100.0, round(progress_rate, 4)),
            "production_batch_no": state_row.get("production_batch_no"),
        }
        self._upsert_order_state(next_row)

    def _require_order(self, order_no: str) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT
                production_order_no,
                material_code,
                material_name,
                material_specification,
                production_qty,
                status,
                planned_start_date,
                planned_end_date,
                source_bill_no,
                material_list_no,
                updated_at
            FROM production_orders
            WHERE production_order_no = ?
            """,
            (order_no,),
        )
        if row is None:
            raise not_found(
                code="ORDER_NOT_FOUND",
                message="Order does not exist.",
                details={"order_no": order_no},
            )
        return row

    def _get_order_state(self, order_no: str) -> dict[str, Any] | None:
        return fetch_one(
            self.connection,
            """
            SELECT
                production_order_no,
                promised_due_date,
                expected_start_date,
                expected_start_time,
                expected_finish_time,
                priority_level,
                urgent_flag,
                lock_flag,
                frozen_flag,
                status,
                order_status,
                completed_qty,
                remaining_qty,
                progress_rate,
                production_batch_no,
                updated_at
            FROM order_pool_state
            WHERE production_order_no = ?
            """,
            (order_no,),
        )

    def _upsert_order_state(self, row: dict[str, Any]) -> None:
        updated_at = utc_now()
        self.connection.execute(
            """
            INSERT INTO order_pool_state (
                production_order_no,
                promised_due_date,
                expected_start_date,
                expected_start_time,
                expected_finish_time,
                priority_level,
                urgent_flag,
                lock_flag,
                frozen_flag,
                status,
                order_status,
                completed_qty,
                remaining_qty,
                progress_rate,
                production_batch_no,
                updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(production_order_no) DO UPDATE SET
                promised_due_date = excluded.promised_due_date,
                expected_start_date = excluded.expected_start_date,
                expected_start_time = excluded.expected_start_time,
                expected_finish_time = excluded.expected_finish_time,
                priority_level = excluded.priority_level,
                urgent_flag = excluded.urgent_flag,
                lock_flag = excluded.lock_flag,
                frozen_flag = excluded.frozen_flag,
                status = excluded.status,
                order_status = excluded.order_status,
                completed_qty = excluded.completed_qty,
                remaining_qty = excluded.remaining_qty,
                progress_rate = excluded.progress_rate,
                production_batch_no = excluded.production_batch_no,
                updated_at = excluded.updated_at
            """,
            (
                row.get("production_order_no"),
                row.get("promised_due_date"),
                row.get("expected_start_date"),
                row.get("expected_start_time"),
                row.get("expected_finish_time"),
                row.get("priority_level"),
                row.get("urgent_flag"),
                row.get("lock_flag"),
                row.get("frozen_flag"),
                row.get("status"),
                row.get("order_status"),
                row.get("completed_qty"),
                row.get("remaining_qty"),
                row.get("progress_rate"),
                row.get("production_batch_no"),
                updated_at,
            ),
        )

    def _resolve_final_process_meta_by_product(
        self,
        product_codes: set[str],
    ) -> dict[str, dict[str, str]]:
        normalized_product_codes = {
            str(product_code or "").strip().upper()
            for product_code in product_codes
            if str(product_code or "").strip()
        }
        if not normalized_product_codes:
            return {}
        placeholders = ",".join("?" for _ in normalized_product_codes)
        rows = fetch_all(
            self.connection,
            f"""
            SELECT
                product_code,
                sequence_no,
                process_code,
                process_name_cn,
                COALESCE(is_final_process, 0) AS is_final_process
            FROM masterdata_process_routes
            WHERE product_code IN ({placeholders})
            ORDER BY product_code ASC, sequence_no ASC
            """,
            tuple(sorted(normalized_product_codes)),
        )
        rows_by_product: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            product_code = str(row.get("product_code") or "").strip().upper()
            process_code = str(row.get("process_code") or "").strip().upper()
            sequence_no = int(_to_number(row.get("sequence_no"), 0))
            if not product_code or not process_code or sequence_no <= 0:
                continue
            rows_by_product.setdefault(product_code, []).append(
                {
                    "sequence_no": sequence_no,
                    "process_code": process_code,
                    "process_name_cn": str(row.get("process_name_cn") or process_code).strip() or process_code,
                    "is_final_process": 1 if int(_to_number(row.get("is_final_process"), 0)) == 1 else 0,
                }
            )

        out: dict[str, dict[str, str]] = {}
        for product_code, route_rows in rows_by_product.items():
            ordered_rows = sorted(route_rows, key=lambda item: int(item["sequence_no"]))
            marked_rows = [row for row in ordered_rows if int(row.get("is_final_process") or 0) == 1]
            if len(marked_rows) > 1:
                raise server_error(
                    code="ROUTE_FINAL_PROCESS_DUPLICATED",
                    message="Route has more than one final process step.",
                    details={"product_code": product_code},
                )
            target_row = marked_rows[0] if marked_rows else ordered_rows[-1]
            process_code = str(target_row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            out[product_code] = {
                "process_code": process_code,
                "process_name_cn": str(target_row.get("process_name_cn") or process_code).strip() or process_code,
            }
        return out

    def _iso_at(self, date_text: str | None, clock_text: str) -> str | None:
        if not date_text:
            return None
        return f"{date_text}T{clock_text}+08:00"

    def _normalize_priority_level(self, value: object, default: int = 5) -> int:
        try:
            level = int(value)
        except (TypeError, ValueError):
            return default
        return max(1, min(5, level))

    def delete_reporting(
        self,
        report_id: str,
        actor: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        existing = fetch_one(
            self.connection,
            """
            SELECT report_id, process_code, workshop_code, line_code, report_time
            FROM work_reports
            WHERE report_id = ?
            """,
            (report_id,),
        )
        if existing is None:
            raise not_found(
                code="REPORT_NOT_FOUND",
                message="Report does not exist.",
                details={"report_id": report_id},
            )
        normalized_actor = actor if isinstance(actor, dict) else {}
        workshop_code = str(existing.get("workshop_code") or "").strip().upper()
        line_code = str(existing.get("line_code") or "").strip().upper()
        if is_workshop_manager(normalized_actor):
            if not workshop_code or not line_code:
                raise forbidden(
                    code="REPORT_LINE_SCOPE_FORBIDDEN",
                    message="Current workshop manager is not allowed to delete this reporting record.",
                    details={"report_id": report_id},
                )
            assert_actor_can_access_line(
                self.connection,
                normalized_actor,
                company_code="COMPANY-MAIN",
                workshop_code=workshop_code,
                line_code=line_code,
                missing_user_error_code="REPORT_ACTOR_USER_REQUIRED",
                forbidden_error_code="REPORT_LINE_SCOPE_FORBIDDEN",
                forbidden_message="Current workshop manager is not allowed to delete this reporting record.",
            )
        with transaction(self.connection):
            self.connection.execute(
                "DELETE FROM work_reports WHERE report_id = ?",
                (report_id,),
            )
        process_code = str(existing.get("process_code") or "").strip().upper()
        report_time_text = str(existing.get("report_time") or "").strip()
        if workshop_code and line_code and process_code and report_time_text:
            local_date = _local_date_from_iso_datetime(report_time_text)
            if local_date:
                self.host.rebuild_line_daily_actual_capacity({"calendar_date": local_date})
        return {"ok": True}
