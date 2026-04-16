from __future__ import annotations

import base64
import binascii
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
import hashlib
from io import BytesIO
import math
from pathlib import Path
import random
import sqlite3
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from ..db import fetch_all, fetch_one, transaction, utc_now
from ..errors import bad_request, forbidden, not_found, server_error
from ..gateway.inventory import ERPInventoryGateway
from ..gateway.masterdata import UpstreamMasterdataGateway
from ..gateway.orders import ERPOrderGateway
from ..gateway.supply import ERPSupplyGateway
from ..json_utils import dumps, loads
from ..repositories.backups import BackupRepository
from .final_process_metrics import build_order_final_process_metrics
from .job_dispatcher import ServiceFactory
from .line_daily_capacity_service import LineDailyCapacityService


RULES_SINGLETON_KEY = "default"
SHIFT_SEQUENCE = ("DAY", "NIGHT")
SHIFT_INDEX_BY_CODE = {"DAY": 0, "NIGHT": 1}
SHIFT_CODE_BY_INDEX = {0: "DAY", 1: "NIGHT"}
WEEKEND_REST_MODES = {"NONE", "SINGLE", "DOUBLE"}
DATE_SHIFT_MODES = {"REST", "DAY", "NIGHT", "BOTH"}
SUPPORTED_STRATEGY_CODES = {
    "KEY_ORDER_FIRST",
    "MAX_CAPACITY_FIRST",
    "MIN_DELAY_FIRST",
}
SUPPORTED_CAPACITY_SOURCE_MODES = {"DEFAULT", "PLANNED", "ACTUAL"}
SUPPORTED_CN_HOLIDAY_YEARS = {2024, 2025, 2026, 2027, 2028}
PRIORITY_LEVEL_MIN = 1
PRIORITY_LEVEL_MAX = 5
ROLE_SCHEDULER = "SCHEDULER"
ROLE_WORKSHOP_MANAGER = "WORKSHOP_MANAGER"
DEFAULT_COMPANY_CODE = "COMPANY-MAIN"
ANGIO_CATHETER_PRODUCT_CODES = frozenset(
    {
        "A006.034.10191",
        "YXN.009.020.1047",
        "YXN.044.02.1028",
        "YXN.067.005.1006",
        "A006.034.6104",
        "YXN.044.02.1020",
    }
)
DEFAULT_ANGIO_WORKSHOP_CODE = "1车间"
DEFAULT_ANGIO_WORKSHOP_NAME = "1车间"
DEFAULT_ANGIO_LINE_CODE = "1产线"
DEFAULT_ANGIO_LINE_NAME = "1产线"
CN_STATUTORY_HOLIDAY_DATE_SET = frozenset(
    {
        # 2024
        "2024-01-01",
        "2024-02-10",
        "2024-02-11",
        "2024-02-12",
        "2024-02-13",
        "2024-02-14",
        "2024-02-15",
        "2024-02-16",
        "2024-02-17",
        "2024-04-04",
        "2024-04-05",
        "2024-04-06",
        "2024-05-01",
        "2024-05-02",
        "2024-05-03",
        "2024-05-04",
        "2024-05-05",
        "2024-06-08",
        "2024-06-09",
        "2024-06-10",
        "2024-09-15",
        "2024-09-16",
        "2024-09-17",
        "2024-10-01",
        "2024-10-02",
        "2024-10-03",
        "2024-10-04",
        "2024-10-05",
        "2024-10-06",
        "2024-10-07",
        # 2025
        "2025-01-01",
        "2025-01-28",
        "2025-01-29",
        "2025-01-30",
        "2025-01-31",
        "2025-02-01",
        "2025-02-02",
        "2025-02-03",
        "2025-02-04",
        "2025-04-04",
        "2025-04-05",
        "2025-04-06",
        "2025-05-01",
        "2025-05-02",
        "2025-05-03",
        "2025-05-04",
        "2025-05-05",
        "2025-05-31",
        "2025-06-01",
        "2025-06-02",
        "2025-10-01",
        "2025-10-02",
        "2025-10-03",
        "2025-10-04",
        "2025-10-05",
        "2025-10-06",
        "2025-10-07",
        "2025-10-08",
        # 2026
        "2026-01-01",
        "2026-01-02",
        "2026-01-03",
        "2026-02-15",
        "2026-02-16",
        "2026-02-17",
        "2026-02-18",
        "2026-02-19",
        "2026-02-20",
        "2026-02-21",
        "2026-02-22",
        "2026-02-23",
        "2026-04-04",
        "2026-04-05",
        "2026-04-06",
        "2026-05-01",
        "2026-05-02",
        "2026-05-03",
        "2026-05-04",
        "2026-05-05",
        "2026-06-19",
        "2026-06-20",
        "2026-06-21",
        "2026-09-25",
        "2026-09-26",
        "2026-09-27",
        "2026-10-01",
        "2026-10-02",
        "2026-10-03",
        "2026-10-04",
        "2026-10-05",
        "2026-10-06",
        "2026-10-07",
    }
)
SCHEDULE_SLOT_SEARCH_GUARD = 20000
SCHEDULE_EPOCH_DAY = date(1970, 1, 1)
SCHEDULE_NUMBER_EPSILON = 1e-9
STATUS_NAME_BY_CODE = {
    "DRAFT": "Draft",
    "PUBLISHED": "Published",
    "ARCHIVED": "Archived",
}
LOCAL_ORDER_STATUS_CODES = frozenset(
    {
        "OPEN",
        "IN_PROGRESS",
        "DONE",
        "COMPLETED",
        "CLOSED",
        "DELAY",
    }
)
LOCAL_TIMEZONE = timezone(timedelta(hours=8))
WORK_REPORT_COLUMNS = (
    "report_id",
    "production_order_no",
    "process_code",
    "process_name",
    "company_code",
    "workshop_code",
    "workshop_name",
    "line_code",
    "line_name",
    "report_qty",
    "report_time",
    "operator_code",
    "operator_name",
    "section_leader_name",
    "dispatch_no",
    "product_code",
    "product_name",
    "product_specification",
    "resource_group_name",
    "resource_name",
    "department_name",
    "source_process_code",
    "source_process_name",
    "mold_code",
    "support_count",
    "weight_kg",
    "cavity_count",
    "total_cycle_time",
    "production_quota",
    "work_duration",
    "clamp_or_assembly_weight",
    "unit_weight",
    "source_sheet_name",
    "source_row_no",
    "source_file_name",
    "daily_capacity_compare_audit_id",
    "daily_capacity_compare_qty",
    "daily_capacity_compare_selected_at",
    "updated_at",
)
REPORTING_IMPORT_REQUIRED_HEADER_MAP = {
    "报工日期": "report_datetime",
    "报工人编码": "operator_code",
    "报工人名称": "operator_name",
    "工段长": "section_leader_name",
    "生产订单号": "production_order_no",
    "生产资源组": "resource_group_name",
    "生产资源": "resource_name",
    "派工单号": "dispatch_no",
    "产品编码": "product_code",
    "产品名称": "product_name",
    "规格": "product_specification",
    "工序编码": "source_process_code",
    "工序名称": "source_process_name",
    "所属部门": "department_name",
    "报工数量": "report_qty",
}
REPORTING_IMPORT_OPTIONAL_HEADER_MAP = {
    "模具编码": "mold_code",
    "支数": "support_count",
    "公斤数": "weight_kg",
    "实腔数": "cavity_count",
    "全程时间": "total_cycle_time",
    "生产定额": "production_quota",
    "工作时长": "work_duration",
    "注塑合模/组装公斤数": "clamp_or_assembly_weight",
    "注塑个数/组装个重": "unit_weight",
}


def _today_text() -> str:
    return date.today().isoformat()


def _normalize_date_text(value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def _parse_date_or_today(value: object) -> date:
    normalized = _normalize_date_text(value)
    if normalized is None:
        return date.today()
    return date.fromisoformat(normalized)


def _local_date_from_iso_datetime(value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None
    return parsed.astimezone(LOCAL_TIMEZONE).date().isoformat()


def _iso_at(date_text: str | None, clock_text: str) -> str | None:
    if not date_text:
        return None
    return f"{date_text}T{clock_text}+08:00"


def _days_between(start_text: object, end_text: object) -> int:
    start = _normalize_date_text(start_text)
    end = _normalize_date_text(end_text)
    if not start or not end:
        return 0
    return max(0, (date.fromisoformat(end) - date.fromisoformat(start)).days)


def _signed_days_between(start_text: object, end_text: object) -> int | None:
    start = _normalize_date_text(start_text)
    end = _normalize_date_text(end_text)
    if not start or not end:
        return None
    return (date.fromisoformat(end) - date.fromisoformat(start)).days


def _parse_local_iso_datetime(value: object) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None
    return parsed.astimezone(LOCAL_TIMEZONE)


def _due_deadline_datetime(date_text: object, clock_text: str = "18:00:00") -> datetime | None:
    normalized = _normalize_date_text(date_text)
    if not normalized:
        return None
    return _parse_local_iso_datetime(_iso_at(normalized, clock_text))


def _event_datetime(value: object, default_clock_text: str = "18:00:00") -> datetime | None:
    parsed = _parse_local_iso_datetime(value)
    if parsed is not None:
        return parsed
    normalized_date = _normalize_date_text(value)
    if not normalized_date:
        return None
    return _parse_local_iso_datetime(_iso_at(normalized_date, default_clock_text))


def _signed_due_gap_days(
    due_date_text: object,
    event_value: object,
    *,
    event_default_clock_text: str = "18:00:00",
) -> int | None:
    due_deadline = _due_deadline_datetime(due_date_text)
    event_time = _event_datetime(event_value, event_default_clock_text)
    if due_deadline is None or event_time is None:
        return None
    day_gap = (event_time.date() - due_deadline.date()).days
    if day_gap == 0 and event_time > due_deadline:
        return 1
    return day_gap


def _days_from_today(date_text: object) -> int | None:
    normalized = _normalize_date_text(date_text)
    if normalized is None:
        return None
    today = datetime.now(LOCAL_TIMEZONE).date()
    return (date.fromisoformat(normalized) - today).days


def _to_number(value: object, fallback: float = 0.0) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return fallback
    return number if number == number else fallback


def _normalize_dashboard_code(value: object) -> str:
    return str(value or "").strip().upper()


def _normalize_dashboard_name(value: object, fallback: str) -> str:
    text = str(value or "").strip()
    return text or fallback


def _build_dashboard_options(
    meta_by_code: dict[str, dict[str, str]],
    *,
    code_key: str,
    name_key: str,
) -> list[dict[str, str]]:
    def option_sort_key(item: tuple[str, dict[str, str]]) -> tuple[str, str]:
        code, meta = item
        name = str(meta.get(name_key) or code).strip() or code
        return (name.lower(), code)

    return [
        {
            code_key: code,
            name_key: str(meta.get(name_key) or code).strip() or code,
        }
        for code, meta in sorted(meta_by_code.items(), key=option_sort_key)
    ]


def _build_dashboard_change_items(
    calendar_dates: list[str],
    totals_by_code: dict[str, dict[str, float]],
    meta_by_code: dict[str, dict[str, str]],
    *,
    code_key: str,
    name_key: str,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for option in _build_dashboard_options(
        meta_by_code,
        code_key=code_key,
        name_key=name_key,
    ):
        code = str(option.get(code_key) or "").strip().upper()
        if not code:
            continue
        previous_planned_capacity_qty: float | None = None
        planned_by_date = totals_by_code.get(code) or {}
        for calendar_date in calendar_dates:
            planned_capacity_qty = round(
                _to_number(planned_by_date.get(calendar_date), 0),
                4,
            )
            planned_capacity_change_qty = (
                0
                if previous_planned_capacity_qty is None
                else round(planned_capacity_qty - previous_planned_capacity_qty, 4)
            )
            items.append(
                {
                    "calendar_date": calendar_date,
                    code_key: code,
                    name_key: str(option.get(name_key) or code).strip() or code,
                    "planned_capacity_qty": planned_capacity_qty,
                    "planned_capacity_change_qty": planned_capacity_change_qty,
                }
            )
            previous_planned_capacity_qty = planned_capacity_qty
    return items


def _normalize_priority_level(value: object, default: int = PRIORITY_LEVEL_MAX) -> int:
    try:
        level = int(value)
    except (TypeError, ValueError):
        return default
    return max(PRIORITY_LEVEL_MIN, min(PRIORITY_LEVEL_MAX, level))


def _parse_enabled_flag(value: object) -> int:
    if isinstance(value, bool):
        return 1 if value else 0
    if isinstance(value, int):
        if value in (0, 1):
            return value
        raise ValueError("enabled_flag must be 0 or 1.")
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "on"}:
        return 1
    if text in {"0", "false", "no", "off"}:
        return 0
    raise ValueError("enabled_flag must be 0/1 or a boolean-like value.")


def _parse_int_in_range(
    value: object,
    *,
    field_name: str,
    min_value: int,
    max_value: int,
) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field_name} must be an integer.")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"{field_name} must be an integer.")
        parsed = int(value)
    else:
        text = str(value or "").strip()
        if not text or any(ch not in "0123456789" for ch in text):
            raise ValueError(f"{field_name} must be an integer.")
        parsed = int(text)

    if parsed < min_value or parsed > max_value:
        raise ValueError(
            f"{field_name} must be between {min_value} and {max_value} (inclusive)."
        )
    return parsed


def _urgent_flag_from_priority_level(priority_level: object) -> int:
    return 1 if _normalize_priority_level(priority_level) == PRIORITY_LEVEL_MIN else 0


def _route_row_id(product_code: str, sequence_no: int) -> str:
    return f"route-{product_code}-{sequence_no}"


def _status_name(status: str) -> str:
    normalized = str(status or "").strip().upper()
    return STATUS_NAME_BY_CODE.get(normalized, normalized or "-")


def _normalize_shift_code(value: object) -> str:
    normalized = str(value or "").strip().upper()
    if normalized == "D":
        return "DAY"
    if normalized == "N":
        return "NIGHT"
    if normalized in SHIFT_INDEX_BY_CODE:
        return normalized
    return "DAY"


def _slot_index_for(date_value: date, shift_code: str) -> int:
    shift_index = SHIFT_INDEX_BY_CODE[_normalize_shift_code(shift_code)]
    day_index = (date_value - SCHEDULE_EPOCH_DAY).days
    return day_index * 2 + shift_index


def _slot_index_from_text(date_text: str, shift_code: str) -> int:
    return _slot_index_for(date.fromisoformat(date_text), shift_code)


def _slot_to_date_shift(slot_index: int) -> tuple[date, str]:
    if slot_index < 0:
        raise ValueError("slot_index must be >= 0")
    day_index = slot_index // 2
    shift_index = slot_index % 2
    return (
        SCHEDULE_EPOCH_DAY + timedelta(days=day_index),
        SHIFT_CODE_BY_INDEX[shift_index],
    )


def _material_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "order_no": row.get("production_order_no"),
        "child_material_code": row["child_material_code"],
        "child_material_name_cn": row["child_material_name"],
        "spec_model": row.get("spec_model"),
        "issue_qty": row.get("issue_qty"),
        "required_qty": row.get("issue_qty"),
        "usage_numerator": row.get("usage_numerator"),
        "usage_denominator": row.get("usage_denominator"),
        "child_unit": row.get("child_unit"),
        "child_material_supply_type": row["supply_type_code"],
        "child_material_supply_type_name_cn": row["supply_type_name"],
        "child_material_inventory_qty": row.get("inventory_qty"),
        "child_material_inventory_status": row["inventory_status"],
    }


def _expected_start_shift_from_datetime_text(value: object) -> str:
    timestamp_text = str(value or "").strip().upper()
    if not timestamp_text:
        return "DAY"
    if "T" in timestamp_text:
        time_text = timestamp_text.split("T", 1)[1][:8]
        try:
            hour = int(time_text.split(":", 1)[0])
        except (TypeError, ValueError):
            hour = None
        if hour is not None and (hour >= 20 or hour < 8):
            return "NIGHT"
    return "DAY"


def _schedule_version_status_label(status: object) -> str:
    normalized = str(status or "").strip().upper()
    if normalized == "PUBLISHED":
        return "已发布"
    if normalized == "DRAFT":
        return "草稿"
    if normalized == "ARCHIVED":
        return "已归档"
    return normalized or "-"


class AppService:
    def __init__(self, connection: sqlite3.Connection) -> None:
        self.connection = connection
        self.factory = ServiceFactory(connection)
        self.line_daily_capacity_service = LineDailyCapacityService(connection)
        self.masterdata_gateway = UpstreamMasterdataGateway()
        self.order_gateway = ERPOrderGateway()
        self.inventory_gateway = ERPInventoryGateway()
        self.supply_gateway = ERPSupplyGateway()

    def list_order_pool(self, *, version_no: str | None = None) -> dict[str, Any]:
        rows = self._list_order_rows()
        order_nos = [str(row["production_order_no"]) for row in rows]
        states = self._get_order_state_map(order_nos)
        capacity_map = self._get_capacity_map(order_nos)
        self._ensure_masterdata_seeded()
        reference_version_no = self._resolve_order_pool_version_no(version_no)
        reference_schedule_context = self._build_reference_schedule_context(reference_version_no)
        published_version_no = self._pick_published_schedule_version_no()
        published_schedule_context = (
            reference_schedule_context
            if published_version_no and published_version_no == reference_version_no
            else self._build_reference_schedule_context(published_version_no)
        )
        draft_version_no = self._pick_latest_draft_schedule_version_no()
        draft_version = self.get_schedule_version(draft_version_no) if draft_version_no else None
        shortage_analysis = self._build_schedule_shortage_analysis(reference_version_no)
        final_process_metrics_by_order = build_order_final_process_metrics(
            self.connection,
            rows,
        )
        route_rows = self._list_route_rows()
        route_map = self._group_routes_by_product(route_rows)
        topology_by_process = self._first_topology_by_process(self._list_line_topology_rows())
        items: list[dict[str, Any]] = []
        for row in rows:
            order_no = str(row["production_order_no"])
            items.append(
                self._build_order_pool_row(
                base_row=row,
                state_row=states.get(order_no),
                capacity_rows=capacity_map.get(order_no, []),
                route_rows=route_map.get(str(row["material_code"]), []),
                topology_by_process=topology_by_process,
                reference_version=reference_schedule_context.get("version"),
                schedule_fact=reference_schedule_context["order_map"].get(order_no),
                published_version=published_schedule_context.get("version"),
                published_schedule_fact=published_schedule_context["order_map"].get(order_no),
                shortage_summary=shortage_analysis["order_map"].get(order_no),
                final_process_metrics=final_process_metrics_by_order.get(order_no),
            )
            )
        reference_version = reference_schedule_context.get("version")
        reference_version_status = str((reference_version or {}).get("status") or "").strip().upper() or None
        reference_version_status_label = _schedule_version_status_label(reference_version_status)
        published_version = published_schedule_context.get("version")
        published_version_status = str((published_version or {}).get("status") or "").strip().upper() or None
        return {
            "reference_version_no": reference_version_no,
            "current_view_version_no": reference_version_no,
            "current_view_version_status": reference_version_status,
            "current_view_version_status_label": reference_version_status_label,
            "current_view_version_label": (
                f"当前查看版 {reference_version_no}（{reference_version_status_label}）"
                if reference_version_no
                else "当前查看版：未选择"
            ),
            "published_version_no": published_version_no,
            "published_version_status": published_version_status,
            "published_version_status_label": _schedule_version_status_label(published_version_status),
            "published_version_label": (
                f"正式执行版 {published_version_no}"
                if published_version_no
                else "正式执行版：未发布"
            ),
            "draft_version_no": draft_version_no,
            "draft_version_status": str((draft_version or {}).get("status") or "").strip().upper() or None,
            "draft_version_status_label": _schedule_version_status_label((draft_version or {}).get("status")),
            "draft_version_label": (
                f"草稿版 {draft_version_no}"
                if draft_version_no
                else "草稿版：暂无"
            ),
            "items": items,
        }

    def get_order_pool_item(self, order_no: str, *, version_no: str | None = None) -> dict[str, Any]:
        base_row = self._require_order(order_no)
        state_row = self._get_order_state(order_no)
        self._ensure_masterdata_seeded()
        normalized_order_no = str(base_row["production_order_no"])
        reference_version_no = self._resolve_order_pool_version_no(version_no)
        reference_schedule_context = self._build_reference_schedule_context(reference_version_no)
        published_version_no = self._pick_published_schedule_version_no()
        published_schedule_context = (
            reference_schedule_context
            if published_version_no and published_version_no == reference_version_no
            else self._build_reference_schedule_context(published_version_no)
        )
        shortage_analysis = self._build_schedule_shortage_analysis(reference_version_no)
        final_process_metrics = build_order_final_process_metrics(
            self.connection,
            [base_row],
        ).get(normalized_order_no)
        route_rows = self._group_routes_by_product(self._list_route_rows()).get(
            str(base_row["material_code"]),
            [],
        )
        topology_by_process = self._first_topology_by_process(self._list_line_topology_rows())
        capacity_rows = self._get_capacity_rows(order_no)
        return self._build_order_pool_row(
            base_row=base_row,
            state_row=state_row,
            capacity_rows=capacity_rows,
            route_rows=route_rows,
            topology_by_process=topology_by_process,
            reference_version=reference_schedule_context.get("version"),
            schedule_fact=reference_schedule_context["order_map"].get(normalized_order_no),
            published_version=published_schedule_context.get("version"),
            published_schedule_fact=published_schedule_context["order_map"].get(normalized_order_no),
            shortage_summary=shortage_analysis["order_map"].get(normalized_order_no),
            final_process_metrics=final_process_metrics,
        )

    def get_order_pool_process_timeline(
        self,
        order_no: str,
        *,
        process_code: str | None = None,
        version_no: str | None = None,
    ) -> dict[str, Any]:
        base_row = self._require_order(order_no)
        normalized_order_no = str(base_row["production_order_no"]).strip()
        normalized_product_code = str(base_row.get("material_code") or "").strip().upper()
        normalized_process_code = str(process_code or "").strip().upper() or None
        if process_code is not None and not normalized_process_code:
            raise bad_request(
                code="PROCESS_CODE_REQUIRED",
                message="process_code must be non-empty when provided.",
            )

        reference_version_no = self._resolve_order_pool_version_no(version_no)
        if reference_version_no is None:
            return {
                "summary": {
                    "reference_version_no": None,
                    "bottleneck_process_code": None,
                    "bottleneck_process_name_cn": None,
                    "message": "暂无排产任务数据",
                },
                "process_items": [],
                "selected_process_detail": self._build_empty_selected_process_detail(
                    process_code=normalized_process_code,
                ),
            }

        order_task_rows = self._list_schedule_task_rows_by_version_order(
            version_no=reference_version_no,
            order_no=normalized_order_no,
        )
        if len(order_task_rows) == 0:
            return {
                "summary": {
                    "reference_version_no": reference_version_no,
                    "bottleneck_process_code": None,
                    "bottleneck_process_name_cn": None,
                    "message": "暂无排产任务数据",
                },
                "process_items": [],
                "selected_process_detail": self._build_empty_selected_process_detail(
                    process_code=normalized_process_code,
                ),
            }

        route_rows = self._group_routes_by_product(self._list_route_rows()).get(
            normalized_product_code,
            [],
        )
        process_stats_by_code = self._build_schedule_process_stats_by_code(
            self._list_schedule_task_detail_rows(reference_version_no),
        )
        process_items = self._build_order_process_timeline_items(
            task_rows=order_task_rows,
            route_rows=route_rows,
            process_stats_by_code=process_stats_by_code,
        )
        bottleneck_item = next(
            (item for item in process_items if int(item.get("is_bottleneck") or 0) == 1),
            None,
        )
        selected_process_detail = self._build_empty_selected_process_detail(
            process_code=normalized_process_code,
        )
        if normalized_process_code:
            selected_process_detail = self._build_selected_process_timeline_detail(
                version_no=reference_version_no,
                order_no=normalized_order_no,
                process_code=normalized_process_code,
                process_items=process_items,
            )

        return {
            "summary": {
                "reference_version_no": reference_version_no,
                "bottleneck_process_code": (
                    str(bottleneck_item.get("process_code") or "").strip().upper()
                    if bottleneck_item
                    else None
                ),
                "bottleneck_process_name_cn": (
                    str(
                        bottleneck_item.get("process_name_cn")
                        or bottleneck_item.get("process_code")
                        or ""
                    ).strip()
                    if bottleneck_item
                    else None
                ),
            },
            "process_items": process_items,
            "selected_process_detail": selected_process_detail,
        }

    def list_order_pool_materials(
        self,
        order_no: str,
        *,
        refresh: bool = False,
    ) -> dict[str, Any]:
        if refresh:
            self.factory.build_material_refresh_service().refresh_order_materials(order_no)
        items = self.factory.build_material_query_service().list_order_materials(order_no)
        return {"items": [_material_row(item) for item in items]}

    def list_material_children(
        self,
        parent_material_code: str,
        *,
        refresh: bool = False,
    ) -> dict[str, Any]:
        if refresh:
            self._refresh_bom_children(parent_material_code)
        items = self.factory.build_material_query_service().list_bom_children(
            parent_material_code
        )
        return {"items": [_material_row(item) for item in items]}

    def patch_order_pool_order(
        self,
        order_no: str,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        base_row = self._require_order(order_no)
        current = self._get_order_state(order_no) or {}
        next_row = {
            "production_order_no": order_no,
            "promised_due_date": current.get("promised_due_date"),
            "expected_start_date": current.get("expected_start_date"),
            "expected_start_time": current.get("expected_start_time"),
            "expected_finish_time": current.get("expected_finish_time"),
            "priority_level": _normalize_priority_level(current.get("priority_level"), PRIORITY_LEVEL_MAX),
            "urgent_flag": _urgent_flag_from_priority_level(current.get("priority_level")),
            "lock_flag": int(current.get("lock_flag") or 0),
            "frozen_flag": int(current.get("frozen_flag") or 0),
            "status": current.get("status") or base_row.get("status"),
            "order_status": current.get("order_status") or current.get("status") or base_row.get("status"),
            "completed_qty": current.get("completed_qty"),
            "remaining_qty": current.get("remaining_qty"),
            "progress_rate": current.get("progress_rate"),
            "production_batch_no": current.get("production_batch_no"),
        }

        expected_start_date = _normalize_date_text(payload.get("expected_start_date"))
        expected_start_shift = str(payload.get("expected_start_shift") or "").strip().upper()
        if expected_start_shift and expected_start_shift not in {"DAY", "NIGHT"}:
            raise bad_request(
                code="ORDER_EXPECTED_START_SHIFT_INVALID",
                message="expected_start_shift must be DAY or NIGHT.",
            )
        if expected_start_date:
            next_row["expected_start_date"] = expected_start_date
            next_row["expected_start_time"] = _iso_at(
                expected_start_date,
                "20:00:00" if expected_start_shift == "NIGHT" else "08:00:00",
            )
        elif expected_start_shift:
            current_start_date = _normalize_date_text(current.get("expected_start_date"))
            if current_start_date is None:
                raise bad_request(
                    code="ORDER_EXPECTED_START_DATE_REQUIRED",
                    message="expected_start_date is required when expected_start_shift is provided.",
                )
            next_row["expected_start_time"] = _iso_at(
                current_start_date,
                "20:00:00" if expected_start_shift == "NIGHT" else "08:00:00",
            )

        if "priority_level" in payload:
            next_row["priority_level"] = _normalize_priority_level(payload.get("priority_level"), PRIORITY_LEVEL_MAX)
            next_row["urgent_flag"] = _urgent_flag_from_priority_level(next_row["priority_level"])
        elif "urgent_flag" in payload:
            next_row["priority_level"] = PRIORITY_LEVEL_MIN if int(payload["urgent_flag"] or 0) == 1 else PRIORITY_LEVEL_MAX
            next_row["urgent_flag"] = _urgent_flag_from_priority_level(next_row["priority_level"])

        for field in ("lock_flag", "frozen_flag"):
            if field in payload:
                next_row[field] = int(payload[field] or 0)

        status_value = str(payload.get("status") or "").strip().upper()
        order_status_value = str(payload.get("order_status") or "").strip().upper()
        if status_value:
            next_row["status"] = status_value
        if order_status_value:
            next_row["order_status"] = order_status_value

        if status_value == "DONE" or order_status_value == "DONE" or int(_to_number(payload.get("revoke_done_flag"), 0)) == 1:
            raise bad_request(
                code="ORDER_MANUAL_COMPLETION_FORBIDDEN",
                message="订单不能通过手工补丁直接完工或撤销完工，请使用末道报工更新订单状态。",
            )

        production_qty = _to_number(base_row.get("production_qty"), 0)
        normalized_status = str(
            next_row.get("order_status") or next_row.get("status") or ""
        ).strip().upper()
        if normalized_status in {"DONE", "COMPLETED"}:
            raise bad_request(
                code="ORDER_MANUAL_COMPLETION_FORBIDDEN",
                message="订单不能通过手工补丁直接完工，请使用末道报工更新订单状态。",
            )
        next_row["remaining_qty"] = max(0, _to_number(next_row.get("remaining_qty"), production_qty))
        next_row["progress_rate"] = max(0, _to_number(next_row.get("progress_rate"), 0))

        if "production_batch_no" in payload:
            next_row["production_batch_no"] = str(payload.get("production_batch_no") or "").strip() or None

        with transaction(self.connection):
            self._upsert_order_state(next_row)
        result = self.get_order_pool_item(order_no)
        result["save_impact"] = self._build_expected_start_save_impact(result)
        return result

    def delete_order_pool_order(self, order_no: str) -> dict[str, Any]:
        self._require_order(order_no)
        raise bad_request(
            code="ORDER_DELETE_FORBIDDEN",
            message="生产订单不允许物理删除，请使用 ERP 同步失效/关闭或后续审批能力处理。",
        )

    def list_mes_reportings(
        self,
        *,
        start_time: str | None = None,
        end_time: str | None = None,
        current_user: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        where_clauses: list[str] = []
        parameters: list[Any] = []
        if start_time:
            where_clauses.append("report_time >= ?")
            parameters.append(start_time)
        if end_time:
            where_clauses.append("report_time <= ?")
            parameters.append(end_time)
        manager_user_id = self._resolve_manager_user_id(current_user)
        if manager_user_id is not None:
            where_clauses.append(
                """
                EXISTS (
                    SELECT 1
                    FROM app_user_line_scopes scope
                    WHERE scope.user_id = ?
                      AND scope.workshop_code = work_reports.workshop_code
                      AND scope.line_code = work_reports.line_code
                )
                """
            )
            parameters.append(manager_user_id)
        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
        rows = fetch_all(
            self.connection,
            f"""
            SELECT
                report_id,
                production_order_no,
                report_scope,
                process_code,
                process_name,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                report_qty,
                report_time,
                operator_code,
                operator_name,
                section_leader_name,
                updated_at,
                date(report_time, '+8 hours') AS report_local_date,
                daily_capacity_compare_audit_id,
                daily_capacity_compare_qty
            FROM work_reports
            {where_sql}
            ORDER BY report_time DESC, report_id DESC
            """,
            tuple(parameters),
        )
        return {
            "items": [
                {
                    "report_id": row["report_id"],
                    "order_no": row["production_order_no"],
                    "process_code": row.get("process_code"),
                    "process_name_cn": row.get("process_name") or row.get("process_code"),
                    "workshop_code": row.get("workshop_code"),
                    "workshop_name": row.get("workshop_name") or row.get("workshop_code"),
                    "line_code": row.get("line_code"),
                    "line_name": row.get("line_name") or row.get("line_code"),
                    "report_qty": row.get("report_qty"),
                    "report_time": row.get("report_time"),
                    "report_local_date": _normalize_date_text(row.get("report_local_date")),
                    "operator_code": row.get("operator_code"),
                    "operator_name_cn": row.get("operator_name") or "系统填报",
                    "section_leader_name": row.get("section_leader_name"),
                    "updated_at": row.get("updated_at"),
                    "daily_capacity_compare_audit_id": row.get("daily_capacity_compare_audit_id"),
                    "daily_capacity_compare_qty": row.get("daily_capacity_compare_qty"),
                }
                for row in rows
            ]
        }

    def list_reporting_import_files(self, *, limit: int = 50) -> dict[str, Any]:
        normalized_limit = max(1, min(200, int(limit or 50)))
        rows = fetch_all(
            self.connection,
            """
            SELECT
                file_sha256,
                source_file_name,
                file_path,
                original_file_name,
                file_size_bytes,
                sheet_names_json,
                imported_by_user_id,
                imported_by_username,
                imported_by_display_name,
                total_row_count,
                imported_count,
                skipped_existing_count,
                failed_count,
                created_missing_order_count,
                created_at,
                last_imported_at
            FROM reporting_import_files
            ORDER BY last_imported_at DESC, file_sha256 DESC
            LIMIT ?
            """,
            (normalized_limit,),
        )
        items: list[dict[str, Any]] = []
        for row in rows:
            sheet_names = row.get("sheet_names_json")
            parsed_sheet_names = (
                loads(sheet_names)
                if isinstance(sheet_names, str) and sheet_names.strip()
                else []
            )
            items.append({**row, "sheet_names": parsed_sheet_names})
        return {"items": items, "total": len(items)}

    def get_order_summary(
        self,
        *,
        start_date: str,
        end_date: str,
        workshop_manager_user_id: str | None = None,
        current_user: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        normalized_start_date = _normalize_date_text(start_date)
        if normalized_start_date is None:
            raise bad_request(
                code="ORDER_SUMMARY_START_DATE_REQUIRED",
                message="start_date must be a valid YYYY-MM-DD date.",
            )
        normalized_end_date = _normalize_date_text(end_date)
        if normalized_end_date is None:
            raise bad_request(
                code="ORDER_SUMMARY_END_DATE_REQUIRED",
                message="end_date must be a valid YYYY-MM-DD date.",
            )
        if normalized_end_date < normalized_start_date:
            raise bad_request(
                code="ORDER_SUMMARY_DATE_RANGE_INVALID",
                message="end_date must be greater than or equal to start_date.",
                details={
                    "start_date": normalized_start_date,
                    "end_date": normalized_end_date,
                },
            )

        filters: list[str] = [
            "date(work_reports.report_time, '+8 hours') >= ?",
            "date(work_reports.report_time, '+8 hours') <= ?",
        ]
        parameters: list[Any] = [normalized_start_date, normalized_end_date]
        manager_user_id = self._resolve_order_summary_scope_manager_user_id(
            current_user=current_user,
            workshop_manager_user_id=workshop_manager_user_id,
        )
        if manager_user_id is not None:
            filters.append(
                """
                EXISTS (
                    SELECT 1
                    FROM app_user_line_scopes scope
                    WHERE scope.user_id = ?
                      AND scope.workshop_code = work_reports.workshop_code
                      AND scope.line_code = work_reports.line_code
                )
                """
            )
            parameters.append(manager_user_id)

        where_sql = " AND ".join(filters)
        grouped_report_rows = fetch_all(
            self.connection,
            f"""
            SELECT
                production_order_no,
                process_code,
                process_name,
                SUM(report_qty) AS report_qty,
                MAX(date(report_time, '+8 hours')) AS last_report_local_date
            FROM work_reports
            WHERE {where_sql}
            GROUP BY production_order_no, process_code
            ORDER BY production_order_no ASC, process_code ASC
            """,
            tuple(parameters),
        )

        process_name_by_code = self._process_name_by_code()
        order_process_qty_map: dict[str, dict[str, float]] = defaultdict(dict)
        order_process_last_date_map: dict[str, dict[str, str | None]] = defaultdict(dict)
        process_agg_map: dict[str, dict[str, Any]] = {}
        order_nos: set[str] = set()
        for row in grouped_report_rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            process_name = (
                str(row.get("process_name") or process_name_by_code.get(process_code) or process_code).strip()
                or process_code
            )
            report_qty = _to_number(row.get("report_qty"), 0)
            last_report_local_date = _normalize_date_text(row.get("last_report_local_date"))
            order_no = str(row.get("production_order_no") or "").strip()

            process_agg = process_agg_map.get(process_code)
            if process_agg is None:
                process_agg = {
                    "process_code": process_code,
                    "process_name_cn": process_name,
                    "report_qty_total": 0.0,
                    "order_nos": set(),
                    "last_report_date": None,
                }
                process_agg_map[process_code] = process_agg
            process_agg["report_qty_total"] = _to_number(process_agg.get("report_qty_total"), 0) + report_qty
            if last_report_local_date and (
                process_agg["last_report_date"] is None or last_report_local_date > process_agg["last_report_date"]
            ):
                process_agg["last_report_date"] = last_report_local_date

            if not order_no:
                continue
            order_nos.add(order_no)
            process_agg["order_nos"].add(order_no)
            order_process_qty_map[order_no][process_code] = report_qty
            order_process_last_date_map[order_no][process_code] = last_report_local_date

        if not order_nos and not process_agg_map:
            return {
                "range": {
                    "start_date": normalized_start_date,
                    "end_date": normalized_end_date,
                },
                "summary": {
                    "order_count": 0,
                    "completed_order_count": 0,
                    "completion_rate": 0,
                    "order_qty_total": 0,
                    "final_process_completed_qty_total": 0,
                    "process_count": 0,
                    "process_report_qty_total": 0,
                },
                "order_items": [],
                "process_items": [],
            }

        order_meta_by_no: dict[str, dict[str, Any]] = {}
        if order_nos:
            placeholders = ",".join("?" for _ in order_nos)
            order_rows = fetch_all(
                self.connection,
                f"""
                SELECT
                    production_order_no,
                    material_code,
                    material_name,
                    production_qty
                FROM production_orders
                WHERE production_order_no IN ({placeholders})
                """,
                tuple(sorted(order_nos)),
            )
            for row in order_rows:
                order_no = str(row.get("production_order_no") or "").strip()
                if not order_no:
                    continue
                order_meta_by_no[order_no] = row

        product_codes = {
            str((order_meta_by_no.get(order_no) or {}).get("material_code") or "").strip().upper()
            for order_no in order_nos
        }
        product_codes.discard("")
        final_process_by_product = self._resolve_final_process_meta_by_product(product_codes)

        order_items: list[dict[str, Any]] = []
        for order_no in sorted(order_nos):
            order_meta = order_meta_by_no.get(order_no) or {}
            product_code = str(order_meta.get("material_code") or "").strip().upper()
            product_name = str(order_meta.get("material_name") or product_code or "-").strip() or "-"
            order_qty = _to_number(order_meta.get("production_qty"), 0)
            final_meta = final_process_by_product.get(product_code) or {}
            final_process_code = str(final_meta.get("process_code") or "").strip().upper() or None
            final_process_name_cn = str(final_meta.get("process_name_cn") or final_process_code or "-").strip() or "-"
            process_qty_by_code = order_process_qty_map.get(order_no) or {}
            process_last_date_by_code = order_process_last_date_map.get(order_no) or {}
            final_process_completed_qty = (
                _to_number(process_qty_by_code.get(final_process_code), 0) if final_process_code else 0.0
            )
            final_process_last_report_date = (
                _normalize_date_text(process_last_date_by_code.get(final_process_code))
                if final_process_code
                else None
            )
            completed_flag = (
                bool(final_process_code)
                and order_qty > 0
                and (final_process_completed_qty + SCHEDULE_NUMBER_EPSILON) >= order_qty
            )
            order_items.append(
                {
                    "order_no": order_no,
                    "product_code": product_code or "-",
                    "product_name_cn": product_name,
                    "order_qty": order_qty,
                    "reported_process_count": len(process_qty_by_code),
                    "final_process_code": final_process_code,
                    "final_process_name_cn": final_process_name_cn,
                    "final_process_completed_qty": final_process_completed_qty,
                    "final_process_last_report_date": final_process_last_report_date,
                    "completed_flag": 1 if completed_flag else 0,
                }
            )

        process_items: list[dict[str, Any]] = []
        for process_code in sorted(process_agg_map):
            process_agg = process_agg_map[process_code]
            order_no_set = set(process_agg.get("order_nos") or set())
            involved_order_count = len(order_no_set)
            completed_order_count = 0
            if involved_order_count > 0:
                for order_no in order_no_set:
                    order_meta = order_meta_by_no.get(order_no) or {}
                    order_qty = _to_number(order_meta.get("production_qty"), 0)
                    process_report_qty = _to_number(
                        (order_process_qty_map.get(order_no) or {}).get(process_code),
                        0,
                    )
                    if order_qty > 0 and (process_report_qty + SCHEDULE_NUMBER_EPSILON) >= order_qty:
                        completed_order_count += 1
            completion_rate = (
                round(completed_order_count / involved_order_count * 100, 2)
                if involved_order_count > 0
                else 0
            )
            process_items.append(
                {
                    "process_code": process_code,
                    "process_name_cn": str(process_agg.get("process_name_cn") or process_code).strip()
                    or process_code,
                    "report_qty_total": _to_number(process_agg.get("report_qty_total"), 0),
                    "involved_order_count": involved_order_count,
                    "completed_order_count": completed_order_count,
                    "completion_rate": completion_rate,
                    "last_report_date": _normalize_date_text(process_agg.get("last_report_date")),
                }
            )

        order_count = len(order_items)
        completed_order_count = sum(1 for item in order_items if int(item.get("completed_flag") or 0) == 1)
        order_completion_rate = (
            round(completed_order_count / order_count * 100, 2) if order_count > 0 else 0
        )
        order_qty_total = round(sum(_to_number(item.get("order_qty"), 0) for item in order_items), 4)
        final_process_completed_qty_total = round(
            sum(_to_number(item.get("final_process_completed_qty"), 0) for item in order_items),
            4,
        )
        process_report_qty_total = round(
            sum(_to_number(item.get("report_qty_total"), 0) for item in process_items),
            4,
        )

        return {
            "range": {
                "start_date": normalized_start_date,
                "end_date": normalized_end_date,
            },
            "summary": {
                "order_count": order_count,
                "completed_order_count": completed_order_count,
                "completion_rate": order_completion_rate,
                "order_qty_total": order_qty_total,
                "final_process_completed_qty_total": final_process_completed_qty_total,
                "process_count": len(process_items),
                "process_report_qty_total": process_report_qty_total,
            },
            "order_items": order_items,
            "process_items": process_items,
        }

    def get_scheduler_dashboard(
        self,
        *,
        start_date: str,
        end_date: str,
        top_n: int = 8,
        current_user: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        normalized_start_date = _normalize_date_text(start_date)
        if normalized_start_date is None:
            raise bad_request(
                code="DASHBOARD_START_DATE_REQUIRED",
                message="start_date must be a valid YYYY-MM-DD date.",
            )
        normalized_end_date = _normalize_date_text(end_date)
        if normalized_end_date is None:
            raise bad_request(
                code="DASHBOARD_END_DATE_REQUIRED",
                message="end_date must be a valid YYYY-MM-DD date.",
            )
        if normalized_end_date < normalized_start_date:
            raise bad_request(
                code="DASHBOARD_DATE_RANGE_INVALID",
                message="end_date must be greater than or equal to start_date.",
                details={
                    "start_date": normalized_start_date,
                    "end_date": normalized_end_date,
                },
            )

        normalized_top_n = int(_to_number(top_n, 0))
        if normalized_top_n < 1 or normalized_top_n > 50:
            raise bad_request(
                code="DASHBOARD_TOP_N_INVALID",
                message="top_n must be between 1 and 50.",
            )

        topology_count_row = fetch_one(
            self.connection,
            """
            SELECT COUNT(1) AS total
            FROM masterdata_line_topology
            """,
        )
        topology_count = int(_to_number((topology_count_row or {}).get("total"), 0))
        if topology_count <= 0:
            raise server_error(
                code="DASHBOARD_TOPOLOGY_EMPTY",
                message="masterdata_line_topology contains no rows.",
            )

        order_summary = self.get_order_summary(
            start_date=normalized_start_date,
            end_date=normalized_end_date,
            current_user=current_user,
        )
        order_items = (
            order_summary.get("order_items")
            if isinstance(order_summary.get("order_items"), list)
            else []
        )
        if len(order_items) == 0:
            raise bad_request(
                code="DASHBOARD_ORDER_DATA_EMPTY",
                message="No order summary rows exist for the requested date range.",
                details={
                    "start_date": normalized_start_date,
                    "end_date": normalized_end_date,
                },
            )

        start_date_value = date.fromisoformat(normalized_start_date)
        end_date_value = date.fromisoformat(normalized_end_date)
        current_date_value = start_date_value
        calendar_dates: list[str] = []
        previous_planned_capacity_qty: float | None = None
        daily_capacity_items: list[dict[str, Any]] = []
        total_default_capacity_qty = 0.0
        total_planned_capacity_qty = 0.0
        total_actual_capacity_qty = 0.0
        failure_row_count = 0
        evaluated_row_count = 0
        line_meta_by_code: dict[str, dict[str, str]] = {}
        process_meta_by_code: dict[str, dict[str, str]] = {}
        line_planned_capacity_by_code: dict[str, dict[str, float]] = defaultdict(dict)
        process_planned_capacity_by_code: dict[str, dict[str, float]] = defaultdict(dict)
        line_daily_stats_by_key: dict[tuple[str, str], dict[str, Any]] = {}
        process_daily_stats_by_key: dict[tuple[str, str], dict[str, Any]] = {}
        daily_pressure_items: list[dict[str, Any]] = []
        process_name_by_code = self._process_name_by_code()

        while current_date_value <= end_date_value:
            calendar_date = current_date_value.isoformat()
            calendar_dates.append(calendar_date)
            capacity_payload = self.list_line_daily_capacity(calendar_date)
            capacity_rows = (
                capacity_payload.get("items")
                if isinstance(capacity_payload.get("items"), list)
                else []
            )
            if len(capacity_rows) == 0:
                raise bad_request(
                    code="DASHBOARD_CAPACITY_DATA_EMPTY",
                    message="No daily capacity rows exist for the requested date.",
                    details={"calendar_date": calendar_date},
                )

            day_default_capacity_qty = 0.0
            day_planned_capacity_qty = 0.0
            day_actual_capacity_qty = 0.0
            day_failure_row_count = 0
            day_line_planned_capacity: dict[str, float] = defaultdict(float)
            day_process_planned_capacity: dict[str, float] = defaultdict(float)
            day_line_default_capacity: dict[str, float] = defaultdict(float)
            day_line_actual_capacity: dict[str, float] = defaultdict(float)
            day_process_default_capacity: dict[str, float] = defaultdict(float)
            day_process_actual_capacity: dict[str, float] = defaultdict(float)
            for row in capacity_rows:
                default_capacity_qty = _to_number(row.get("default_capacity_qty"), 0)
                planned_capacity_qty = _to_number(row.get("planned_capacity_qty"), 0)
                actual_capacity_qty = _to_number(row.get("actual_capacity_qty"), 0)
                machine_count = int(_to_number(row.get("machine_count"), 0))
                required_machines = int(_to_number(row.get("required_machines"), 0))
                line_code = _normalize_dashboard_code(row.get("line_code"))
                if line_code:
                    line_name = _normalize_dashboard_name(row.get("line_name"), line_code)
                    line_meta_by_code.setdefault(
                        line_code,
                        {
                            "line_code": line_code,
                            "line_name": line_name,
                        },
                    )
                    day_line_planned_capacity[line_code] += planned_capacity_qty
                    day_line_default_capacity[line_code] += default_capacity_qty
                    day_line_actual_capacity[line_code] += actual_capacity_qty
                process_code = _normalize_dashboard_code(row.get("process_code"))
                if process_code:
                    process_name_cn = _normalize_dashboard_name(
                        process_name_by_code.get(process_code),
                        process_code,
                    )
                    process_meta_by_code.setdefault(
                        process_code,
                        {
                            "process_code": process_code,
                            "process_name_cn": process_name_cn,
                        },
                    )
                    day_process_planned_capacity[process_code] += planned_capacity_qty
                    day_process_default_capacity[process_code] += default_capacity_qty
                    day_process_actual_capacity[process_code] += actual_capacity_qty
                day_default_capacity_qty += default_capacity_qty
                day_planned_capacity_qty += planned_capacity_qty
                day_actual_capacity_qty += actual_capacity_qty
                if required_machines > 0:
                    evaluated_row_count += 1
                    if machine_count < required_machines:
                        day_failure_row_count += 1
                        failure_row_count += 1

            planned_capacity_change_qty = (
                0
                if previous_planned_capacity_qty is None
                else round(day_planned_capacity_qty - previous_planned_capacity_qty, 4)
            )
            previous_planned_capacity_qty = day_planned_capacity_qty

            day_failure_rate = (
                round(day_failure_row_count / len(capacity_rows) * 100, 2)
                if len(capacity_rows) > 0
                else 0
            )
            daily_capacity_items.append(
                {
                    "calendar_date": calendar_date,
                    "default_capacity_qty": round(day_default_capacity_qty, 4),
                    "planned_capacity_qty": round(day_planned_capacity_qty, 4),
                    "actual_capacity_qty": round(day_actual_capacity_qty, 4),
                    "planned_capacity_change_qty": planned_capacity_change_qty,
                    "failure_row_count": day_failure_row_count,
                    "row_count": len(capacity_rows),
                    "failure_rate": day_failure_rate,
                }
            )
            for line_code, planned_capacity in day_line_planned_capacity.items():
                line_planned_capacity_by_code[line_code][calendar_date] = planned_capacity
                line_daily_stats_by_key[(calendar_date, line_code)] = {
                    "calendar_date": calendar_date,
                    "line_code": line_code,
                    "line_name": str(
                        (line_meta_by_code.get(line_code) or {}).get("line_name") or line_code
                    ).strip()
                    or line_code,
                    "default_capacity_qty": round(day_line_default_capacity.get(line_code, 0), 4),
                    "planned_capacity_qty": round(planned_capacity, 4),
                    "actual_capacity_qty": round(day_line_actual_capacity.get(line_code, 0), 4),
                }
            for process_code, planned_capacity in day_process_planned_capacity.items():
                process_planned_capacity_by_code[process_code][calendar_date] = planned_capacity
                process_daily_stats_by_key[(calendar_date, process_code)] = {
                    "calendar_date": calendar_date,
                    "process_code": process_code,
                    "process_name_cn": str(
                        (process_meta_by_code.get(process_code) or {}).get("process_name_cn")
                        or process_code
                    ).strip()
                    or process_code,
                    "default_capacity_qty": round(day_process_default_capacity.get(process_code, 0), 4),
                    "planned_capacity_qty": round(planned_capacity, 4),
                    "actual_capacity_qty": round(day_process_actual_capacity.get(process_code, 0), 4),
                }
            day_line_stats = [line_daily_stats_by_key[(calendar_date, code)] for code in day_line_planned_capacity]
            overload_line_count = sum(
                1
                for item in day_line_stats
                if _to_number(item.get("planned_capacity_qty"), 0)
                > _to_number(item.get("default_capacity_qty"), 0) + SCHEDULE_NUMBER_EPSILON
            )
            idle_line_count = sum(
                1
                for item in day_line_stats
                if _to_number(item.get("planned_capacity_qty"), 0)
                + SCHEDULE_NUMBER_EPSILON
                < _to_number(item.get("default_capacity_qty"), 0)
            )
            utilization_rate = (
                round(day_planned_capacity_qty / day_default_capacity_qty * 100, 2)
                if day_default_capacity_qty > SCHEDULE_NUMBER_EPSILON
                else 0
            )
            daily_pressure_items.append(
                {
                    "calendar_date": calendar_date,
                    "default_capacity_qty": round(day_default_capacity_qty, 4),
                    "planned_capacity_qty": round(day_planned_capacity_qty, 4),
                    "actual_capacity_qty": round(day_actual_capacity_qty, 4),
                    "utilization_rate": utilization_rate,
                    "overload_qty": round(max(0.0, day_planned_capacity_qty - day_default_capacity_qty), 4),
                    "idle_qty": round(max(0.0, day_default_capacity_qty - day_planned_capacity_qty), 4),
                    "overload_line_count": overload_line_count,
                    "idle_line_count": idle_line_count,
                }
            )
            total_default_capacity_qty += day_default_capacity_qty
            total_planned_capacity_qty += day_planned_capacity_qty
            total_actual_capacity_qty += day_actual_capacity_qty
            current_date_value = current_date_value + timedelta(days=1)

        if evaluated_row_count <= 0:
            raise server_error(
                code="DASHBOARD_MACHINE_REQUIREMENT_EMPTY",
                message="No machine-driven capacity rows exist in the requested range.",
            )
        if len(line_meta_by_code) == 0:
            raise server_error(
                code="DASHBOARD_LINE_SERIES_EMPTY",
                message="No line capacity series can be built for the requested range.",
            )
        if len(process_meta_by_code) == 0:
            raise server_error(
                code="DASHBOARD_PROCESS_SERIES_EMPTY",
                message="No process capacity series can be built for the requested range.",
            )

        equipment_failure_rate = round(failure_row_count / evaluated_row_count * 100, 2)
        summary_row = (
            order_summary.get("summary")
            if isinstance(order_summary.get("summary"), dict)
            else {}
        )
        order_count = int(_to_number(summary_row.get("order_count"), 0))
        completed_order_count = int(_to_number(summary_row.get("completed_order_count"), 0))
        order_completion_rate = _to_number(summary_row.get("completion_rate"), 0)

        order_no_set: set[str] = set()
        for item in order_items:
            order_no = str(item.get("order_no") or "").strip()
            if not order_no:
                continue
            order_no_set.add(order_no)

        material_rows: list[dict[str, Any]] = []
        if order_no_set:
            placeholders = ",".join("?" for _ in order_no_set)
            material_rows = fetch_all(
                self.connection,
                f"""
                SELECT
                    production_order_no,
                    child_material_code,
                    child_material_name,
                    child_unit,
                    issue_qty
                FROM material_issue_items
                WHERE production_order_no IN ({placeholders})
                """,
                tuple(sorted(order_no_set)),
            )

        if len(material_rows) == 0:
            raise bad_request(
                code="DASHBOARD_MATERIAL_DATA_EMPTY",
                message="No material_issue_items rows exist for the requested orders.",
                details={
                    "order_count": len(order_no_set),
                },
            )

        material_aggregate_map: dict[str, dict[str, Any]] = {}
        for row in material_rows:
            order_no = str(row.get("production_order_no") or "").strip()
            material_code = str(row.get("child_material_code") or "").strip().upper()
            if not order_no or not material_code:
                continue
            if order_no not in order_no_set:
                continue
            issue_qty = _to_number(row.get("issue_qty"), 0)
            estimated_issue_qty = issue_qty
            if estimated_issue_qty <= SCHEDULE_NUMBER_EPSILON:
                continue
            aggregate_row = material_aggregate_map.get(material_code)
            if aggregate_row is None:
                aggregate_row = {
                    "material_code": material_code,
                    "material_name_cn": str(row.get("child_material_name") or material_code).strip()
                    or material_code,
                    "unit": str(row.get("child_unit") or "").strip(),
                    "estimated_issue_qty": 0.0,
                    "order_no_set": set(),
                }
                material_aggregate_map[material_code] = aggregate_row
            aggregate_row["estimated_issue_qty"] = (
                _to_number(aggregate_row.get("estimated_issue_qty"), 0) + estimated_issue_qty
            )
            cast_order_set = aggregate_row.get("order_no_set")
            if isinstance(cast_order_set, set):
                cast_order_set.add(order_no)

        ranked_material_items = sorted(
            material_aggregate_map.values(),
            key=lambda item: (
                -_to_number(item.get("estimated_issue_qty"), 0),
                str(item.get("material_code") or ""),
            ),
        )
        if len(ranked_material_items) == 0:
            raise bad_request(
                code="DASHBOARD_MATERIAL_CONSUMPTION_EMPTY",
                message="Material consumption ranking is empty in the requested range.",
            )
        material_ranking = ranked_material_items[:normalized_top_n]
        material_consumption_items = [
            {
                "rank": index + 1,
                "material_code": str(item.get("material_code") or ""),
                "material_name_cn": str(item.get("material_name_cn") or ""),
                "unit": str(item.get("unit") or ""),
                "estimated_issue_qty": round(_to_number(item.get("estimated_issue_qty"), 0), 4),
                "order_count": len(item.get("order_no_set") or set()),
            }
            for index, item in enumerate(material_ranking)
        ]

        line_overload_items = []
        for item in line_daily_stats_by_key.values():
            default_capacity_qty = _to_number(item.get("default_capacity_qty"), 0)
            planned_capacity_qty = _to_number(item.get("planned_capacity_qty"), 0)
            actual_capacity_qty = _to_number(item.get("actual_capacity_qty"), 0)
            utilization_rate = (
                round(planned_capacity_qty / default_capacity_qty * 100, 2)
                if default_capacity_qty > SCHEDULE_NUMBER_EPSILON
                else 0
            )
            overload_qty = max(0.0, planned_capacity_qty - default_capacity_qty)
            idle_qty = max(0.0, default_capacity_qty - planned_capacity_qty)
            line_overload_items.append(
                {
                    **item,
                    "utilization_rate": utilization_rate,
                    "overload_qty": round(overload_qty, 4),
                    "idle_qty": round(idle_qty, 4),
                    "actual_capacity_qty": round(actual_capacity_qty, 4),
                }
            )
        line_overload_items.sort(
            key=lambda item: (
                -_to_number(item.get("overload_qty"), 0),
                -_to_number(item.get("utilization_rate"), 0),
                str(item.get("calendar_date") or ""),
                str(item.get("line_code") or ""),
            )
        )

        process_bottleneck_items = []
        process_grouped_daily: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for item in process_daily_stats_by_key.values():
            process_grouped_daily[str(item.get("process_code") or "")].append(item)
        for process_code, grouped_items in process_grouped_daily.items():
            peak_item = None
            utilization_rates: list[float] = []
            peak_utilization_rate = 0.0
            for item in grouped_items:
                default_capacity_qty = _to_number(item.get("default_capacity_qty"), 0)
                planned_capacity_qty = _to_number(item.get("planned_capacity_qty"), 0)
                utilization_rate = (
                    round(planned_capacity_qty / default_capacity_qty * 100, 2)
                    if default_capacity_qty > SCHEDULE_NUMBER_EPSILON
                    else 0
                )
                utilization_rates.append(utilization_rate)
                if peak_item is None or utilization_rate > peak_utilization_rate:
                    peak_item = item
                    peak_utilization_rate = utilization_rate
            average_utilization_rate = (
                round(sum(utilization_rates) / len(utilization_rates), 2)
                if utilization_rates
                else 0
            )
            process_bottleneck_items.append(
                {
                    "process_code": process_code,
                    "process_name_cn": str(
                        (peak_item or {}).get("process_name_cn") or process_code
                    ).strip()
                    or process_code,
                    "peak_utilization_rate": peak_utilization_rate,
                    "average_utilization_rate": average_utilization_rate,
                    "peak_date": (peak_item or {}).get("calendar_date"),
                    "peak_planned_capacity_qty": round(
                        _to_number((peak_item or {}).get("planned_capacity_qty"), 0),
                        4,
                    ),
                    "peak_default_capacity_qty": round(
                        _to_number((peak_item or {}).get("default_capacity_qty"), 0),
                        4,
                    ),
                }
            )
        process_bottleneck_items.sort(
            key=lambda item: (
                -_to_number(item.get("peak_utilization_rate"), 0),
                -_to_number(item.get("average_utilization_rate"), 0),
                str(item.get("process_code") or ""),
            )
        )

        peak_overload_day = max(
            daily_pressure_items,
            key=lambda item: (
                _to_number(item.get("overload_qty"), 0),
                _to_number(item.get("utilization_rate"), 0),
            ),
            default={},
        )
        peak_idle_day = max(
            daily_pressure_items,
            key=lambda item: (
                _to_number(item.get("idle_qty"), 0),
                -_to_number(item.get("utilization_rate"), 0),
            ),
            default={},
        )
        top_overload_line = line_overload_items[0] if line_overload_items else {}
        top_bottleneck_process = process_bottleneck_items[0] if process_bottleneck_items else {}

        total_days = (end_date_value - start_date_value).days + 1
        return {
            "range": {
                "start_date": normalized_start_date,
                "end_date": normalized_end_date,
                "total_days": total_days,
            },
            "summary": {
                "order_count": order_count,
                "completed_order_count": completed_order_count,
                "order_completion_rate": round(order_completion_rate, 2),
                "evaluated_machine_row_count": evaluated_row_count,
                "failure_row_count": failure_row_count,
                "equipment_failure_rate": equipment_failure_rate,
                "total_capacity_qty": round(total_planned_capacity_qty, 4),
                "total_default_capacity_qty": round(total_default_capacity_qty, 4),
                "total_planned_capacity_qty": round(total_planned_capacity_qty, 4),
                "total_actual_capacity_qty": round(total_actual_capacity_qty, 4),
            },
            "daily_capacity": {
                "items": daily_capacity_items,
            },
            "capacity_change_by_line": {
                "options": _build_dashboard_options(
                    line_meta_by_code,
                    code_key="line_code",
                    name_key="line_name",
                ),
                "items": _build_dashboard_change_items(
                    calendar_dates,
                    line_planned_capacity_by_code,
                    line_meta_by_code,
                    code_key="line_code",
                    name_key="line_name",
                ),
            },
            "capacity_change_by_process": {
                "options": _build_dashboard_options(
                    process_meta_by_code,
                    code_key="process_code",
                    name_key="process_name_cn",
                ),
                "items": _build_dashboard_change_items(
                    calendar_dates,
                    process_planned_capacity_by_code,
                    process_meta_by_code,
                    code_key="process_code",
                    name_key="process_name_cn",
                ),
            },
            "material_consumption": {
                "top_n": normalized_top_n,
                "items": material_consumption_items,
            },
            "cockpit": {
                "summary": {
                    "peak_overload_date": peak_overload_day.get("calendar_date"),
                    "peak_overload_qty": round(
                        _to_number(peak_overload_day.get("overload_qty"), 0),
                        4,
                    ),
                    "peak_idle_date": peak_idle_day.get("calendar_date"),
                    "peak_idle_qty": round(_to_number(peak_idle_day.get("idle_qty"), 0), 4),
                    "top_overload_line_code": top_overload_line.get("line_code"),
                    "top_overload_line_name": top_overload_line.get("line_name"),
                    "top_overload_line_date": top_overload_line.get("calendar_date"),
                    "top_overload_line_utilization_rate": _to_number(
                        top_overload_line.get("utilization_rate"),
                        0,
                    ),
                    "top_bottleneck_process_code": top_bottleneck_process.get("process_code"),
                    "top_bottleneck_process_name_cn": top_bottleneck_process.get(
                        "process_name_cn"
                    ),
                    "top_bottleneck_process_peak_date": top_bottleneck_process.get("peak_date"),
                    "top_bottleneck_process_peak_utilization_rate": _to_number(
                        top_bottleneck_process.get("peak_utilization_rate"),
                        0,
                    ),
                },
                "line_overload_items": line_overload_items[:8],
                "process_bottleneck_items": process_bottleneck_items[:8],
                "daily_pressure_items": daily_pressure_items,
            },
        }

    def list_order_summary_workshop_managers(
        self,
        *,
        current_user: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        if self._current_user_role_code(current_user) != ROLE_SCHEDULER:
            raise forbidden(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_FILTER_FORBIDDEN",
                message="Only scheduler can list order summary workshop manager filters.",
            )
        rows = fetch_all(
            self.connection,
            """
            SELECT
                users.user_id,
                users.username,
                users.display_name,
                COUNT(scope.line_code) AS line_scope_count
            FROM app_users users
            LEFT JOIN masterdata_workshop_manager_visibility visibility
              ON visibility.user_id = users.user_id
            JOIN app_user_line_scopes scope
              ON scope.user_id = users.user_id
            WHERE users.role_code = ?
              AND users.enabled_flag = 1
              AND COALESCE(visibility.visible_flag, 1) = 1
            GROUP BY users.user_id, users.username, users.display_name
            HAVING COUNT(scope.line_code) > 0
            ORDER BY
                LOWER(COALESCE(NULLIF(TRIM(users.display_name), ''), NULLIF(TRIM(users.username), ''), users.user_id)) ASC,
                LOWER(COALESCE(users.username, '')) ASC,
                users.user_id ASC
            """,
            (ROLE_WORKSHOP_MANAGER,),
        )
        return {
            "items": [
                {
                    "user_id": str(row.get("user_id") or "").strip(),
                    "username": str(row.get("username") or "").strip(),
                    "display_name": str(row.get("display_name") or "").strip(),
                    "line_scope_count": int(_to_number(row.get("line_scope_count"), 0)),
                }
                for row in rows
            ]
        }

    def list_line_daily_capacity(
        self,
        calendar_date: str,
        *,
        workshop_code: str | None = None,
        line_code: str | None = None,
        process_code: str | None = None,
        current_user: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        return self.line_daily_capacity_service.list_line_daily_capacity(
            calendar_date,
            workshop_code=workshop_code,
            line_code=line_code,
            process_code=process_code,
            current_user=current_user,
        )

    def save_line_daily_capacity(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self.line_daily_capacity_service.save_line_daily_capacity(payload)

    def list_line_daily_capacity_audits(
        self,
        calendar_date: str,
        *,
        workshop_code: str | None = None,
        line_code: str | None = None,
        process_code: str | None = None,
        operator_keyword: str | None = None,
        changed_only: bool = False,
        current_user: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        return self.line_daily_capacity_service.list_line_daily_capacity_audits(
            calendar_date,
            workshop_code=workshop_code,
            line_code=line_code,
            process_code=process_code,
            operator_keyword=operator_keyword,
            changed_only=changed_only,
            current_user=current_user,
        )

    def rebuild_line_daily_actual_capacity(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self.line_daily_capacity_service.rebuild_line_daily_actual_capacity(payload)

    def _rebuild_line_daily_actual_capacity_rows(
        self,
        normalized_date: str,
    ) -> tuple[int, int]:
        return self.line_daily_capacity_service._rebuild_line_daily_actual_capacity_rows(
            normalized_date
        )

    def create_reporting(self, payload: dict[str, Any]) -> dict[str, Any]:
        order_no = str(payload.get("order_no") or "").strip() or None
        process_code = str(payload.get("process_code") or "").strip().upper()
        report_qty = _to_number(payload.get("report_qty"), -1)
        report_scope = str(payload.get("report_scope") or "").strip().upper()
        if not process_code:
            raise bad_request(
                code="PROCESS_CODE_REQUIRED",
                message="process_code is required.",
            )
        if report_qty <= 0:
            raise bad_request(
                code="REPORT_QTY_INVALID",
                message="report_qty must be greater than 0.",
            )
        if report_scope and report_scope not in {"ORDER", "LINE_OUTPUT"}:
            raise bad_request(
                code="REPORT_SCOPE_INVALID",
                message="report_scope must be ORDER or LINE_OUTPUT.",
            )
        if order_no:
            self._require_order(order_no)
        elif report_scope == "ORDER":
            raise bad_request(
                code="REPORT_ORDER_REQUIRED",
                message="订单报工必须指定 order_no。",
            )
        normalized_report_scope = report_scope or ("ORDER" if order_no else "LINE_OUTPUT")
        process_name_by_code = self._process_name_by_code()
        report_id = f"RPT-{uuid4().hex[:10].upper()}"
        report_time_text = str(payload.get("report_time") or "").strip()
        calendar_date_input = payload.get("calendar_date")
        calendar_date_text = str(calendar_date_input or "").strip()
        normalized_calendar_date = _normalize_date_text(calendar_date_input)
        if calendar_date_text and normalized_calendar_date is None:
            raise bad_request(
                code="CALENDAR_DATE_INVALID",
                message="calendar_date must be a valid YYYY-MM-DD date.",
            )
        local_timezone = LOCAL_TIMEZONE
        if report_time_text:
            try:
                report_time_dt = datetime.fromisoformat(report_time_text)
            except ValueError as exc:
                raise bad_request(
                    code="REPORT_TIME_INVALID",
                    message="report_time must be a valid ISO datetime.",
                ) from exc
            if report_time_dt.tzinfo is None or report_time_dt.utcoffset() is None:
                raise bad_request(
                    code="REPORT_TIME_TIMEZONE_REQUIRED",
                    message="report_time must include timezone offset.",
                )
            report_time = report_time_dt.astimezone(timezone.utc).replace(microsecond=0).isoformat()
            reporting_local_date = report_time_dt.astimezone(local_timezone).date().isoformat()
        elif normalized_calendar_date:
            reporting_local_date = normalized_calendar_date
            report_time = (
                datetime.fromisoformat(f"{normalized_calendar_date}T10:00:00+08:00")
                .astimezone(timezone.utc)
                .replace(microsecond=0)
                .isoformat()
            )
        else:
            simulation_state = self._get_simulation_state()
            reporting_local_date = _normalize_date_text(simulation_state.get("current_date"))
            if reporting_local_date is None:
                raise server_error(
                    code="SIMULATION_CURRENT_DATE_INVALID",
                    message="Current simulation date is invalid.",
                    details={"current_date": simulation_state.get("current_date")},
                )
            report_time = (
                datetime.fromisoformat(f"{reporting_local_date}T10:00:00+08:00")
                .astimezone(timezone.utc)
                .replace(microsecond=0)
                .isoformat()
            )
        operator_name = (
            str(
                payload.get("operator_name_cn")
                or payload.get("operator_name")
                or "系统填报"
            ).strip()
            or "系统填报"
        )
        workshop_code = str(payload.get("workshop_code") or "").strip().upper() or None
        workshop_name = str(payload.get("workshop_name") or workshop_code or "").strip() or workshop_code
        line_code = str(payload.get("line_code") or "").strip().upper() or None
        line_name = str(payload.get("line_name") or line_code or "").strip() or line_code
        company_code = str(payload.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() or DEFAULT_COMPANY_CODE
        actor = payload.get("actor") if isinstance(payload.get("actor"), dict) else {}
        if self._is_workshop_manager(actor):
            if not workshop_code or not line_code:
                raise bad_request(
                    code="REPORT_LINE_REQUIRED",
                    message="workshop_code and line_code are required for workshop manager reporting.",
                )
            self._assert_actor_can_access_line(
                actor,
                company_code=company_code,
                workshop_code=workshop_code,
                line_code=line_code,
                missing_user_error_code="REPORT_ACTOR_USER_REQUIRED",
                forbidden_error_code="REPORT_LINE_SCOPE_FORBIDDEN",
                forbidden_message="Current workshop manager is not allowed to report on this line.",
            )
        with transaction(self.connection):
            self.connection.execute(
                """
                INSERT INTO work_reports (
                    report_id,
                    production_order_no,
                    process_code,
                    process_name,
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    report_id,
                    order_no,
                    normalized_report_scope,
                    process_code,
                    process_name_by_code.get(process_code, process_code),
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    report_time,
                ),
            )
            if order_no:
                self._sync_order_state_from_reporting(
                    order_no=order_no,
                    process_code=process_code,
                )
        if workshop_code and line_code:
            self.rebuild_line_daily_actual_capacity({"calendar_date": reporting_local_date})
        return {
            "report_id": report_id,
            "order_no": order_no,
            "report_scope": normalized_report_scope,
            "process_code": process_code,
            "process_name_cn": process_name_by_code.get(process_code, process_code),
            "report_qty": report_qty,
            "report_time": report_time,
            "operator_name_cn": operator_name,
        }

    def _sync_order_state_from_reporting(self, *, order_no: str, process_code: str) -> None:
        base_row = self._require_order(order_no)
        state_row = self._get_order_state(order_no) or {}
        product_code = str(base_row.get("material_code") or "").strip().upper()
        final_meta = self._resolve_final_process_meta_by_product({product_code}).get(product_code) or {}
        final_process_code = str(final_meta.get("process_code") or "").strip().upper()
        if not final_process_code or final_process_code != process_code:
            return

        final_report_row = fetch_one(
            self.connection,
            """
            SELECT COALESCE(SUM(report_qty), 0) AS total_qty
            FROM work_reports
            WHERE production_order_no = ?
              AND UPPER(TRIM(COALESCE(process_code, ''))) = ?
            """,
            (order_no, final_process_code),
        )
        completed_qty = _to_number((final_report_row or {}).get("total_qty"), 0)
        production_qty = _to_number(base_row.get("production_qty"), 0)
        remaining_qty = max(0.0, production_qty - completed_qty)
        progress_rate = (completed_qty / production_qty * 100) if production_qty > 0 else 0
        normalized_status = "DONE" if production_qty > 0 and completed_qty + SCHEDULE_NUMBER_EPSILON >= production_qty else "IN_PROGRESS" if completed_qty > 0 else "OPEN"
        next_row = {
            "production_order_no": order_no,
            "promised_due_date": state_row.get("promised_due_date") or base_row.get("planned_end_date"),
            "expected_start_date": state_row.get("expected_start_date") or base_row.get("planned_start_date"),
            "expected_start_time": state_row.get("expected_start_time")
            or _iso_at(state_row.get("expected_start_date") or base_row.get("planned_start_date"), "08:00:00"),
            "expected_finish_time": state_row.get("expected_finish_time")
            or _iso_at(state_row.get("promised_due_date") or base_row.get("planned_end_date"), "18:00:00"),
            "priority_level": _normalize_priority_level(state_row.get("priority_level"), PRIORITY_LEVEL_MAX),
            "urgent_flag": _urgent_flag_from_priority_level(state_row.get("priority_level")),
            "lock_flag": int(_to_number(state_row.get("lock_flag"), 0)),
            "frozen_flag": int(_to_number(state_row.get("frozen_flag"), 0)),
            "status": normalized_status,
            "order_status": normalized_status,
            "completed_qty": completed_qty,
            "remaining_qty": remaining_qty,
            "progress_rate": min(100.0, round(progress_rate, 4)),
            "production_batch_no": state_row.get("production_batch_no"),
        }
        self._upsert_order_state(next_row)

    def import_mes_reportings_from_xlsx(self, payload: dict[str, Any]) -> dict[str, Any]:
        file_path = str(payload.get("file_path") or "").strip()
        if not file_path:
            raise bad_request(
                code="REPORTING_IMPORT_FILE_PATH_REQUIRED",
                message="file_path is required.",
            )
        if not file_path.lower().endswith(".xlsx"):
            raise bad_request(
                code="REPORTING_IMPORT_FILE_TYPE_INVALID",
                message="Only .xlsx files are supported.",
                details={"file_path": file_path},
            )
        workbook_path = Path(file_path)
        if not workbook_path.exists():
            raise bad_request(
                code="REPORTING_IMPORT_FILE_NOT_FOUND",
                message="xlsx file does not exist.",
                details={"file_path": file_path},
            )

        def _normalize_sha256(value: object) -> str | None:
            raw = str(value or "").strip().lower()
            if not raw:
                return None
            text = raw[7:] if raw.startswith("sha256:") else raw
            if len(text) != 64:
                return None
            try:
                binascii.unhexlify(text)
            except binascii.Error:
                return None
            return text

        file_sha256 = _normalize_sha256(payload.get("file_sha256"))
        if file_sha256 is None:
            file_sha256 = _normalize_sha256(payload.get("source_file_name"))
        if file_sha256 is None:
            hasher = hashlib.sha256()
            with workbook_path.open("rb") as reader:
                for chunk in iter(lambda: reader.read(1024 * 1024), b""):
                    hasher.update(chunk)
            file_sha256 = hasher.hexdigest()

        sha_source_file_name = f"sha256:{file_sha256}"
        explicit_source_file_name = str(payload.get("source_file_name") or "").strip()
        source_file_names: list[str] = []
        for candidate in (sha_source_file_name, explicit_source_file_name, file_path):
            if candidate and candidate not in source_file_names:
                source_file_names.append(candidate)

        original_file_name = str(payload.get("original_file_name") or "").strip() or workbook_path.name
        file_size_bytes = int(workbook_path.stat().st_size or 0)
        size_input = str(payload.get("file_size_bytes") or "").strip()
        if size_input:
            try:
                file_size_bytes = max(0, int(size_input))
            except ValueError:
                file_size_bytes = int(workbook_path.stat().st_size or 0)

        company_code = (
            str(payload.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper()
            or DEFAULT_COMPANY_CODE
        )
        create_missing_orders = False
        if payload.get("create_missing_orders") is not None:
            try:
                create_missing_orders = _parse_enabled_flag(
                    payload.get("create_missing_orders")
                ) == 1
            except ValueError as exc:
                raise bad_request(
                    code="REPORTING_IMPORT_CREATE_MISSING_ORDERS_INVALID",
                    message=str(exc),
                ) from exc
        sheet_names_payload = payload.get("sheet_names")
        sheet_names: list[str] | None = None
        if sheet_names_payload is not None:
            if not isinstance(sheet_names_payload, list):
                raise bad_request(
                    code="REPORTING_IMPORT_SHEET_NAMES_INVALID",
                    message="sheet_names must be an array of strings.",
                )
            normalized_sheet_names: list[str] = []
            for item in sheet_names_payload:
                name = str(item or "").strip()
                if not name:
                    raise bad_request(
                        code="REPORTING_IMPORT_SHEET_NAMES_INVALID",
                        message="sheet_names must not contain empty names.",
                    )
                if name not in normalized_sheet_names:
                    normalized_sheet_names.append(name)
            sheet_names = normalized_sheet_names

        def parse_optional_text(value: object) -> str | None:
            text = str(value or "").strip()
            return text or None

        def parse_optional_number(value: object) -> float | None:
            if value is None:
                return None
            if isinstance(value, (int, float)):
                number = float(value)
                return number if number == number else None
            text = str(value or "").strip()
            if not text:
                return None
            try:
                number = float(text)
            except (TypeError, ValueError):
                return None
            return number if number == number else None

        def parse_required_positive_number(value: object) -> float | None:
            number = parse_optional_number(value)
            if number is None or number <= 0:
                return None
            return number

        def parse_report_datetime(value: object) -> datetime | None:
            if isinstance(value, datetime):
                return value
            if isinstance(value, date):
                return datetime.combine(value, datetime.min.time())
            text = str(value or "").strip()
            if not text:
                return None
            try:
                return datetime.fromisoformat(text)
            except ValueError:
                return None

        def to_utc_iso(value: datetime) -> str:
            parsed = value
            if parsed.tzinfo is None or parsed.utcoffset() is None:
                parsed = parsed.replace(tzinfo=LOCAL_TIMEZONE)
            return (
                parsed.astimezone(timezone.utc).replace(microsecond=0).isoformat()
            )

        required_headers = ("报工日期", "生产订单号", "工序编码", "报工数量")
        header_to_field: dict[str, str] = {
            **REPORTING_IMPORT_REQUIRED_HEADER_MAP,
            **REPORTING_IMPORT_OPTIONAL_HEADER_MAP,
            "工序编码": "process_code",
            "工序名称": "process_name",
        }

        try:
            workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        except Exception as exc:
            raise bad_request(
                code="REPORTING_IMPORT_WORKBOOK_LOAD_FAILED",
                message="Failed to read xlsx workbook.",
                details={"file_path": file_path, "error": str(exc)},
            ) from exc

        try:
            available_sheet_names = list(workbook.sheetnames)
            selected_sheet_names = sheet_names or available_sheet_names
            if not selected_sheet_names:
                raise bad_request(
                    code="REPORTING_IMPORT_SHEETS_EMPTY",
                    message="xlsx workbook has no sheets.",
                    details={"file_path": file_path},
                )
            missing_sheets = [
                name
                for name in selected_sheet_names
                if name not in available_sheet_names
            ]
            if missing_sheets:
                raise bad_request(
                    code="REPORTING_IMPORT_SHEET_NOT_FOUND",
                    message="Some sheets do not exist in workbook.",
                    details={
                        "file_path": file_path,
                        "missing_sheet_names": missing_sheets,
                        "available_sheet_names": available_sheet_names,
                    },
                )

            imported_count = 0
            skipped_existing_count = 0
            failed_count = 0
            total_row_count = 0
            failures: list[dict[str, Any]] = []
            order_exists_cache: dict[str, bool] = {}
            created_missing_order_count = 0
            created_missing_order_nos: list[str] = []

            for sheet_name in selected_sheet_names:
                sheet = workbook[sheet_name]
                header_cells = [
                    str(sheet.cell(1, c).value or "").strip()
                    for c in range(1, sheet.max_column + 1)
                ]
                header_index: dict[str, int] = {}
                for col_index, header in enumerate(header_cells, start=1):
                    if header and header not in header_index:
                        header_index[header] = col_index

                missing_headers = [
                    header for header in required_headers if header not in header_index
                ]
                if missing_headers:
                    raise bad_request(
                        code="REPORTING_IMPORT_REQUIRED_HEADERS_MISSING",
                        message=f"Missing required header(s) in sheet '{sheet_name}'.",
                        details={
                            "sheet_name": sheet_name,
                            "missing_headers": missing_headers,
                        },
                    )

                sheet_columns: list[tuple[str, int]] = []
                for header, field_name in header_to_field.items():
                    col = header_index.get(header)
                    if col is None:
                        continue
                    sheet_columns.append((field_name, col))

                for row_no in range(2, sheet.max_row + 1):
                    raw: dict[str, Any] = {
                        field_name: sheet.cell(row_no, col).value
                        for field_name, col in sheet_columns
                    }
                    report_datetime_raw = raw.get("report_datetime")
                    order_no_raw = raw.get("production_order_no")
                    process_code_raw = raw.get("process_code")
                    report_qty_raw = raw.get("report_qty")
                    if (
                        report_datetime_raw is None
                        and not parse_optional_text(order_no_raw)
                        and not parse_optional_text(process_code_raw)
                        and parse_optional_text(report_qty_raw) is None
                    ):
                        continue

                    total_row_count += 1
                    source_placeholder_sql = ",".join("?" for _ in source_file_names)
                    existing = fetch_one(
                        self.connection,
                        f"""
                        SELECT report_id
                        FROM work_reports
                        WHERE source_file_name IN ({source_placeholder_sql})
                          AND source_sheet_name = ?
                          AND source_row_no = ?
                        LIMIT 1
                        """,
                        tuple(source_file_names + [sheet_name, row_no]),
                    )
                    if existing is not None:
                        skipped_existing_count += 1
                        continue

                    report_dt = parse_report_datetime(report_datetime_raw)
                    if report_dt is None:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "error": "报工日期无法解析。",
                            }
                        )
                        continue
                    report_time = to_utc_iso(report_dt)

                    order_no_text = parse_optional_text(order_no_raw)
                    if not order_no_text:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "error": "生产订单号为空。",
                            }
                        )
                        continue
                    base_order_no = str(order_no_text.split("-")[0]).strip()
                    if not base_order_no:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "order_no": order_no_text,
                                "error": "生产订单号格式无效。",
                            }
                        )
                        continue

                    process_code_text = parse_optional_text(process_code_raw)
                    if not process_code_text:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "order_no": order_no_text,
                                "error": "工序编码为空。",
                            }
                        )
                        continue
                    process_code = process_code_text.upper()

                    report_qty = parse_required_positive_number(report_qty_raw)
                    if report_qty is None:
                        failed_count += 1
                        failures.append(
                            {
                                "sheet_name": sheet_name,
                                "row_no": row_no,
                                "order_no": order_no_text,
                                "process_code": process_code,
                                "error": "报工数量必须大于 0。",
                            }
                        )
                        continue

                    if base_order_no not in order_exists_cache:
                        order_exists_cache[base_order_no] = (
                            fetch_one(
                                self.connection,
                                """
                                SELECT production_order_no
                                FROM production_orders
                                WHERE production_order_no = ?
                                LIMIT 1
                                """,
                                (base_order_no,),
                            )
                            is not None
                        )

                    if not order_exists_cache[base_order_no]:
                        if not create_missing_orders:
                            failed_count += 1
                            failures.append(
                                {
                                    "sheet_name": sheet_name,
                                    "row_no": row_no,
                                    "order_no": order_no_text,
                                    "base_order_no": base_order_no,
                                    "process_code": process_code,
                                    "error": (
                                        "生产订单不存在，请先同步/导入生产订单后再导入报工。"
                                    ),
                                }
                            )
                            continue
                        material_code = parse_optional_text(raw.get("product_code")) or "UNKNOWN"
                        material_name = (
                            parse_optional_text(raw.get("product_name"))
                            or material_code
                            or base_order_no
                        )
                        material_specification = parse_optional_text(
                            raw.get("product_specification")
                        )
                        with transaction(self.connection):
                            self.connection.execute(
                                """
                                INSERT INTO production_orders (
                                    production_order_no,
                                    material_code,
                                    material_name,
                                    material_specification,
                                    production_qty,
                                    status,
                                    planned_start_date,
                                    planned_end_date,
                                    source_bill_no,
                                    material_list_no,
                                    updated_at
                                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                ON CONFLICT(production_order_no) DO NOTHING
                                """,
                                (
                                    base_order_no,
                                    material_code,
                                    material_name,
                                    material_specification,
                                    0.0,
                                    "OPEN",
                                    None,
                                    None,
                                    None,
                                    None,
                                    report_time,
                                ),
                            )
                        order_exists_cache[base_order_no] = True
                        created_missing_order_count += 1
                        if len(created_missing_order_nos) < 50:
                            created_missing_order_nos.append(base_order_no)

                    process_name = parse_optional_text(raw.get("process_name"))
                    report_id = f"RPT-{uuid4().hex[:10].upper()}"
                    row_payload: dict[str, Any] = {
                        "report_id": report_id,
                        "production_order_no": base_order_no,
                        "process_code": process_code,
                        "process_name": process_name,
                        "company_code": company_code,
                        "workshop_code": None,
                        "workshop_name": None,
                        "line_code": None,
                        "line_name": None,
                        "report_qty": float(report_qty),
                        "report_time": report_time,
                        "operator_code": parse_optional_text(raw.get("operator_code")),
                        "operator_name": parse_optional_text(raw.get("operator_name")),
                        "section_leader_name": parse_optional_text(
                            raw.get("section_leader_name")
                        ),
                        "dispatch_no": parse_optional_text(raw.get("dispatch_no")),
                        "product_code": parse_optional_text(raw.get("product_code")),
                        "product_name": parse_optional_text(raw.get("product_name")),
                        "product_specification": parse_optional_text(
                            raw.get("product_specification")
                        ),
                        "resource_group_name": parse_optional_text(
                            raw.get("resource_group_name")
                        ),
                        "resource_name": parse_optional_text(raw.get("resource_name")),
                        "department_name": parse_optional_text(
                            raw.get("department_name")
                        ),
                        "source_process_code": process_code,
                        "source_process_name": process_name,
                        "mold_code": parse_optional_text(raw.get("mold_code")),
                        "support_count": parse_optional_number(raw.get("support_count")),
                        "weight_kg": parse_optional_number(raw.get("weight_kg")),
                        "cavity_count": parse_optional_number(raw.get("cavity_count")),
                        "total_cycle_time": parse_optional_number(
                            raw.get("total_cycle_time")
                        ),
                        "production_quota": parse_optional_number(
                            raw.get("production_quota")
                        ),
                        "work_duration": parse_optional_number(raw.get("work_duration")),
                        "clamp_or_assembly_weight": parse_optional_number(
                            raw.get("clamp_or_assembly_weight")
                        ),
                        "unit_weight": parse_optional_number(raw.get("unit_weight")),
                        "source_sheet_name": sheet_name,
                        "source_row_no": row_no,
                        "source_file_name": sha_source_file_name,
                        "updated_at": report_time,
                    }
                    columns_sql = ", ".join(WORK_REPORT_COLUMNS)
                    placeholders = ", ".join("?" for _ in WORK_REPORT_COLUMNS)
                    values = tuple(row_payload.get(column) for column in WORK_REPORT_COLUMNS)
                    with transaction(self.connection):
                        self.connection.execute(
                            f"INSERT INTO work_reports ({columns_sql}) VALUES ({placeholders})",
                            values,
                        )
                    imported_count += 1

            actor = payload.get("actor") if isinstance(payload.get("actor"), dict) else {}
            imported_by_user_id = str(actor.get("user_id") or "").strip() or None
            imported_by_username = str(actor.get("username") or "").strip() or None
            imported_by_display_name = (
                str(actor.get("display_name") or "").strip() or None
            )
            imported_at = utc_now()
            previous_imported_at: str | None = None
            existing_file_record = fetch_one(
                self.connection,
                """
                SELECT last_imported_at
                FROM reporting_import_files
                WHERE file_sha256 = ?
                """,
                (file_sha256,),
            )
            if existing_file_record is not None:
                previous_imported_at = str(
                    existing_file_record.get("last_imported_at") or ""
                ).strip() or None

            with transaction(self.connection):
                self.connection.execute(
                    """
                    INSERT INTO reporting_import_files (
                        file_sha256,
                        source_file_name,
                        file_path,
                        original_file_name,
                        file_size_bytes,
                        sheet_names_json,
                        imported_by_user_id,
                        imported_by_username,
                        imported_by_display_name,
                        total_row_count,
                        imported_count,
                        skipped_existing_count,
                        failed_count,
                        created_missing_order_count,
                        created_at,
                        last_imported_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(file_sha256) DO UPDATE SET
                        source_file_name = excluded.source_file_name,
                        file_path = excluded.file_path,
                        original_file_name = excluded.original_file_name,
                        file_size_bytes = excluded.file_size_bytes,
                        sheet_names_json = excluded.sheet_names_json,
                        imported_by_user_id = excluded.imported_by_user_id,
                        imported_by_username = excluded.imported_by_username,
                        imported_by_display_name = excluded.imported_by_display_name,
                        total_row_count = excluded.total_row_count,
                        imported_count = excluded.imported_count,
                        skipped_existing_count = excluded.skipped_existing_count,
                        failed_count = excluded.failed_count,
                        created_missing_order_count = excluded.created_missing_order_count,
                        last_imported_at = excluded.last_imported_at
                    """,
                    (
                        file_sha256,
                        sha_source_file_name,
                        file_path,
                        original_file_name,
                        int(file_size_bytes),
                        dumps(selected_sheet_names),
                        imported_by_user_id,
                        imported_by_username,
                        imported_by_display_name,
                        int(total_row_count),
                        int(imported_count),
                        int(skipped_existing_count),
                        int(failed_count),
                        int(created_missing_order_count),
                        imported_at,
                        imported_at,
                    ),
                )

            return {
                "file_path": file_path,
                "source_file_name": sha_source_file_name,
                "file_sha256": file_sha256,
                "original_file_name": original_file_name,
                "file_size_bytes": int(file_size_bytes),
                "previous_imported_at": previous_imported_at,
                "sheet_names": selected_sheet_names,
                "total_row_count": total_row_count,
                "imported_count": imported_count,
                "skipped_existing_count": skipped_existing_count,
                "failed_count": failed_count,
                "created_missing_order_count": created_missing_order_count,
                "created_missing_order_nos": created_missing_order_nos,
                "failures": failures,
            }
        finally:
            workbook.close()

    def select_reporting_capacity_compare(self, payload: dict[str, Any]) -> dict[str, Any]:
        report_id = str(payload.get("report_id") or "").strip()
        if not report_id:
            raise bad_request(
                code="REPORT_ID_REQUIRED",
                message="report_id is required.",
            )
        audit_id = str(payload.get("audit_id") or "").strip()
        if not audit_id:
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_AUDIT_ID_REQUIRED",
                message="audit_id is required.",
            )
        actor = payload.get("actor") if isinstance(payload.get("actor"), dict) else {}
        existing = fetch_one(
            self.connection,
            """
            SELECT
                report_id,
                process_code,
                workshop_code,
                line_code,
                report_time,
                daily_capacity_compare_audit_id,
                daily_capacity_compare_qty
            FROM work_reports
            WHERE report_id = ?
            """,
            (report_id,),
        )
        if existing is None:
            raise not_found(
                code="REPORT_NOT_FOUND",
                message="Report does not exist.",
                details={"report_id": report_id},
            )

        workshop_code = str(existing.get("workshop_code") or "").strip().upper()
        line_code = str(existing.get("line_code") or "").strip().upper()
        process_code = str(existing.get("process_code") or "").strip().upper()
        report_local_date = _local_date_from_iso_datetime(existing.get("report_time"))
        if not report_local_date or not workshop_code or not line_code or not process_code:
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_REPORT_KEY_INVALID",
                message=(
                    "The report record is missing report_time, workshop_code, line_code, "
                    "or process_code required for capacity comparison."
                ),
                details={"report_id": report_id},
            )

        if self._is_workshop_manager(actor):
            self._assert_actor_can_access_line(
                actor,
                company_code=DEFAULT_COMPANY_CODE,
                workshop_code=workshop_code,
                line_code=line_code,
                missing_user_error_code="REPORT_ACTOR_USER_REQUIRED",
                forbidden_error_code="REPORT_LINE_SCOPE_FORBIDDEN",
                forbidden_message=(
                    "Current workshop manager is not allowed to select a capacity comparison for this reporting record."
                ),
            )

        candidate_rows = fetch_all(
            self.connection,
            """
            SELECT
                audit_id,
                new_planned_capacity_qty
            FROM daily_line_capacity_plan_audit
            WHERE calendar_date = ?
              AND company_code = ?
              AND workshop_code = ?
              AND line_code = ?
              AND process_code = ?
            ORDER BY changed_at DESC, audit_id DESC
            """,
            (
                report_local_date,
                DEFAULT_COMPANY_CODE,
                workshop_code,
                line_code,
                process_code,
            ),
        )
        if len(candidate_rows) == 0:
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_AUDITS_EMPTY",
                message="No same-day daily capacity audit records exist for this report.",
                details={
                    "report_id": report_id,
                    "report_local_date": report_local_date,
                    "workshop_code": workshop_code,
                    "line_code": line_code,
                    "process_code": process_code,
                },
            )

        selected_audit = next(
            (
                row
                for row in candidate_rows
                if str(row.get("audit_id") or "").strip() == audit_id
            ),
            None,
        )
        if selected_audit is None:
            any_audit = fetch_one(
                self.connection,
                """
                SELECT audit_id
                FROM daily_line_capacity_plan_audit
                WHERE audit_id = ?
                LIMIT 1
                """,
                (audit_id,),
            )
            if any_audit is None:
                raise not_found(
                    code="DAILY_CAPACITY_AUDIT_NOT_FOUND",
                    message="Daily capacity audit does not exist.",
                    details={"audit_id": audit_id},
                )
            raise bad_request(
                code="REPORT_CAPACITY_COMPARE_AUDIT_SCOPE_MISMATCH",
                message=(
                    "The selected daily capacity audit does not match the report date, "
                    "workshop, line, and process."
                ),
                details={
                    "report_id": report_id,
                    "audit_id": audit_id,
                    "report_local_date": report_local_date,
                    "workshop_code": workshop_code,
                    "line_code": line_code,
                    "process_code": process_code,
                },
            )

        compare_qty = _to_number(selected_audit.get("new_planned_capacity_qty"), 0)
        selected_at = utc_now()
        with transaction(self.connection):
            self.connection.execute(
                """
                UPDATE work_reports
                SET daily_capacity_compare_audit_id = ?,
                    daily_capacity_compare_qty = ?,
                    daily_capacity_compare_selected_at = ?
                WHERE report_id = ?
                """,
                (audit_id, compare_qty, selected_at, report_id),
            )
        return {
            "report_id": report_id,
            "report_local_date": report_local_date,
            "daily_capacity_compare_audit_id": audit_id,
            "daily_capacity_compare_qty": compare_qty,
            "daily_capacity_compare_selected_at": selected_at,
        }

    def delete_reporting(
        self,
        report_id: str,
        actor: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        existing = fetch_one(
            self.connection,
            """
            SELECT report_id, process_code, workshop_code, line_code, report_time
            FROM work_reports
            WHERE report_id = ?
            """,
            (report_id,),
        )
        if existing is None:
            raise not_found(
                code="REPORT_NOT_FOUND",
                message="Report does not exist.",
                details={"report_id": report_id},
            )
        normalized_actor = actor if isinstance(actor, dict) else {}
        workshop_code = str(existing.get("workshop_code") or "").strip().upper()
        line_code = str(existing.get("line_code") or "").strip().upper()
        if self._is_workshop_manager(normalized_actor):
            if not workshop_code or not line_code:
                raise forbidden(
                    code="REPORT_LINE_SCOPE_FORBIDDEN",
                    message="Current workshop manager is not allowed to delete this reporting record.",
                    details={"report_id": report_id},
                )
            self._assert_actor_can_access_line(
                normalized_actor,
                company_code=DEFAULT_COMPANY_CODE,
                workshop_code=workshop_code,
                line_code=line_code,
                missing_user_error_code="REPORT_ACTOR_USER_REQUIRED",
                forbidden_error_code="REPORT_LINE_SCOPE_FORBIDDEN",
                forbidden_message="Current workshop manager is not allowed to delete this reporting record.",
            )
        with transaction(self.connection):
            self.connection.execute(
                "DELETE FROM work_reports WHERE report_id = ?",
                (report_id,),
            )
        process_code = str(existing.get("process_code") or "").strip().upper()
        report_time_text = str(existing.get("report_time") or "").strip()
        if workshop_code and line_code and process_code and report_time_text:
            local_date = _local_date_from_iso_datetime(report_time_text)
            if local_date:
                self.rebuild_line_daily_actual_capacity({"calendar_date": local_date})
        return {"ok": True}

    def list_schedule_versions(self) -> dict[str, Any]:
        rows = fetch_all(
            self.connection,
            """
            SELECT version_no, status, status_name_cn, strategy_code, created_at, published_at
            FROM schedule_versions
            ORDER BY
                created_at ASC,
                CAST(
                    CASE
                        WHEN INSTR(version_no, '-D') > 0 THEN SUBSTR(version_no, INSTR(version_no, '-D') + 2)
                        ELSE '0'
                    END AS INTEGER
                ) ASC,
                version_no ASC
            """,
        )
        items: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            item["status_label"] = _schedule_version_status_label(row.get("status"))
            items.append(item)
        return {"items": items}

    def get_schedule_version(self, version_no: str) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT version_no, status, status_name_cn, strategy_code, created_at, published_at
            FROM schedule_versions
            WHERE version_no = ?
            """,
            (version_no,),
        )
        if row is None:
            raise not_found(
                code="SCHEDULE_VERSION_NOT_FOUND",
                message="Schedule version does not exist.",
                details={"version_no": version_no},
            )
        item = dict(row)
        item["status_label"] = _schedule_version_status_label(row.get("status"))
        return item

    def list_schedule_tasks(self, version_no: str) -> dict[str, Any]:
        self.get_schedule_version(version_no)
        rows = fetch_all(
            self.connection,
            """
            SELECT
                production_order_no,
                process_code,
                process_name_cn,
                workshop_code,
                line_code,
                calendar_date,
                shift_code,
                plan_qty,
                plan_start_time
            FROM schedule_tasks
            WHERE version_no = ?
            ORDER BY task_no ASC
            """,
            (version_no,),
        )
        return {
            "items": [
                {
                    "order_no": row["production_order_no"],
                    "process_code": row["process_code"],
                    "process_name_cn": row.get("process_name_cn") or row["process_code"],
                    "workshop_code": row.get("workshop_code"),
                    "line_code": row.get("line_code"),
                    "calendar_date": row["calendar_date"],
                    "shift_code": row["shift_code"],
                    "plan_qty": row["plan_qty"],
                    "plan_start_time": row.get("plan_start_time"),
                }
                for row in rows
            ]
        }

    def get_schedule_algorithm(self, version_no: str) -> dict[str, Any]:
        row = self.get_schedule_version(version_no)
        return {
            "version_no": row["version_no"],
            "strategy_code": row["strategy_code"],
            "algorithm_name_cn": "MVP Schedule",
        }

    def get_schedule_diff(
        self,
        version_no: str,
        compare_with: str | None,
    ) -> dict[str, Any]:
        current_version = self.get_schedule_version(version_no)
        if compare_with:
            compare_version = self.get_schedule_version(compare_with)
        else:
            compare_version = self._pick_schedule_compare_version(version_no)
        if compare_version is None:
            raise bad_request(
                code="SCHEDULE_COMPARE_VERSION_REQUIRED",
                message="compare_with is required when no other comparable version exists.",
            )

        current_tasks = self._list_schedule_task_detail_rows(version_no)
        compare_tasks = self._list_schedule_task_detail_rows(str(compare_version["version_no"]))
        current_orders = self._build_schedule_order_summary_map(current_tasks)
        compare_orders = self._build_schedule_order_summary_map(compare_tasks)

        delivery_changes = self._build_schedule_delivery_changes(
            current_orders=current_orders,
            compare_orders=compare_orders,
        )
        schedule_changes = self._build_schedule_schedule_changes(
            current_orders=current_orders,
            compare_orders=compare_orders,
        )
        line_changes = self._build_schedule_line_changes(
            current_tasks=current_tasks,
            compare_tasks=compare_tasks,
        )
        material_changes = self._build_schedule_material_changes(
            current_orders=current_orders,
            compare_orders=compare_orders,
        )

        summary = {
            "selected_version_no": str(current_version["version_no"]),
            "compare_version_no": str(compare_version["version_no"]),
            "changed_order_count": len(schedule_changes["items"]),
            "added_order_count": sum(1 for item in schedule_changes["items"] if item["change_type"] == "ADDED"),
            "removed_order_count": sum(1 for item in schedule_changes["items"] if item["change_type"] == "REMOVED"),
            "earlier_finish_count": sum(
                1 for item in delivery_changes["items"] if item["change_type"] == "EARLIER_FINISH"
            ),
            "later_finish_count": sum(
                1 for item in delivery_changes["items"] if item["change_type"] == "LATER_FINISH"
            ),
            "start_changed_count": sum(
                1
                for item in schedule_changes["items"]
                if item["selected_start_date"] != item["compare_start_date"]
            ),
            "material_risk_increase_count": sum(
                1
                for item in material_changes["items"]
                if item["risk_change"] in {"NEW_SHORTAGE", "SHORTAGE_WORSE"}
            ),
            "line_change_available": bool(line_changes["available"]),
            "line_change_count": len(line_changes["items"]),
        }

        return {
            "selected_version": current_version,
            "compare_version": compare_version,
            "summary": summary,
            "delivery_changes": delivery_changes,
            "schedule_changes": schedule_changes,
            "line_changes": line_changes,
            "material_changes": material_changes,
        }

    def get_schedule_material_shortages(self, version_no: str) -> dict[str, Any]:
        current_version = self.get_schedule_version(version_no)
        shortage_analysis = self._build_schedule_shortage_analysis(version_no)
        return {
            "summary": {
                "version_no": str(current_version["version_no"]),
                "shortage_material_count": int(
                    shortage_analysis["summary"].get("shortage_material_count") or 0
                ),
                "impacted_order_count": int(
                    shortage_analysis["summary"].get("impacted_order_count") or 0
                ),
            },
            "items": shortage_analysis["items"],
        }

    def _pick_reference_schedule_version_no(self) -> str | None:
        row = fetch_one(
            self.connection,
            """
            SELECT version_no
            FROM schedule_versions
            WHERE UPPER(TRIM(COALESCE(status, ''))) = 'PUBLISHED'
            ORDER BY
                COALESCE(NULLIF(TRIM(COALESCE(published_at, '')), ''), created_at) DESC,
                created_at DESC,
                CAST(
                    CASE
                        WHEN INSTR(version_no, '-D') > 0 THEN SUBSTR(version_no, INSTR(version_no, '-D') + 2)
                        ELSE '0'
                    END AS INTEGER
                ) DESC,
                version_no DESC
            LIMIT 1
            """,
        )
        if row is not None:
            version_no = str(row.get("version_no") or "").strip()
            if version_no:
                return version_no

        row = fetch_one(
            self.connection,
            """
            SELECT version_no
            FROM schedule_versions
            ORDER BY
                created_at DESC,
                CAST(
                    CASE
                        WHEN INSTR(version_no, '-D') > 0 THEN SUBSTR(version_no, INSTR(version_no, '-D') + 2)
                        ELSE '0'
                    END AS INTEGER
                ) DESC,
                version_no DESC
            LIMIT 1
            """,
        )
        if row is None:
            return None
        version_no = str(row.get("version_no") or "").strip()
        return version_no or None

    def _resolve_order_pool_version_no(self, version_no: str | None) -> str | None:
        normalized_version_no = str(version_no or "").strip()
        if normalized_version_no:
            self.get_schedule_version(normalized_version_no)
            return normalized_version_no
        return self._pick_reference_schedule_version_no()

    def _pick_published_schedule_version_no(self) -> str | None:
        row = fetch_one(
            self.connection,
            """
            SELECT version_no
            FROM schedule_versions
            WHERE UPPER(TRIM(COALESCE(status, ''))) = 'PUBLISHED'
            ORDER BY
                COALESCE(NULLIF(TRIM(COALESCE(published_at, '')), ''), created_at) DESC,
                created_at DESC,
                version_no DESC
            LIMIT 1
            """,
        )
        if row is None:
            return None
        version_no = str(row.get("version_no") or "").strip()
        return version_no or None

    def _pick_latest_draft_schedule_version_no(self) -> str | None:
        row = fetch_one(
            self.connection,
            """
            SELECT version_no
            FROM schedule_versions
            WHERE UPPER(TRIM(COALESCE(status, ''))) = 'DRAFT'
            ORDER BY created_at DESC, version_no DESC
            LIMIT 1
            """,
        )
        if row is None:
            return None
        version_no = str(row.get("version_no") or "").strip()
        return version_no or None

    def _list_schedule_task_rows_by_version_order(
        self,
        *,
        version_no: str,
        order_no: str,
    ) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                st.task_no,
                st.version_no,
                st.production_order_no,
                st.process_code,
                COALESCE(st.process_name_cn, st.process_code) AS process_name_cn,
                st.workshop_code,
                st.line_code,
                st.calendar_date,
                st.shift_code,
                st.plan_qty,
                st.plan_start_time,
                po.material_code,
                COALESCE(po.material_name, po.material_code) AS material_name,
                po.production_qty
            FROM schedule_tasks st
            JOIN production_orders po
              ON po.production_order_no = st.production_order_no
            WHERE st.version_no = ?
              AND st.production_order_no = ?
            ORDER BY st.task_no ASC
            """,
            (version_no, order_no),
        )

    def _list_schedule_task_rows_by_version_process(
        self,
        *,
        version_no: str,
        process_code: str,
    ) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                st.task_no,
                st.version_no,
                st.production_order_no,
                st.process_code,
                COALESCE(st.process_name_cn, st.process_code) AS process_name_cn,
                st.workshop_code,
                st.line_code,
                st.calendar_date,
                st.shift_code,
                st.plan_qty,
                st.plan_start_time,
                po.material_code,
                COALESCE(po.material_name, po.material_code) AS material_name,
                po.production_qty
            FROM schedule_tasks st
            JOIN production_orders po
              ON po.production_order_no = st.production_order_no
            WHERE st.version_no = ?
              AND UPPER(TRIM(COALESCE(st.process_code, ''))) = ?
            ORDER BY st.task_no ASC
            """,
            (version_no, process_code),
        )

    def _parse_iso_datetime_strict(
        self,
        value: object,
        *,
        error_code: str,
        message: str,
        details: dict[str, Any] | None = None,
    ) -> datetime:
        text = str(value or "").strip()
        if not text:
            raise server_error(
                code=error_code,
                message=message,
                details=details or {},
            )
        try:
            return datetime.fromisoformat(text)
        except ValueError as exc:
            raise server_error(
                code=error_code,
                message=message,
                details={
                    **(details or {}),
                    "value": text,
                },
            ) from exc

    def _task_start_time_from_schedule_row(self, row: dict[str, Any]) -> str:
        plan_start_time = str(row.get("plan_start_time") or "").strip()
        if plan_start_time:
            return plan_start_time

        calendar_date = _normalize_date_text(row.get("calendar_date"))
        if not calendar_date:
            raise server_error(
                code="SCHEDULE_TASK_CALENDAR_DATE_INVALID",
                message="schedule task calendar_date is required.",
                details={
                    "version_no": row.get("version_no"),
                    "task_no": row.get("task_no"),
                    "order_no": row.get("production_order_no"),
                },
            )
        shift_code = _normalize_shift_code(row.get("shift_code"))
        return str(
            _iso_at(
                calendar_date,
                "08:00:00" if shift_code == "DAY" else "20:00:00",
            )
            or ""
        )

    def _task_finish_time_from_schedule_row(self, row: dict[str, Any]) -> str:
        calendar_date = _normalize_date_text(row.get("calendar_date"))
        if not calendar_date:
            raise server_error(
                code="SCHEDULE_TASK_CALENDAR_DATE_INVALID",
                message="schedule task calendar_date is required.",
                details={
                    "version_no": row.get("version_no"),
                    "task_no": row.get("task_no"),
                    "order_no": row.get("production_order_no"),
                },
            )
        calendar_day = date.fromisoformat(calendar_date)
        shift_code = _normalize_shift_code(row.get("shift_code"))
        if shift_code == "NIGHT":
            return str(
                _iso_at(
                    (calendar_day + timedelta(days=1)).isoformat(),
                    "08:00:00",
                )
                or ""
            )
        return str(_iso_at(calendar_day.isoformat(), "20:00:00") or "")

    def _build_order_process_timeline_items(
        self,
        *,
        task_rows: list[dict[str, Any]],
        route_rows: list[dict[str, Any]],
        process_stats_by_code: dict[str, dict[str, Any]],
    ) -> list[dict[str, Any]]:
        route_by_process: dict[str, dict[str, Any]] = {}
        for route_row in route_rows:
            process_code = str(route_row.get("process_code") or "").strip().upper()
            if process_code and process_code not in route_by_process:
                route_by_process[process_code] = route_row

        grouped: dict[str, dict[str, Any]] = {}
        for task_row in task_rows:
            process_code = str(task_row.get("process_code") or "").strip().upper()
            if not process_code:
                raise server_error(
                    code="SCHEDULE_TASK_PROCESS_CODE_MISSING",
                    message="schedule task process_code is required.",
                    details={
                        "version_no": task_row.get("version_no"),
                        "task_no": task_row.get("task_no"),
                        "order_no": task_row.get("production_order_no"),
                    },
                )
            route_row = route_by_process.get(process_code, {})
            process_name_cn = str(
                task_row.get("process_name_cn")
                or route_row.get("process_name_cn")
                or process_code
            ).strip() or process_code
            start_time = self._task_start_time_from_schedule_row(task_row)
            finish_time = self._task_finish_time_from_schedule_row(task_row)
            start_dt = self._parse_iso_datetime_strict(
                start_time,
                error_code="SCHEDULE_TASK_START_TIME_INVALID",
                message="schedule task start_time is invalid.",
                details={
                    "version_no": task_row.get("version_no"),
                    "task_no": task_row.get("task_no"),
                    "order_no": task_row.get("production_order_no"),
                    "process_code": process_code,
                },
            )
            finish_dt = self._parse_iso_datetime_strict(
                finish_time,
                error_code="SCHEDULE_TASK_FINISH_TIME_INVALID",
                message="schedule task finish_time is invalid.",
                details={
                    "version_no": task_row.get("version_no"),
                    "task_no": task_row.get("task_no"),
                    "order_no": task_row.get("production_order_no"),
                    "process_code": process_code,
                },
            )
            if finish_dt < start_dt:
                raise server_error(
                    code="SCHEDULE_TASK_TIME_RANGE_INVALID",
                    message="schedule task finish_time must be >= start_time.",
                    details={
                        "version_no": task_row.get("version_no"),
                        "task_no": task_row.get("task_no"),
                        "order_no": task_row.get("production_order_no"),
                        "process_code": process_code,
                        "start_time": start_time,
                        "finish_time": finish_time,
                    },
                )
            current = grouped.get(process_code)
            if current is None:
                grouped[process_code] = {
                    "process_code": process_code,
                    "process_name_cn": process_name_cn,
                    "sequence_no": int(_to_number(route_row.get("sequence_no"), 10**9)),
                    "start_time": start_time,
                    "finish_time": finish_time,
                    "start_dt": start_dt,
                    "finish_dt": finish_dt,
                    "plan_qty_total": _to_number(task_row.get("plan_qty"), 0),
                }
                continue
            if start_dt < current["start_dt"]:
                current["start_dt"] = start_dt
                current["start_time"] = start_time
            if finish_dt > current["finish_dt"]:
                current["finish_dt"] = finish_dt
                current["finish_time"] = finish_time
            current["plan_qty_total"] = _to_number(current.get("plan_qty_total"), 0) + _to_number(
                task_row.get("plan_qty"),
                0,
            )

        items: list[dict[str, Any]] = []
        for current in grouped.values():
            process_code = str(current.get("process_code") or "").strip().upper()
            process_stats = process_stats_by_code.get(process_code)
            if process_stats is None:
                raise server_error(
                    code="SCHEDULE_PROCESS_STATS_MISSING",
                    message="schedule process stats are required.",
                    details={
                        "process_code": process_code,
                    },
                )
            duration_hours = (
                (current["finish_dt"] - current["start_dt"]).total_seconds() / 3600
            )
            items.append(
                {
                    "process_code": process_code,
                    "process_name_cn": current["process_name_cn"],
                    "start_time": current["start_time"],
                    "finish_time": current["finish_time"],
                    "duration_hours": round(duration_hours, 4),
                    "related_order_count": int(
                        _to_number(process_stats.get("related_order_count"), 0)
                    ),
                    "plan_qty_total": round(
                        _to_number(process_stats.get("plan_qty_total"), 0),
                        4,
                    ),
                    "is_bottleneck": 0,
                    "sequence_no": int(_to_number(current.get("sequence_no"), 10**9)),
                }
            )

        items.sort(
            key=lambda item: (
                int(_to_number(item.get("sequence_no"), 10**9)),
                str(item.get("process_code") or ""),
            )
        )
        if items:
            max_duration_hours = max(
                _to_number(item.get("duration_hours"), 0)
                for item in items
            )
            bottleneck_candidates = [
                item
                for item in items
                if _to_number(item.get("duration_hours"), 0) == max_duration_hours
            ]
            bottleneck = min(
                bottleneck_candidates,
                key=lambda item: (
                    int(_to_number(item.get("sequence_no"), 10**9)),
                    str(item.get("process_code") or ""),
                ),
            )
            bottleneck_code = str(bottleneck.get("process_code") or "").strip().upper()
            for item in items:
                item["is_bottleneck"] = (
                    1
                    if str(item.get("process_code") or "").strip().upper()
                    == bottleneck_code
                    else 0
                )
        for item in items:
            item.pop("sequence_no", None)
        return items

    def _build_schedule_process_stats_by_code(
        self,
        task_rows: list[dict[str, Any]],
    ) -> dict[str, dict[str, Any]]:
        grouped: dict[str, dict[str, Any]] = {}
        for task_row in task_rows:
            process_code = str(task_row.get("process_code") or "").strip().upper()
            if not process_code:
                raise server_error(
                    code="SCHEDULE_TASK_PROCESS_CODE_MISSING",
                    message="schedule task process_code is required.",
                    details={
                        "version_no": task_row.get("version_no"),
                        "task_no": task_row.get("task_no"),
                        "order_no": task_row.get("production_order_no"),
                    },
                )
            order_no = str(task_row.get("production_order_no") or "").strip()
            if not order_no:
                raise server_error(
                    code="SCHEDULE_TASK_ORDER_NO_MISSING",
                    message="schedule task production_order_no is required.",
                    details={
                        "version_no": task_row.get("version_no"),
                        "task_no": task_row.get("task_no"),
                        "process_code": process_code,
                    },
                )
            current = grouped.get(process_code)
            if current is None:
                current = {
                    "plan_qty_total": 0.0,
                    "order_no_set": set(),
                }
                grouped[process_code] = current
            current["plan_qty_total"] = _to_number(
                current.get("plan_qty_total"),
                0,
            ) + _to_number(task_row.get("plan_qty"), 0)
            order_no_set = current.get("order_no_set")
            if not isinstance(order_no_set, set):
                raise server_error(
                    code="SCHEDULE_PROCESS_ORDER_SET_INVALID",
                    message="schedule process order set is invalid.",
                    details={
                        "process_code": process_code,
                    },
                )
            order_no_set.add(order_no)

        return {
            process_code: {
                "related_order_count": len(current.get("order_no_set") or set()),
                "plan_qty_total": round(_to_number(current.get("plan_qty_total"), 0), 4),
            }
            for process_code, current in grouped.items()
        }

    def _build_empty_selected_process_detail(
        self,
        *,
        process_code: str | None,
    ) -> dict[str, Any] | None:
        if not process_code:
            return None
        return {
            "process_code": process_code,
            "process_name_cn": process_code,
            "current_order_finish_time": None,
            "order_timeline_items": [],
            "message": "暂无排产任务数据",
        }

    def _build_selected_process_timeline_detail(
        self,
        *,
        version_no: str,
        order_no: str,
        process_code: str,
        process_items: list[dict[str, Any]],
    ) -> dict[str, Any]:
        normalized_process_code = str(process_code or "").strip().upper()
        selected_process_row = next(
            (
                item
                for item in process_items
                if str(item.get("process_code") or "").strip().upper() == normalized_process_code
            ),
            None,
        )
        process_name_cn = str(
            (selected_process_row or {}).get("process_name_cn") or normalized_process_code
        ).strip() or normalized_process_code

        task_rows = self._list_schedule_task_rows_by_version_process(
            version_no=version_no,
            process_code=normalized_process_code,
        )
        grouped: dict[str, dict[str, Any]] = {}
        for task_row in task_rows:
            current_order_no = str(task_row.get("production_order_no") or "").strip()
            if not current_order_no:
                continue
            start_time = self._task_start_time_from_schedule_row(task_row)
            finish_time = self._task_finish_time_from_schedule_row(task_row)
            start_dt = self._parse_iso_datetime_strict(
                start_time,
                error_code="SCHEDULE_TASK_START_TIME_INVALID",
                message="schedule task start_time is invalid.",
                details={
                    "version_no": task_row.get("version_no"),
                    "task_no": task_row.get("task_no"),
                    "order_no": current_order_no,
                    "process_code": normalized_process_code,
                },
            )
            finish_dt = self._parse_iso_datetime_strict(
                finish_time,
                error_code="SCHEDULE_TASK_FINISH_TIME_INVALID",
                message="schedule task finish_time is invalid.",
                details={
                    "version_no": task_row.get("version_no"),
                    "task_no": task_row.get("task_no"),
                    "order_no": current_order_no,
                    "process_code": normalized_process_code,
                },
            )
            if finish_dt < start_dt:
                raise server_error(
                    code="SCHEDULE_TASK_TIME_RANGE_INVALID",
                    message="schedule task finish_time must be >= start_time.",
                    details={
                        "version_no": task_row.get("version_no"),
                        "task_no": task_row.get("task_no"),
                        "order_no": current_order_no,
                        "process_code": normalized_process_code,
                        "start_time": start_time,
                        "finish_time": finish_time,
                    },
                )

            current = grouped.get(current_order_no)
            if current is None:
                grouped[current_order_no] = {
                    "order_no": current_order_no,
                    "product_code": str(task_row.get("material_code") or "").strip().upper() or None,
                    "product_name_cn": str(
                        task_row.get("material_name") or task_row.get("material_code") or ""
                    ).strip() or None,
                    "start_time": start_time,
                    "finish_time": finish_time,
                    "start_dt": start_dt,
                    "finish_dt": finish_dt,
                    "plan_qty_total": _to_number(task_row.get("plan_qty"), 0),
                }
                continue
            if start_dt < current["start_dt"]:
                current["start_dt"] = start_dt
                current["start_time"] = start_time
            if finish_dt > current["finish_dt"]:
                current["finish_dt"] = finish_dt
                current["finish_time"] = finish_time
            current["plan_qty_total"] = _to_number(current.get("plan_qty_total"), 0) + _to_number(
                task_row.get("plan_qty"),
                0,
            )

        order_timeline_items: list[dict[str, Any]] = []
        for row in grouped.values():
            duration_hours = (row["finish_dt"] - row["start_dt"]).total_seconds() / 3600
            order_timeline_items.append(
                {
                    "order_no": row["order_no"],
                    "product_code": row["product_code"],
                    "product_name_cn": row["product_name_cn"],
                    "start_time": row["start_time"],
                    "finish_time": row["finish_time"],
                    "duration_hours": round(duration_hours, 4),
                    "plan_qty_total": _to_number(row.get("plan_qty_total"), 0),
                }
            )
        order_timeline_items.sort(
            key=lambda item: (
                str(item.get("start_time") or ""),
                str(item.get("order_no") or ""),
            )
        )

        current_order_item = next(
            (
                item
                for item in order_timeline_items
                if str(item.get("order_no") or "").strip() == order_no
            ),
            None,
        )
        current_order_finish_time = (
            str(current_order_item.get("finish_time") or "").strip()
            if current_order_item
            else (
                str((selected_process_row or {}).get("finish_time") or "").strip()
                or None
            )
        )

        detail: dict[str, Any] = {
            "process_code": normalized_process_code,
            "process_name_cn": process_name_cn,
            "current_order_finish_time": current_order_finish_time,
            "order_timeline_items": order_timeline_items,
        }
        if len(order_timeline_items) == 0:
            detail["message"] = "暂无排产任务数据"
        return detail

    def _pick_schedule_compare_version(self, version_no: str) -> dict[str, Any] | None:
        rows = fetch_all(
            self.connection,
            """
            SELECT version_no, status, status_name_cn, strategy_code, created_at, published_at
            FROM schedule_versions
            WHERE version_no <> ?
            ORDER BY created_at ASC, version_no ASC
            """,
            (version_no,),
        )
        if not rows:
            return None
        published = [row for row in rows if str(row.get("status") or "").strip().upper() == "PUBLISHED"]
        if published:
            return published[-1]
        return rows[-1]

    def _list_schedule_task_detail_rows(self, version_no: str) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                st.version_no,
                st.task_no,
                st.production_order_no,
                st.process_code,
                COALESCE(st.process_name_cn, st.process_code) AS process_name_cn,
                st.workshop_code,
                st.line_code,
                st.calendar_date,
                st.shift_code,
                st.plan_qty,
                st.plan_start_time,
                po.material_code,
                COALESCE(po.material_name, po.material_code) AS material_name,
                po.production_qty,
                po.planned_end_date
            FROM schedule_tasks st
            JOIN production_orders po
              ON po.production_order_no = st.production_order_no
            WHERE st.version_no = ?
            ORDER BY st.task_no ASC
            """,
            (version_no,),
        )

    def _build_schedule_order_summary_map(
        self,
        task_rows: list[dict[str, Any]],
    ) -> dict[str, dict[str, Any]]:
        out: dict[str, dict[str, Any]] = {}
        for row in task_rows:
            order_no = str(row.get("production_order_no") or "").strip()
            if not order_no:
                continue
            current = out.get(order_no)
            if current is None:
                current = {
                    "order_no": order_no,
                    "material_code": str(row.get("material_code") or "").strip(),
                    "material_name": str(row.get("material_name") or row.get("material_code") or "").strip(),
                    "order_qty": _to_number(row.get("production_qty"), 0),
                    "due_date": _normalize_date_text(row.get("planned_end_date")),
                    "start_date": None,
                    "finish_date": None,
                    "task_count": 0,
                    "planned_qty": 0.0,
                    "signature_parts": [],
                    "process_codes": set(),
                    "date_set": set(),
                    "shift_set": set(),
                }
                out[order_no] = current
            calendar_date = _normalize_date_text(row.get("calendar_date"))
            if calendar_date and (current["start_date"] is None or calendar_date < current["start_date"]):
                current["start_date"] = calendar_date
            if calendar_date and (current["finish_date"] is None or calendar_date > current["finish_date"]):
                current["finish_date"] = calendar_date
            process_code = str(row.get("process_code") or "").strip().upper()
            shift_code = _normalize_shift_code(row.get("shift_code"))
            plan_qty = _to_number(row.get("plan_qty"), 0)
            current["task_count"] += 1
            current["planned_qty"] += plan_qty
            if process_code:
                current["process_codes"].add(process_code)
            if calendar_date:
                current["date_set"].add(calendar_date)
            if shift_code:
                current["shift_set"].add(shift_code)
            current["signature_parts"].append(
                "|".join(
                    [
                        calendar_date or "",
                        process_code,
                        shift_code,
                        f"{plan_qty:.6f}",
                    ]
                )
            )
        for row in out.values():
            row["signature"] = ";".join(sorted(row["signature_parts"]))
            row["process_codes"] = sorted(row["process_codes"])
            row["date_set"] = sorted(row["date_set"])
            row["shift_set"] = sorted(row["shift_set"])
            order_qty = _to_number(row.get("order_qty"), 0)
            if order_qty > 0:
                row["planned_ratio"] = max(0.0, min(1.0, _to_number(row["planned_qty"], 0) / order_qty))
            else:
                row["planned_ratio"] = 0.0
        return out

    def _build_schedule_delivery_changes(
        self,
        *,
        current_orders: dict[str, dict[str, Any]],
        compare_orders: dict[str, dict[str, Any]],
    ) -> dict[str, Any]:
        items: list[dict[str, Any]] = []
        for order_no in sorted(set(current_orders) | set(compare_orders)):
            current = current_orders.get(order_no)
            compare = compare_orders.get(order_no)
            change_type = "UNCHANGED"
            if current and not compare:
                change_type = "ADDED"
            elif compare and not current:
                change_type = "REMOVED"
            elif current and compare:
                current_finish = current.get("finish_date")
                compare_finish = compare.get("finish_date")
                if current_finish and compare_finish and current_finish < compare_finish:
                    change_type = "EARLIER_FINISH"
                elif current_finish and compare_finish and current_finish > compare_finish:
                    change_type = "LATER_FINISH"
            if change_type == "UNCHANGED":
                continue
            selected_finish = current.get("finish_date") if current else None
            compare_finish = compare.get("finish_date") if compare else None
            selected_due = current.get("due_date") if current else compare.get("due_date") if compare else None
            compare_due = compare.get("due_date") if compare else current.get("due_date") if current else None
            items.append(
                {
                    "order_no": order_no,
                    "material_code": (current or compare or {}).get("material_code"),
                    "material_name": (current or compare or {}).get("material_name"),
                    "change_type": change_type,
                    "selected_finish_date": selected_finish,
                    "compare_finish_date": compare_finish,
                    "selected_due_date": selected_due,
                    "compare_due_date": compare_due,
                    "selected_overdue_days": _days_between(selected_due, selected_finish),
                    "compare_overdue_days": _days_between(compare_due, compare_finish),
                }
            )
        return {"items": items}

    def _build_schedule_schedule_changes(
        self,
        *,
        current_orders: dict[str, dict[str, Any]],
        compare_orders: dict[str, dict[str, Any]],
    ) -> dict[str, Any]:
        items: list[dict[str, Any]] = []
        for order_no in sorted(set(current_orders) | set(compare_orders)):
            current = current_orders.get(order_no)
            compare = compare_orders.get(order_no)
            change_type = "UNCHANGED"
            if current and not compare:
                change_type = "ADDED"
            elif compare and not current:
                change_type = "REMOVED"
            elif current and compare:
                if current.get("start_date") != compare.get("start_date"):
                    change_type = "START_CHANGED"
                elif current.get("finish_date") != compare.get("finish_date"):
                    change_type = "FINISH_CHANGED"
                elif current.get("signature") != compare.get("signature"):
                    change_type = "TASK_CHANGED"
            if change_type == "UNCHANGED":
                continue
            items.append(
                {
                    "order_no": order_no,
                    "material_code": (current or compare or {}).get("material_code"),
                    "material_name": (current or compare or {}).get("material_name"),
                    "change_type": change_type,
                    "selected_start_date": current.get("start_date") if current else None,
                    "compare_start_date": compare.get("start_date") if compare else None,
                    "selected_finish_date": current.get("finish_date") if current else None,
                    "compare_finish_date": compare.get("finish_date") if compare else None,
                    "selected_task_count": int(current.get("task_count") or 0) if current else 0,
                    "compare_task_count": int(compare.get("task_count") or 0) if compare else 0,
                    "selected_planned_qty": _to_number(current.get("planned_qty"), 0) if current else 0,
                    "compare_planned_qty": _to_number(compare.get("planned_qty"), 0) if compare else 0,
                }
            )
        return {"items": items}

    def _build_schedule_line_changes(
        self,
        *,
        current_tasks: list[dict[str, Any]],
        compare_tasks: list[dict[str, Any]],
    ) -> dict[str, Any]:
        binding_rows = fetch_all(
            self.connection,
            """
            SELECT production_order_no, process_code, workshop_code, line_code
            FROM capacity_bindings
            """
        )
        if len(binding_rows) == 0:
            return {
                "available": False,
                "reason": "schedule_tasks does not store workshop_code/line_code, and capacity_bindings is empty.",
                "items": [],
            }

        binding_map: dict[tuple[str, str], tuple[str, str]] = {}
        for row in binding_rows:
            key = (
                str(row.get("production_order_no") or "").strip(),
                str(row.get("process_code") or "").strip().upper(),
            )
            value = (
                str(row.get("workshop_code") or "").strip().upper(),
                str(row.get("line_code") or "").strip().upper(),
            )
            if key in binding_map and binding_map[key] != value:
                return {
                    "available": False,
                    "reason": "capacity_bindings contains multiple workshop/line mappings for the same order and process.",
                    "items": [],
                }
            binding_map[key] = value

        def summarize(task_rows: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
            out: dict[tuple[str, str], dict[str, Any]] = {}
            for row in task_rows:
                order_no = str(row.get("production_order_no") or "").strip()
                process_code = str(row.get("process_code") or "").strip().upper()
                if not order_no or not process_code:
                    continue
                mapping = binding_map.get((order_no, process_code))
                if mapping is None:
                    continue
                workshop_code, line_code = mapping
                key = (order_no, process_code)
                current = out.get(key)
                if current is None:
                    current = {
                        "order_no": order_no,
                        "process_code": process_code,
                        "process_name_cn": str(row.get("process_name_cn") or process_code),
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                        "start_date": None,
                        "finish_date": None,
                    }
                    out[key] = current
                date_text = _normalize_date_text(row.get("calendar_date"))
                if date_text and (current["start_date"] is None or date_text < current["start_date"]):
                    current["start_date"] = date_text
                if date_text and (current["finish_date"] is None or date_text > current["finish_date"]):
                    current["finish_date"] = date_text
            return out

        current_map = summarize(current_tasks)
        compare_map = summarize(compare_tasks)
        items: list[dict[str, Any]] = []
        for key in sorted(set(current_map) | set(compare_map)):
            current = current_map.get(key)
            compare = compare_map.get(key)
            if current is None or compare is None:
                continue
            if (
                current["workshop_code"] == compare["workshop_code"]
                and current["line_code"] == compare["line_code"]
                and current["start_date"] == compare["start_date"]
                and current["finish_date"] == compare["finish_date"]
            ):
                continue
            items.append(
                {
                    "order_no": key[0],
                    "process_code": key[1],
                    "process_name_cn": current.get("process_name_cn") or compare.get("process_name_cn"),
                    "selected_workshop_code": current.get("workshop_code"),
                    "selected_line_code": current.get("line_code"),
                    "compare_workshop_code": compare.get("workshop_code"),
                    "compare_line_code": compare.get("line_code"),
                    "selected_start_date": current.get("start_date"),
                    "selected_finish_date": current.get("finish_date"),
                    "compare_start_date": compare.get("start_date"),
                    "compare_finish_date": compare.get("finish_date"),
                }
            )
        return {"available": True, "reason": "", "items": items}

    def _list_schedule_material_rows(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                mi.production_order_no,
                mi.child_material_code,
                mi.child_material_name,
                mi.issue_qty,
                COALESCE(sc.supply_type_code, mi.supply_type_code, '') AS supply_type_code,
                COALESCE(sc.supply_type_name, mi.supply_type_name, '-') AS supply_type_name,
                COALESCE(ic.inventory_qty, mi.inventory_qty, 0) AS inventory_qty,
                mi.expandable
            FROM material_issue_items mi
            LEFT JOIN inventory_cache ic
              ON ic.material_code = mi.child_material_code
            LEFT JOIN material_supply_cache sc
              ON sc.material_code = mi.child_material_code
            """
        )

    def _list_schedule_bom_child_rows(self, parent_material_code: str) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                bc.parent_material_code,
                bc.child_material_code,
                bc.child_material_name,
                bc.usage_numerator,
                bc.usage_denominator,
                COALESCE(sc.supply_type_code, bc.supply_type_code, '') AS supply_type_code,
                COALESCE(sc.supply_type_name, bc.supply_type_name, '-') AS supply_type_name,
                COALESCE(ic.inventory_qty, bc.inventory_qty, 0) AS inventory_qty,
                bc.expandable
            FROM bom_children bc
            LEFT JOIN inventory_cache ic
              ON ic.material_code = bc.child_material_code
            LEFT JOIN material_supply_cache sc
              ON sc.material_code = bc.child_material_code
            WHERE bc.parent_material_code = ?
            ORDER BY bc.display_order, bc.child_material_code
            """,
            (parent_material_code,),
        )

    def _expand_schedule_material_children(
        self,
        *,
        production_order_no: str,
        parent_material_code: str,
        parent_issue_qty: float,
        bom_child_cache: dict[str, list[dict[str, Any]]],
        lineage: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        normalized_parent_code = str(parent_material_code or "").strip().upper()
        if not normalized_parent_code or parent_issue_qty <= SCHEDULE_NUMBER_EPSILON:
            return []
        child_rows = bom_child_cache.get(normalized_parent_code)
        if child_rows is None:
            child_rows = self._list_schedule_bom_child_rows(normalized_parent_code)
            bom_child_cache[normalized_parent_code] = child_rows
        if not child_rows:
            return []

        out: list[dict[str, Any]] = []
        for child_row in child_rows:
            child_material_code = str(child_row.get("child_material_code") or "").strip().upper()
            if not child_material_code:
                continue
            if child_material_code in lineage:
                raise server_error(
                    code="BOM_CHILDREN_CYCLE_DETECTED",
                    message="BOM children contain a cycle and cannot be used for shortage analysis.",
                    details={
                        "parent_material_code": normalized_parent_code,
                        "child_material_code": child_material_code,
                        "lineage": list(lineage),
                    },
                )

            usage_denominator = _to_number(child_row.get("usage_denominator"), 0)
            usage_numerator = _to_number(child_row.get("usage_numerator"), 0)
            if usage_denominator <= SCHEDULE_NUMBER_EPSILON or usage_numerator <= SCHEDULE_NUMBER_EPSILON:
                continue

            child_issue_qty = parent_issue_qty * usage_numerator / usage_denominator
            if child_issue_qty <= SCHEDULE_NUMBER_EPSILON:
                continue

            derived_row = {
                "production_order_no": production_order_no,
                "child_material_code": child_material_code,
                "child_material_name": child_row.get("child_material_name"),
                "issue_qty": child_issue_qty,
                "supply_type_code": child_row.get("supply_type_code"),
                "supply_type_name": child_row.get("supply_type_name"),
                "inventory_qty": child_row.get("inventory_qty"),
                "expandable": child_row.get("expandable"),
            }
            out.append(derived_row)

            supply_type_code = str(child_row.get("supply_type_code") or "").strip().upper()
            is_expandable = supply_type_code == "SELF_MADE" or int(child_row.get("expandable") or 0) == 1
            if is_expandable:
                out.extend(
                    self._expand_schedule_material_children(
                        production_order_no=production_order_no,
                        parent_material_code=child_material_code,
                        parent_issue_qty=child_issue_qty,
                        bom_child_cache=bom_child_cache,
                        lineage=(*lineage, child_material_code),
                    )
                )
        return out

    def _list_expanded_schedule_material_rows(self) -> list[dict[str, Any]]:
        root_rows = self._list_schedule_material_rows()
        bom_child_cache: dict[str, list[dict[str, Any]]] = {}
        expanded_rows = list(root_rows)
        for row in root_rows:
            production_order_no = str(row.get("production_order_no") or "").strip()
            parent_material_code = str(row.get("child_material_code") or "").strip().upper()
            parent_issue_qty = _to_number(row.get("issue_qty"), 0)
            supply_type_code = str(row.get("supply_type_code") or "").strip().upper()
            is_expandable = supply_type_code == "SELF_MADE" or int(row.get("expandable") or 0) == 1
            if not production_order_no or not parent_material_code or not is_expandable:
                continue
            expanded_rows.extend(
                self._expand_schedule_material_children(
                    production_order_no=production_order_no,
                    parent_material_code=parent_material_code,
                    parent_issue_qty=parent_issue_qty,
                    bom_child_cache=bom_child_cache,
                    lineage=(parent_material_code,),
                )
            )
        return expanded_rows

    def _build_schedule_material_usage_map(
        self,
        order_map: dict[str, dict[str, Any]],
        *,
        material_rows: list[dict[str, Any]] | None = None,
    ) -> dict[str, dict[str, Any]]:
        source_rows = material_rows if material_rows is not None else self._list_expanded_schedule_material_rows()
        out: dict[str, dict[str, Any]] = {}
        for row in source_rows:
            order_no = str(row.get("production_order_no") or "").strip()
            order_summary = order_map.get(order_no)
            if order_summary is None:
                continue
            consume_ratio = _to_number(order_summary.get("planned_ratio"), 0)
            if consume_ratio <= 0:
                continue
            material_code = str(row.get("child_material_code") or "").strip().upper()
            if not material_code:
                continue
            supply_type_code = str(row.get("supply_type_code") or "").strip().upper()
            if supply_type_code == "SELF_MADE":
                continue
            current = out.get(material_code)
            if current is None:
                current = {
                    "material_code": material_code,
                    "material_name": str(row.get("child_material_name") or material_code).strip() or material_code,
                    "supply_type_code": supply_type_code,
                    "supply_type_name": str(row.get("supply_type_name") or "-").strip() or "-",
                    "inventory_qty": _to_number(row.get("inventory_qty"), 0),
                    "planned_consume_qty": 0.0,
                    "order_nos": set(),
                }
                out[material_code] = current
            current["inventory_qty"] = max(current["inventory_qty"], _to_number(row.get("inventory_qty"), 0))
            current["planned_consume_qty"] += _to_number(row.get("issue_qty"), 0) * consume_ratio
            current["order_nos"].add(order_no)
        for current in out.values():
            current["estimated_inventory_qty"] = current["inventory_qty"] - current["planned_consume_qty"]
            current["order_nos"] = sorted(current["order_nos"])
        return out

    def _build_schedule_shortage_analysis(self, version_no: str | None) -> dict[str, Any]:
        if not version_no:
            return {
                "summary": {
                    "shortage_material_count": 0,
                    "impacted_order_count": 0,
                },
                "items": [],
                "order_map": {},
            }

        task_rows = self._list_schedule_task_detail_rows(version_no)
        if len(task_rows) == 0:
            return {
                "summary": {
                    "shortage_material_count": 0,
                    "impacted_order_count": 0,
                },
                "items": [],
                "order_map": {},
            }

        current_orders = self._build_schedule_order_summary_map(task_rows)
        first_task_by_order: dict[str, dict[str, Any]] = {}
        for row in task_rows:
            order_no = str(row.get("production_order_no") or "").strip()
            if order_no and order_no not in first_task_by_order:
                first_task_by_order[order_no] = row

        material_map = self._build_schedule_material_usage_map(
            current_orders,
            material_rows=self._list_expanded_schedule_material_rows(),
        )
        impacted_order_nos: set[str] = set()
        items: list[dict[str, Any]] = []
        order_map: dict[str, dict[str, Any]] = {}

        def impacted_order_sort_key(item: dict[str, Any]) -> tuple[str, str]:
            start_date = str(item.get("shortage_start_date") or "").strip()
            return (start_date or "9999-12-31", str(item.get("order_no") or "").strip())

        for material_code in sorted(material_map):
            current = material_map[material_code]
            estimated_inventory_qty = _to_number(current.get("estimated_inventory_qty"), 0)
            if estimated_inventory_qty >= 0:
                continue
            order_nos = sorted(
                {
                    str(order_no).strip()
                    for order_no in current.get("order_nos", [])
                    if str(order_no).strip()
                }
            )
            impacted_order_nos.update(order_nos)
            impacted_orders: list[dict[str, Any]] = []
            for order_no in order_nos:
                order_summary = current_orders.get(order_no) or {}
                first_task = first_task_by_order.get(order_no) or {}
                impacted_order = {
                    "order_no": order_no,
                    "shortage_start_date": _normalize_date_text(order_summary.get("start_date")),
                    "first_impacted_process_code": str(
                        first_task.get("process_code") or ""
                    ).strip().upper()
                    or None,
                    "first_impacted_process_name_cn": str(
                        first_task.get("process_name_cn")
                        or first_task.get("process_code")
                        or ""
                    ).strip()
                    or None,
                }
                impacted_orders.append(impacted_order)

                current_order_summary = order_map.get(order_no)
                if current_order_summary is None:
                    current_order_summary = {
                        "shortage_material_count": 0,
                        "items": [],
                    }
                    order_map[order_no] = current_order_summary
                current_order_summary["items"].append(
                    {
                        "material_code": material_code,
                        "material_name": str(current.get("material_name") or material_code).strip()
                        or material_code,
                        "shortage_qty": abs(estimated_inventory_qty),
                        **impacted_order,
                    }
                )

            impacted_orders.sort(key=impacted_order_sort_key)
            first_impacted_order = impacted_orders[0] if impacted_orders else {}
            items.append(
                {
                    "material_code": material_code,
                    "material_name": current.get("material_name"),
                    "supply_type_name": current.get("supply_type_name"),
                    "inventory_qty": _to_number(current.get("inventory_qty"), 0),
                    "planned_consume_qty": _to_number(current.get("planned_consume_qty"), 0),
                    "estimated_inventory_qty": estimated_inventory_qty,
                    "shortage_qty": abs(estimated_inventory_qty),
                    "order_nos": order_nos,
                    "impacted_orders": impacted_orders,
                    "first_impacted_order_no": first_impacted_order.get("order_no"),
                    "shortage_start_date": first_impacted_order.get("shortage_start_date"),
                    "first_impacted_process_code": first_impacted_order.get("first_impacted_process_code"),
                    "first_impacted_process_name_cn": first_impacted_order.get(
                        "first_impacted_process_name_cn"
                    ),
                }
            )

        items.sort(
            key=lambda item: (
                -_to_number(item.get("shortage_qty"), 0),
                str(item.get("material_code") or ""),
            )
        )

        for order_summary in order_map.values():
            material_items = (
                order_summary.get("items")
                if isinstance(order_summary.get("items"), list)
                else []
            )
            material_items.sort(
                key=lambda item: (
                    str(item.get("shortage_start_date") or "").strip() or "9999-12-31",
                    -_to_number(item.get("shortage_qty"), 0),
                    str(item.get("material_code") or ""),
                )
            )
            first_item = material_items[0] if material_items else {}
            order_summary["shortage_material_count"] = len(material_items)
            order_summary["first_material_code"] = first_item.get("material_code")
            order_summary["first_material_name"] = first_item.get("material_name")
            order_summary["shortage_start_date"] = first_item.get("shortage_start_date")
            order_summary["first_impacted_process_code"] = first_item.get("first_impacted_process_code")
            order_summary["first_impacted_process_name_cn"] = first_item.get(
                "first_impacted_process_name_cn"
            )
            if material_items:
                first_material_name = str(
                    first_item.get("material_name") or first_item.get("material_code") or "物料"
                ).strip() or "物料"
                if len(material_items) == 1:
                    order_summary["summary_text"] = first_material_name
                else:
                    order_summary["summary_text"] = (
                        f"{first_material_name} 等 {len(material_items)} 项"
                    )
            else:
                order_summary["summary_text"] = ""

        return {
            "summary": {
                "shortage_material_count": len(items),
                "impacted_order_count": len(impacted_order_nos),
            },
            "items": items,
            "order_map": order_map,
        }

    def _build_schedule_material_changes(
        self,
        *,
        current_orders: dict[str, dict[str, Any]],
        compare_orders: dict[str, dict[str, Any]],
    ) -> dict[str, Any]:
        material_rows = self._list_expanded_schedule_material_rows()
        current_map = self._build_schedule_material_usage_map(
            current_orders,
            material_rows=material_rows,
        )
        compare_map = self._build_schedule_material_usage_map(
            compare_orders,
            material_rows=material_rows,
        )
        items: list[dict[str, Any]] = []
        for material_code in sorted(set(current_map) | set(compare_map)):
            current = current_map.get(material_code)
            compare = compare_map.get(material_code)
            selected_inventory = _to_number(current.get("inventory_qty"), 0) if current else _to_number(compare.get("inventory_qty"), 0)
            compare_inventory = _to_number(compare.get("inventory_qty"), 0) if compare else selected_inventory
            selected_estimated = _to_number(current.get("estimated_inventory_qty"), selected_inventory) if current else selected_inventory
            compare_estimated = _to_number(compare.get("estimated_inventory_qty"), compare_inventory) if compare else compare_inventory
            selected_consume = _to_number(current.get("planned_consume_qty"), 0) if current else 0
            compare_consume = _to_number(compare.get("planned_consume_qty"), 0) if compare else 0
            if (
                abs(selected_consume - compare_consume) < SCHEDULE_NUMBER_EPSILON
                and abs(selected_estimated - compare_estimated) < SCHEDULE_NUMBER_EPSILON
            ):
                continue
            risk_change = "UNCHANGED"
            if compare_estimated >= 0 > selected_estimated:
                risk_change = "NEW_SHORTAGE"
            elif compare_estimated < 0 and selected_estimated < compare_estimated:
                risk_change = "SHORTAGE_WORSE"
            elif compare_estimated < 0 <= selected_estimated:
                risk_change = "SHORTAGE_RESOLVED"
            elif selected_estimated > compare_estimated:
                risk_change = "PRESSURE_RELIEVED"
            elif selected_estimated < compare_estimated:
                risk_change = "PRESSURE_INCREASED"
            items.append(
                {
                    "material_code": material_code,
                    "material_name": (current or compare or {}).get("material_name"),
                    "supply_type_name": (current or compare or {}).get("supply_type_name"),
                    "selected_inventory_qty": selected_inventory,
                    "compare_inventory_qty": compare_inventory,
                    "selected_planned_consume_qty": selected_consume,
                    "compare_planned_consume_qty": compare_consume,
                    "selected_estimated_inventory_qty": selected_estimated,
                    "compare_estimated_inventory_qty": compare_estimated,
                    "risk_change": risk_change,
                    "selected_order_nos": current.get("order_nos", []) if current else [],
                    "compare_order_nos": compare.get("order_nos", []) if compare else [],
                }
            )
        items.sort(
            key=lambda item: (
                0 if item["risk_change"] in {"NEW_SHORTAGE", "SHORTAGE_WORSE"} else 1,
                _to_number(item["selected_estimated_inventory_qty"], 0),
                str(item["material_code"]),
            )
        )
        return {"items": items}

    def get_schedule_daily_process_load(self, version_no: str) -> dict[str, Any]:
        self.get_schedule_version(version_no)
        rows = fetch_all(
            self.connection,
            """
            SELECT
                calendar_date,
                process_code,
                COALESCE(MAX(process_name_cn), process_code) AS process_name_cn,
                COUNT(1) AS task_count,
                SUM(plan_qty) AS plan_qty
            FROM schedule_tasks
            WHERE version_no = ?
            GROUP BY calendar_date, process_code
            ORDER BY calendar_date ASC, process_code ASC
            """,
            (version_no,),
        )
        return {"items": rows}

    def get_masterdata_config(self) -> dict[str, Any]:
        self._ensure_masterdata_seeded()
        rules = self.get_schedule_calendar_rules()["data"]
        line_skeletons = self._list_line_skeleton_rows()
        line_topology = self._list_line_topology_rows()
        workshop_manager_users = self._list_enabled_workshop_manager_users()
        workshop_manager_visible_by_user_id = {
            str(row.get("user_id") or "").strip(): (
                1 if int(row.get("visible_flag") or 0) == 1 else 0
            )
            for row in self._list_workshop_manager_visibility_rows()
            if str(row.get("user_id") or "").strip()
        }
        workshop_manager_users = [
            {
                **row,
                "visible_flag": workshop_manager_visible_by_user_id.get(
                    str(row.get("user_id") or "").strip(),
                    1,
                ),
            }
            for row in workshop_manager_users
        ]
        workshop_manager_line_scopes = self._list_workshop_manager_line_scope_rows()
        route_rows = self._list_route_rows()
        process_seen: dict[str, str] = {}
        for row in route_rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if process_code and process_code not in process_seen:
                process_seen[process_code] = str(
                    row.get("process_name_cn") or process_code
                )
        for row in line_topology:
            process_code = str(row.get("process_code") or "").strip().upper()
            if process_code and process_code not in process_seen:
                process_seen[process_code] = process_code
        process_configs = [
            {"process_code": code, "process_name_cn": name}
            for code, name in sorted(process_seen.items())
        ]
        backup_repository = BackupRepository(self.connection)
        backup_config = backup_repository.get_backup_config()
        backup_records = backup_repository.list_backup_records()
        return {
            "data": {
                "horizon_start_date": rules["horizon_start_date"],
                "horizon_days": rules["horizon_days"],
                "skip_statutory_holidays": rules["skip_statutory_holidays"],
                "weekend_rest_mode": rules["weekend_rest_mode"],
                "date_shift_mode_by_date": rules["date_shift_mode_by_date"],
                "process_configs": process_configs,
                "line_skeletons": line_skeletons,
                "line_topology": line_topology,
                "workshop_manager_users": workshop_manager_users,
                "workshop_manager_line_scopes": workshop_manager_line_scopes,
                "resource_pool": [],
                "material_availability": [],
                "backup_config": backup_config,
                "backup_records": backup_records,
            }
        }

    def save_masterdata_config(self, payload: dict[str, Any]) -> dict[str, Any]:
        line_skeletons = payload.get("line_skeletons")
        if not isinstance(line_skeletons, list) or len(line_skeletons) == 0:
            raise bad_request(
                code="LINE_SKELETONS_REQUIRED",
                message="line_skeletons must be a non-empty array.",
            )
        line_topology = payload.get("line_topology")
        if line_topology is None:
            line_topology = []
        if not isinstance(line_topology, list):
            raise bad_request(
                code="LINE_TOPOLOGY_REQUIRED",
                message="line_topology must be an array.",
            )
        workshop_manager_line_scopes = payload.get("workshop_manager_line_scopes")
        if not isinstance(workshop_manager_line_scopes, list):
            raise bad_request(
                code="WORKSHOP_MANAGER_LINE_SCOPES_REQUIRED",
                message="workshop_manager_line_scopes must be an array.",
            )
        workshop_manager_users_payload = payload.get("workshop_manager_users")
        if not isinstance(workshop_manager_users_payload, list):
            raise bad_request(
                code="WORKSHOP_MANAGER_USERS_REQUIRED",
                message="workshop_manager_users must be an array.",
            )
        workshop_manager_users = self._list_enabled_workshop_manager_users()
        workshop_manager_user_ids = {
            str(row.get("user_id") or "").strip()
            for row in workshop_manager_users
            if str(row.get("user_id") or "").strip()
        }
        updated_at = utc_now()
        backup_config_payload = payload.get("backup_config")
        backup_config_update: tuple[int, int, int] | None = None
        if backup_config_payload is not None:
            if not isinstance(backup_config_payload, dict):
                raise bad_request(
                    code="BACKUP_CONFIG_INVALID",
                    message="backup_config must be an object.",
                )
            try:
                enabled_flag = _parse_enabled_flag(
                    backup_config_payload.get("enabled_flag")
                )
            except ValueError as exc:
                raise bad_request(
                    code="BACKUP_CONFIG_ENABLED_FLAG_INVALID",
                    message=str(exc),
                )
            try:
                frequency_minutes = _parse_int_in_range(
                    backup_config_payload.get("frequency_minutes"),
                    field_name="frequency_minutes",
                    min_value=1,
                    max_value=525600,
                )
            except ValueError as exc:
                raise bad_request(
                    code="BACKUP_CONFIG_FREQUENCY_INVALID",
                    message=str(exc),
                )
            try:
                max_backups = _parse_int_in_range(
                    backup_config_payload.get("max_backups"),
                    field_name="max_backups",
                    min_value=1,
                    max_value=1000,
                )
            except ValueError as exc:
                raise bad_request(
                    code="BACKUP_CONFIG_MAX_BACKUPS_INVALID",
                    message=str(exc),
                )
            backup_config_update = (enabled_flag, frequency_minutes, max_backups)
        payload_workshop_manager_user_ids: set[str] = set()
        workshop_manager_visibility_rows: list[tuple[Any, ...]] = []
        for item in workshop_manager_users_payload:
            user_id = str(item.get("user_id") or "").strip()
            if not user_id:
                raise bad_request(
                    code="WORKSHOP_MANAGER_USER_ROW_INVALID",
                    message="workshop_manager_users row must include user_id.",
                )
            if user_id not in workshop_manager_user_ids:
                raise bad_request(
                    code="WORKSHOP_MANAGER_USER_INVALID",
                    message="workshop_manager_users contains unknown or disabled workshop manager user.",
                    details={"user_id": user_id},
                )
            if user_id in payload_workshop_manager_user_ids:
                raise bad_request(
                    code="WORKSHOP_MANAGER_USER_DUPLICATE",
                    message="workshop_manager_users contains duplicate rows.",
                    details={"user_id": user_id},
                )
            payload_workshop_manager_user_ids.add(user_id)
            workshop_manager_visibility_rows.append(
                (
                    user_id,
                    1 if int(item.get("visible_flag") or 0) == 1 else 0,
                    updated_at,
                )
            )
        if payload_workshop_manager_user_ids != workshop_manager_user_ids:
            raise bad_request(
                code="WORKSHOP_MANAGER_USERS_MISMATCH",
                message="workshop_manager_users must include every enabled workshop manager exactly once.",
                details={
                    "missing_user_ids": sorted(workshop_manager_user_ids - payload_workshop_manager_user_ids),
                    "extra_user_ids": sorted(payload_workshop_manager_user_ids - workshop_manager_user_ids),
                },
            )
        skeleton_rows: list[tuple[Any, ...]] = []
        skeleton_map: dict[tuple[str, str, str], dict[str, Any]] = {}
        for item in line_skeletons:
            company_code = str(item.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
            workshop_code = str(item.get("workshop_code") or "").strip().upper()
            line_code = str(item.get("line_code") or "").strip().upper()
            if not workshop_code or not line_code:
                raise bad_request(
                    code="LINE_SKELETON_ROW_INVALID",
                    message="workshop_code and line_code are required in line_skeletons.",
                )
            key = (company_code, workshop_code, line_code)
            if key in skeleton_map:
                raise bad_request(
                    code="LINE_SKELETON_DUPLICATE",
                    message="line_skeletons contains duplicate workshop/line rows.",
                    details={
                        "company_code": company_code,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                    },
                )
            skeleton_row = {
                "company_code": company_code,
                "workshop_code": workshop_code,
                "workshop_name": str(item.get("workshop_name") or workshop_code).strip() or workshop_code,
                "line_code": line_code,
                "line_name": str(item.get("line_name") or line_code).strip() or line_code,
                "enabled_flag": 1 if int(item.get("enabled_flag") or 0) == 1 else 0,
            }
            skeleton_map[key] = skeleton_row
            skeleton_rows.append(
                (
                    company_code,
                    skeleton_row["workshop_code"],
                    skeleton_row["workshop_name"],
                    skeleton_row["line_code"],
                    skeleton_row["line_name"],
                    skeleton_row["enabled_flag"],
                    updated_at,
                )
            )
        rows = []
        for item in line_topology:
            company_code = str(item.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
            workshop_code = str(item.get("workshop_code") or "").strip().upper()
            line_code = str(item.get("line_code") or "").strip().upper()
            process_code = str(item.get("process_code") or "").strip().upper()
            if not workshop_code or not line_code or not process_code:
                raise bad_request(
                    code="LINE_TOPOLOGY_ROW_INVALID",
                    message="workshop_code, line_code and process_code are required.",
                )
            skeleton_row = skeleton_map.get((company_code, workshop_code, line_code))
            if skeleton_row is None:
                raise bad_request(
                    code="LINE_TOPOLOGY_SKELETON_MISSING",
                    message="line_topology row must reference an existing line_skeleton.",
                    details={
                        "company_code": company_code,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                        "process_code": process_code,
                    },
                )
            capacity_per_shift = _to_number(item.get("capacity_per_shift"), 0)
            required_workers = int(_to_number(item.get("required_workers"), 0))
            required_machines = int(_to_number(item.get("required_machines"), 0))
            if capacity_per_shift <= 0 or required_workers <= 0 or required_machines < 0:
                raise bad_request(
                    code="LINE_TOPOLOGY_CAPACITY_INVALID",
                    message="capacity_per_shift and required_workers must be greater than 0, required_machines must be >= 0.",
                    details={
                        "company_code": company_code,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                        "process_code": process_code,
                    },
                )
            rows.append(
                (
                    company_code,
                    workshop_code,
                    skeleton_row["workshop_name"],
                    line_code,
                    skeleton_row["line_name"],
                    process_code,
                    capacity_per_shift,
                    required_workers,
                    required_machines,
                    int(item.get("enabled_flag") or 0),
                    updated_at,
                )
            )
        scope_rows: list[tuple[Any, ...]] = []
        scope_seen: set[tuple[str, str, str, str]] = set()
        for item in workshop_manager_line_scopes:
            user_id = str(item.get("user_id") or "").strip()
            company_code = str(item.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() or DEFAULT_COMPANY_CODE
            workshop_code = str(item.get("workshop_code") or "").strip().upper()
            line_code = str(item.get("line_code") or "").strip().upper()
            if not user_id or not workshop_code or not line_code:
                raise bad_request(
                    code="WORKSHOP_MANAGER_LINE_SCOPE_ROW_INVALID",
                    message="user_id, workshop_code and line_code are required in workshop_manager_line_scopes.",
                )
            if user_id not in workshop_manager_user_ids:
                raise bad_request(
                    code="WORKSHOP_MANAGER_USER_INVALID",
                    message="workshop_manager_line_scopes contains unknown or disabled workshop manager user.",
                    details={"user_id": user_id},
                )
            scope_key = (user_id, company_code, workshop_code, line_code)
            if scope_key in scope_seen:
                raise bad_request(
                    code="WORKSHOP_MANAGER_LINE_SCOPE_DUPLICATE",
                    message="workshop_manager_line_scopes contains duplicate rows.",
                    details={
                        "user_id": user_id,
                        "company_code": company_code,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                    },
                )
            scope_seen.add(scope_key)
            if (company_code, workshop_code, line_code) not in skeleton_map:
                raise bad_request(
                    code="WORKSHOP_MANAGER_LINE_SCOPE_SKELETON_MISSING",
                    message="workshop_manager_line_scopes row must reference an existing line_skeleton.",
                    details={
                        "user_id": user_id,
                        "company_code": company_code,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                    },
                )
            scope_rows.append(
                (
                    user_id,
                    company_code,
                    workshop_code,
                    line_code,
                    updated_at,
                )
            )
        with transaction(self.connection):
            self.connection.execute("DELETE FROM masterdata_line_skeletons")
            self.connection.executemany(
                """
                INSERT INTO masterdata_line_skeletons (
                    company_code,
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    enabled_flag,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                skeleton_rows,
            )
            self.connection.execute("DELETE FROM masterdata_line_topology")
            if rows:
                self.connection.executemany(
                    """
                    INSERT INTO masterdata_line_topology (
                        company_code,
                        workshop_code,
                        workshop_name,
                        line_code,
                        line_name,
                        process_code,
                        capacity_per_shift,
                        required_workers,
                        required_machines,
                        enabled_flag,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    rows,
                )
            self.connection.execute("DELETE FROM app_user_line_scopes")
            if scope_rows:
                self.connection.executemany(
                    """
                    INSERT INTO app_user_line_scopes (
                        user_id,
                        company_code,
                        workshop_code,
                        line_code,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?)
                    """,
                    scope_rows,
                )
            self.connection.execute("DELETE FROM masterdata_workshop_manager_visibility")
            if workshop_manager_visibility_rows:
                self.connection.executemany(
                    """
                    INSERT INTO masterdata_workshop_manager_visibility (
                        user_id,
                        visible_flag,
                        updated_at
                    ) VALUES (?, ?, ?)
                    """,
                    workshop_manager_visibility_rows,
                )
            if backup_config_update is not None:
                enabled_flag, frequency_minutes, max_backups = backup_config_update
                BackupRepository(self.connection).upsert_backup_config(
                    enabled_flag=enabled_flag,
                    frequency_minutes=frequency_minutes,
                    max_backups=max_backups,
                    updated_at=updated_at,
                )
        return self.get_masterdata_config()

    def list_process_routes(self) -> dict[str, Any]:
        self._ensure_masterdata_seeded()
        return {"items": self._list_route_rows()}

    def create_process_routes(self, payload: dict[str, Any]) -> dict[str, Any]:
        product_code = str(payload.get("product_code") or "").strip().upper()
        if not product_code:
            raise bad_request(
                code="PRODUCT_CODE_REQUIRED",
                message="product_code is required.",
            )
        existing = fetch_one(
            self.connection,
            """
            SELECT 1
            FROM masterdata_process_routes
            WHERE product_code = ?
            LIMIT 1
            """,
            (product_code,),
        )
        if existing is not None:
            raise bad_request(
                code="ROUTE_ALREADY_EXISTS",
                message="Route already exists for product.",
                details={"product_code": product_code},
            )
        self._replace_process_routes(product_code, payload.get("steps"), source_product_code=None)
        return {"ok": True}

    def update_process_routes(self, payload: dict[str, Any]) -> dict[str, Any]:
        product_code = str(payload.get("product_code") or "").strip().upper()
        if not product_code:
            raise bad_request(
                code="PRODUCT_CODE_REQUIRED",
                message="product_code is required.",
            )
        self._replace_process_routes(product_code, payload.get("steps"), source_product_code=product_code)
        return {"ok": True}

    def copy_process_routes(self, payload: dict[str, Any]) -> dict[str, Any]:
        source_product_code = str(payload.get("source_product_code") or "").strip().upper()
        target_product_code = str(payload.get("target_product_code") or "").strip().upper()
        if not source_product_code or not target_product_code:
            raise bad_request(
                code="PRODUCT_CODE_REQUIRED",
                message="source_product_code and target_product_code are required.",
            )
        self._replace_process_routes(
            target_product_code,
            payload.get("steps"),
            source_product_code=source_product_code,
        )
        return {"ok": True}

    def delete_process_routes(self, payload: dict[str, Any]) -> dict[str, Any]:
        product_code = str(payload.get("product_code") or "").strip().upper()
        if not product_code:
            raise bad_request(
                code="PRODUCT_CODE_REQUIRED",
                message="product_code is required.",
            )
        with transaction(self.connection):
            self.connection.execute(
                "DELETE FROM masterdata_process_routes WHERE product_code = ?",
                (product_code,),
            )
        return {"ok": True}

    def get_schedule_calendar_rules(self) -> dict[str, Any]:
        row = self._get_rules_row()
        simulation_state = self._get_simulation_state()
        return {
            "data": {
                "horizon_start_date": row["horizon_start_date"],
                "horizon_days": row["horizon_days"],
                "skip_statutory_holidays": bool(row["skip_statutory_holidays"]),
                "weekend_rest_mode": row["weekend_rest_mode"],
                "date_shift_mode_by_date": loads(row["date_shift_mode_by_date_json"]) or {},
                "current_date": simulation_state["current_date"],
            }
        }

    def save_schedule_calendar_rules(self, payload: dict[str, Any]) -> dict[str, Any]:
        current = self._get_rules_row()
        weekend_rest_mode = self._normalize_weekend_rest_mode(
            payload.get("weekend_rest_mode")
            if "weekend_rest_mode" in payload
            else current["weekend_rest_mode"]
        )
        date_shift_mode_by_date = self._normalize_date_shift_mode_by_date(
            payload.get("date_shift_mode_by_date")
            if "date_shift_mode_by_date" in payload
            else (loads(current["date_shift_mode_by_date_json"]) or {})
        )
        next_row = {
            "singleton_key": RULES_SINGLETON_KEY,
            "horizon_start_date": _normalize_date_text(
                payload.get("horizon_start_date") or current["horizon_start_date"]
            )
            or _today_text(),
            "horizon_days": int(
                _to_number(payload.get("horizon_days"), current["horizon_days"])
            )
            or 31,
            "skip_statutory_holidays": 1
            if payload.get("skip_statutory_holidays") is True
            else (
                current["skip_statutory_holidays"]
                if "skip_statutory_holidays" not in payload
                else 0
            ),
            "weekend_rest_mode": weekend_rest_mode,
            "date_shift_mode_by_date_json": dumps(date_shift_mode_by_date),
            "updated_at": utc_now(),
        }
        with transaction(self.connection):
            self.connection.execute(
                """
                INSERT INTO schedule_calendar_rules (
                    singleton_key,
                    horizon_start_date,
                    horizon_days,
                    skip_statutory_holidays,
                    weekend_rest_mode,
                    date_shift_mode_by_date_json,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(singleton_key) DO UPDATE SET
                    horizon_start_date = excluded.horizon_start_date,
                    horizon_days = excluded.horizon_days,
                    skip_statutory_holidays = excluded.skip_statutory_holidays,
                    weekend_rest_mode = excluded.weekend_rest_mode,
                    date_shift_mode_by_date_json = excluded.date_shift_mode_by_date_json,
                    updated_at = excluded.updated_at
                """,
                (
                    next_row["singleton_key"],
                    next_row["horizon_start_date"],
                    next_row["horizon_days"],
                    next_row["skip_statutory_holidays"],
                    next_row["weekend_rest_mode"],
                    next_row["date_shift_mode_by_date_json"],
                    next_row["updated_at"],
                ),
            )
        return self.get_schedule_calendar_rules()

    def generate_schedule(self, payload: dict[str, Any]) -> dict[str, Any]:
        self.factory.build_inventory_refresh_service().refresh_inventory([])
        self._ensure_masterdata_seeded()
        strategy_code = self._normalize_strategy_code(payload.get("strategy_code"))
        capacity_source_mode = self._normalize_capacity_source_mode(
            payload.get("capacity_source_mode")
        )
        use_order_state_window = self._normalize_use_order_state_window(
            payload.get("use_order_state_window")
        )
        base_version_no = str(payload.get("base_version_no") or "").strip()
        base_schedule_hints: dict[str, dict[str, Any]] = {}
        if base_version_no:
            self.get_schedule_version(base_version_no)
            base_schedule_hints = self._build_base_schedule_hints(base_version_no)
        planning_rules = self._build_planning_rules_from_current_config()
        version_no = self._next_schedule_version_no()
        order_rows = self._list_order_rows()
        route_rows_by_product = self._group_routes_by_product(self._list_route_rows())
        capacity_map = self._get_capacity_map(
            [str(row["production_order_no"]) for row in order_rows]
        )
        states = self._get_order_state_map(
            [str(row["production_order_no"]) for row in order_rows]
        )
        topology_by_process = self._enabled_topology_by_process(
            self._list_line_topology_rows()
        )
        capacity_resolver = self._build_schedule_capacity_resolver(
            capacity_source_mode=capacity_source_mode
        )

        schedule_candidates: list[dict[str, Any]] = []
        fixed_orders: dict[str, dict[str, Any]] = {}
        missing_fixed_order_nos: list[str] = []
        day_mode_cache: dict[str, str] = {}
        simulation_state = self._get_simulation_state()
        simulation_start = _parse_date_or_today(simulation_state.get("current_date"))
        for order_row in order_rows:
            order_no = str(order_row["production_order_no"])
            state_row = states.get(order_no) or {}
            status = str(
                state_row.get("order_status")
                or state_row.get("status")
                or order_row.get("status")
                or ""
            ).strip().upper()
            if status == "DONE":
                continue

            remaining_qty = state_row.get("remaining_qty")
            if remaining_qty is None:
                production_qty = _to_number(order_row.get("production_qty"), 0)
                completed_qty = _to_number(state_row.get("completed_qty"), 0)
                remaining_qty = max(0.0, production_qty - completed_qty)
            remaining_qty_value = max(0.0, _to_number(remaining_qty, 0))
            if remaining_qty_value <= SCHEDULE_NUMBER_EPSILON:
                continue

            lock_flag = int(_to_number(state_row.get("lock_flag"), 0))
            frozen_flag = int(_to_number(state_row.get("frozen_flag"), 0))
            base_hint = base_schedule_hints.get(order_no) or {}
            if lock_flag == 1 or frozen_flag == 1:
                if base_version_no and base_hint:
                    fixed_orders[order_no] = {
                        "order_no": order_no,
                        "product_code": str(order_row["material_code"]),
                    }
                else:
                    missing_fixed_order_nos.append(order_no)
                continue

            process_contexts = self._build_schedule_process_contexts(
                order_no=order_no,
                product_code=str(order_row["material_code"]),
                capacity_rows=capacity_map.get(order_no, []),
                route_rows=route_rows_by_product.get(str(order_row["material_code"]), []),
                topology_by_process=topology_by_process,
                capacity_resolver=capacity_resolver,
            )

            if use_order_state_window:
                start_date_source = (
                    state_row.get("expected_start_date")
                    or order_row.get("planned_start_date")
                    or simulation_start.isoformat()
                )
                due_date_source = (
                    state_row.get("promised_due_date")
                    or order_row.get("planned_end_date")
                    or start_date_source
                )
            else:
                start_date_source = (
                    order_row.get("planned_start_date") or simulation_start.isoformat()
                )
                due_date_source = order_row.get("planned_end_date") or start_date_source
            start_date = _parse_date_or_today(start_date_source)
            due_date = _parse_date_or_today(due_date_source)

            priority_level = _normalize_priority_level(state_row.get("priority_level"), PRIORITY_LEVEL_MAX)
            expected_start_shift = _expected_start_shift_from_datetime_text(
                state_row.get("expected_start_time")
            )
            start_slot = _slot_index_for(start_date, expected_start_shift)

            required_shifts = 0
            min_capacity = None
            total_capacity = 0.0
            for context in process_contexts:
                capacity_per_shift = self._resolve_effective_capacity_per_shift(
                    process_context=context,
                    calendar_date=start_date.isoformat(),
                    capacity_resolver=capacity_resolver,
                )
                required_shifts += self._estimate_required_shifts_for_process(
                    process_context=context,
                    required_qty=remaining_qty_value,
                    start_slot=start_slot,
                    planning_rules=planning_rules,
                    day_mode_cache=day_mode_cache,
                    capacity_resolver=capacity_resolver,
                )
                min_capacity = (
                    capacity_per_shift
                    if min_capacity is None
                    else min(min_capacity, capacity_per_shift)
                )
                total_capacity += capacity_per_shift
            slack_days = (due_date - start_date).days - required_shifts

            schedule_candidates.append(
                {
                    "order_no": order_no,
                    "product_code": str(order_row["material_code"]),
                    "remaining_qty": remaining_qty_value,
                    "start_date": start_date,
                    "due_date": due_date,
                    "start_slot": start_slot,
                    "priority_level": priority_level,
                    "lock_flag": lock_flag,
                    "frozen_flag": frozen_flag,
                    "updated_at": str(order_row.get("updated_at") or ""),
                    "process_contexts": process_contexts,
                    "required_shifts": required_shifts,
                    "slack_days": slack_days,
                    "min_capacity_per_shift": _to_number(min_capacity, 0),
                    "total_capacity_per_shift": total_capacity,
                    "base_first_task_no": int(base_hint.get("first_task_no") or 10**9),
                    "base_process_first_slot": base_hint.get("process_first_slot", {}),
                }
            )

        if missing_fixed_order_nos:
            raise bad_request(
                code="SCHEDULE_FIXED_ORDER_MISSING_FROM_BASE",
                message="存在锁定或冻结订单未出现在基准版本中，已禁止继续重排。",
                details={"order_nos": sorted(set(missing_fixed_order_nos))},
            )

        pending_candidates = self._sort_schedule_candidates(
            strategy_code=strategy_code,
            candidates=schedule_candidates,
        )

        tasks: list[tuple[Any, ...]] = []
        used_capacity_by_slot: dict[tuple[Any, ...], float] = {}
        task_no = 1
        if base_version_no and fixed_orders:
            fixed_order_nos = list(fixed_orders.keys())
            placeholders = ",".join("?" for _ in fixed_order_nos)
            fixed_task_rows = fetch_all(
                self.connection,
                f"""
                SELECT
                    task_no,
                    production_order_no,
                    report_scope,
                    process_code,
                    process_name_cn,
                    workshop_code,
                    line_code,
                    calendar_date,
                    shift_code,
                    plan_qty,
                    plan_start_time
                FROM schedule_tasks
                WHERE version_no = ?
                  AND production_order_no IN ({placeholders})
                ORDER BY task_no ASC
                """,
                tuple([base_version_no, *fixed_order_nos]),
            )

            fixed_process_locations: dict[str, dict[str, set[tuple[str, str]]]] = {}
            for order_no, meta in fixed_orders.items():
                contexts = self._build_schedule_process_contexts(
                    order_no=order_no,
                    product_code=str(meta.get("product_code") or ""),
                    capacity_rows=capacity_map.get(order_no, []),
                    route_rows=route_rows_by_product.get(str(meta.get("product_code") or ""), []),
                    topology_by_process=topology_by_process,
                    capacity_resolver=capacity_resolver,
                )
                location_map: dict[str, set[tuple[str, str]]] = {}
                for context in contexts:
                    process_code = str(context.get("process_code") or "").strip().upper()
                    candidate_contexts = context.get("candidate_contexts") if isinstance(context.get("candidate_contexts"), list) else []
                    if not process_code:
                        continue
                    location_map[process_code] = {
                        (
                            str(item.get("workshop_code") or "").strip().upper(),
                            str(item.get("line_code") or "").strip().upper(),
                        )
                        for item in candidate_contexts
                        if str(item.get("workshop_code") or "").strip() and str(item.get("line_code") or "").strip()
                    }
                fixed_process_locations[order_no] = location_map

            for row in fixed_task_rows:
                order_no = str(row.get("production_order_no") or "").strip()
                process_code = str(row.get("process_code") or "").strip().upper()
                calendar_date = _normalize_date_text(row.get("calendar_date"))
                if not order_no or not process_code or not calendar_date:
                    raise server_error(
                        code="BASE_SCHEDULE_TASK_INVALID",
                        message="Base schedule task is invalid.",
                        details={
                            "version_no": base_version_no,
                            "task_no": row.get("task_no"),
                        },
                    )
                shift_code = _normalize_shift_code(row.get("shift_code"))
                plan_qty = _to_number(row.get("plan_qty"), 0)
                if plan_qty < -SCHEDULE_NUMBER_EPSILON:
                    raise server_error(
                        code="BASE_SCHEDULE_TASK_PLAN_QTY_INVALID",
                        message="Base schedule task plan_qty must be non-negative.",
                        details={
                            "version_no": base_version_no,
                            "task_no": row.get("task_no"),
                            "order_no": order_no,
                            "process_code": process_code,
                            "plan_qty": row.get("plan_qty"),
                        },
                    )

                tasks.append(
                    (
                        version_no,
                        task_no,
                        order_no,
                        process_code,
                        str(row.get("process_name_cn") or process_code),
                        str(row.get("workshop_code") or "").strip().upper() or None,
                        str(row.get("line_code") or "").strip().upper() or None,
                        calendar_date,
                        shift_code,
                        plan_qty,
                        row.get("plan_start_time"),
                    )
                )
                task_no += 1

                workshop_code = str(row.get("workshop_code") or "").strip().upper()
                line_code = str(row.get("line_code") or "").strip().upper()
                allowed_mappings = fixed_process_locations.get(order_no, {}).get(process_code)
                if not workshop_code or not line_code:
                    raise server_error(
                        code="BASE_SCHEDULE_TASK_LOCATION_MISSING",
                        message="Base schedule task workshop/line mapping is missing for fixed schedule.",
                        details={
                            "version_no": base_version_no,
                            "task_no": row.get("task_no"),
                            "order_no": order_no,
                            "process_code": process_code,
                        },
                    )
                if allowed_mappings is not None and (workshop_code, line_code) not in allowed_mappings:
                    raise server_error(
                        code="BASE_SCHEDULE_TASK_LOCATION_INVALID",
                        message="Base schedule task workshop/line is no longer available in candidate lines.",
                        details={
                            "version_no": base_version_no,
                            "task_no": row.get("task_no"),
                            "order_no": order_no,
                            "process_code": process_code,
                            "workshop_code": workshop_code,
                            "line_code": line_code,
                        },
                    )
                slot_index = _slot_index_from_text(calendar_date, shift_code)
                key = self._capacity_usage_key(
                    process_context={
                        "company_code": DEFAULT_COMPANY_CODE,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                        "process_code": process_code,
                    },
                    slot_index=slot_index,
                    calendar_date=calendar_date,
                    capacity_resolver=capacity_resolver,
                )
                used_capacity_by_slot[key] = _to_number(used_capacity_by_slot.get(key), 0) + plan_qty

        while pending_candidates:
            selected_index = self._select_next_candidate_index(
                strategy_code=strategy_code,
                candidates=pending_candidates,
                planning_rules=planning_rules,
                day_mode_cache=day_mode_cache,
                used_capacity_by_slot=used_capacity_by_slot,
                capacity_resolver=capacity_resolver,
            )
            candidate = pending_candidates.pop(selected_index)
            order_no = str(candidate["order_no"])
            next_start_slot = int(candidate["start_slot"])
            process_base_first_slot = candidate.get("base_process_first_slot", {})
            for context in candidate["process_contexts"]:
                context_start_slot = next_start_slot
                process_code = str(context["process_code"])
                base_process_slot = process_base_first_slot.get(process_code)
                if (
                    (int(candidate["lock_flag"]) == 1 or int(candidate["frozen_flag"]) == 1)
                    and base_process_slot is not None
                ):
                    context_start_slot = max(context_start_slot, int(base_process_slot))

                allocations, last_slot = self._allocate_process_tasks(
                    order_no=order_no,
                    process_context=context,
                    required_qty=_to_number(candidate["remaining_qty"], 0),
                    first_slot=context_start_slot,
                    planning_rules=planning_rules,
                    day_mode_cache=day_mode_cache,
                    used_capacity_by_slot=used_capacity_by_slot,
                    capacity_resolver=capacity_resolver,
                )
                for allocation in allocations:
                    calendar_date = allocation["calendar_date"]
                    shift_code = allocation["shift_code"]
                    tasks.append(
                        (
                            version_no,
                            task_no,
                            order_no,
                            process_code,
                            str(context["process_name_cn"]),
                            allocation.get("workshop_code"),
                            allocation.get("line_code"),
                            calendar_date,
                            shift_code,
                            allocation["plan_qty"],
                            _iso_at(
                                calendar_date,
                                "08:00:00" if shift_code == "DAY" else "20:00:00",
                            ),
                        )
                    )
                    task_no += 1
                next_start_slot = last_slot + 1

        order_summary_map = self._build_schedule_order_summary_map_from_generated_tasks(
            tasks=tasks,
            order_rows=order_rows,
        )
        shortage_result = self._build_generated_schedule_material_shortages(
            order_summary_map=order_summary_map,
        )
        if shortage_result["items"]:
            preview = ", ".join(
                f"{item['material_code']}(-{round(_to_number(item.get('shortage_qty'), 0), 4)})"
                for item in shortage_result["items"][:3]
            )
            raise bad_request(
                code="SCHEDULE_MATERIAL_SHORTAGE_BLOCKED",
                message=(
                    f"排产失败：存在 {shortage_result['summary']['shortage_material_count']} 项缺料，"
                    f"影响 {shortage_result['summary']['impacted_order_count']} 张订单。"
                    f"{' 缺料示例：' + preview if preview else ''}"
                ),
                details=shortage_result,
            )

        created_at = utc_now()
        with transaction(self.connection):
            self.connection.execute(
                """
                INSERT INTO schedule_versions (
                    version_no,
                    status,
                    status_name_cn,
                    strategy_code,
                    created_at,
                    published_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    version_no,
                    "DRAFT",
                    _status_name("DRAFT"),
                    strategy_code,
                    created_at,
                    None,
                ),
            )
            if tasks:
                self.connection.executemany(
                    """
                    INSERT INTO schedule_tasks (
                        version_no,
                        task_no,
                        production_order_no,
                        process_code,
                        process_name_cn,
                        workshop_code,
                        line_code,
                        calendar_date,
                        shift_code,
                        plan_qty,
                        plan_start_time
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    tasks,
                )
                self._sync_order_pool_state_schedule_window_from_tasks(
                    tasks=tasks,
                    order_rows=order_rows,
                    states=states,
                )
        return {"version_no": version_no, "capacity_source_mode": capacity_source_mode}

    def publish_schedule_version(self, version_no: str) -> dict[str, Any]:
        self.get_schedule_version(version_no)
        published_at = utc_now()
        with transaction(self.connection):
            self.connection.execute(
                """
                UPDATE schedule_versions
                SET status = 'ARCHIVED',
                    status_name_cn = ?,
                    published_at = COALESCE(published_at, ?)
                WHERE status = 'PUBLISHED'
                  AND version_no <> ?
                """,
                (_status_name("ARCHIVED"), published_at, version_no),
            )
            self.connection.execute(
                """
                UPDATE schedule_versions
                SET status = 'PUBLISHED',
                    status_name_cn = ?,
                    published_at = ?
                WHERE version_no = ?
                """,
                (_status_name("PUBLISHED"), published_at, version_no),
            )
        return {"ok": True}

    def create_dispatch_command(self, payload: dict[str, Any]) -> dict[str, Any]:
        order_no = str(payload.get("target_order_no") or "").strip()
        command_type = str(payload.get("command_type") or "").strip().upper()
        if not order_no or not command_type:
            raise bad_request(
                code="DISPATCH_COMMAND_INVALID",
                message="target_order_no and command_type are required.",
            )
        self._require_order(order_no)
        with transaction(self.connection):
            command_id = self._insert_dispatch_command_record(
                target_order_no=order_no,
                command_type=command_type,
                payload=payload,
            )
        return {"command_id": command_id}

    def approve_dispatch_command(
        self,
        command_id: str,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT command_id, target_order_no, command_type, status
            FROM dispatch_commands
            WHERE command_id = ?
            """,
            (command_id,),
        )
        if row is None:
            raise not_found(
                code="DISPATCH_COMMAND_NOT_FOUND",
                message="Dispatch command does not exist.",
                details={"command_id": command_id},
            )
        with transaction(self.connection):
            self._approve_dispatch_command_record(
                command_id=command_id,
                target_order_no=str(row["target_order_no"]),
                command_type=str(row["command_type"]),
                previous_status=str(row["status"] or ""),
                payload=payload,
            )
        return {"ok": True}

    def batch_dispatch_commands(self, payload: dict[str, Any]) -> dict[str, Any]:
        command_type = str(payload.get("command_type") or "").strip().upper()
        if command_type not in {"LOCK", "UNLOCK", "PRIORITY_UP"}:
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_COMMAND_INVALID",
                message="command_type must be LOCK, UNLOCK or PRIORITY_UP.",
                details={"command_type": command_type or None},
            )

        order_nos = self._normalize_batch_dispatch_order_nos(payload.get("order_nos"))
        if len(order_nos) == 0:
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_EMPTY",
                message="order_nos must contain at least one order.",
            )

        order_rows_by_no = self._get_order_rows_by_nos(order_nos)
        missing_order_nos = [order_no for order_no in order_nos if order_no not in order_rows_by_no]
        if len(missing_order_nos) > 0:
            raise not_found(
                code="ORDER_BATCH_DISPATCH_ORDER_NOT_FOUND",
                message="Some orders do not exist.",
                details={"order_nos": missing_order_nos},
            )

        state_rows_by_no = self._get_order_state_map(order_nos)
        completed_order_nos: list[str] = []
        frozen_order_nos: list[str] = []
        invalid_lock_state_order_nos: list[str] = []
        invalid_priority_state_order_nos: list[str] = []
        for order_no in order_nos:
            order_row = order_rows_by_no[order_no]
            state_row = state_rows_by_no.get(order_no)
            if self._is_order_completed_for_dispatch(order_row, state_row):
                completed_order_nos.append(order_no)
            if int(_to_number((state_row or {}).get("frozen_flag"), 0)) == 1:
                frozen_order_nos.append(order_no)
            is_locked = int(_to_number((state_row or {}).get("lock_flag"), 0)) == 1
            if command_type == "LOCK" and is_locked:
                invalid_lock_state_order_nos.append(order_no)
            if command_type == "UNLOCK" and not is_locked:
                invalid_lock_state_order_nos.append(order_no)
            if (
                command_type == "PRIORITY_UP"
                and _normalize_priority_level((state_row or {}).get("priority_level"), PRIORITY_LEVEL_MAX)
                <= PRIORITY_LEVEL_MIN
            ):
                invalid_priority_state_order_nos.append(order_no)

        if len(completed_order_nos) > 0:
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_COMPLETED",
                message="Completed orders cannot be batch dispatched.",
                details={"order_nos": completed_order_nos, "command_type": command_type},
            )
        if len(frozen_order_nos) > 0:
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_FROZEN",
                message="Frozen orders cannot be batch dispatched.",
                details={"order_nos": frozen_order_nos, "command_type": command_type},
            )
        if len(invalid_lock_state_order_nos) > 0:
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_LOCK_STATE_INVALID",
                message=(
                    "Selected orders are already locked."
                    if command_type == "LOCK"
                    else "Selected orders are not locked."
                ),
                details={"order_nos": invalid_lock_state_order_nos, "command_type": command_type},
            )
        if len(invalid_priority_state_order_nos) > 0:
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_PRIORITY_STATE_INVALID",
                message="Selected orders are already at the highest priority.",
                details={"order_nos": invalid_priority_state_order_nos, "command_type": command_type},
            )

        actor_name = self._resolve_dispatch_actor_name(payload.get("actor"))
        reason = str(payload.get("reason") or "").strip() or (
            "Batch lock production orders"
            if command_type == "LOCK"
            else "Batch unlock production orders"
            if command_type == "UNLOCK"
            else "Batch priority-up production orders"
        )
        decision_reason = str(payload.get("decision_reason") or "").strip() or (
            "Batch dispatch auto approval"
        )
        effective_time = payload.get("effective_time") or utc_now()
        decision_time = payload.get("decision_time") or effective_time
        command_ids: list[str] = []
        with transaction(self.connection):
            for order_no in order_nos:
                command_id = self._insert_dispatch_command_record(
                    target_order_no=order_no,
                    command_type=command_type,
                    payload={
                        "effective_time": effective_time,
                        "reason": reason,
                        "created_by": actor_name,
                    },
                )
                self._approve_dispatch_command_record(
                    command_id=command_id,
                    target_order_no=order_no,
                    command_type=command_type,
                    previous_status="PENDING",
                    payload={
                        "approver": actor_name,
                        "decision": "APPROVED",
                        "decision_reason": decision_reason,
                        "decision_time": decision_time,
                    },
                )
                command_ids.append(command_id)
        return {
            "command_type": command_type,
            "order_nos": order_nos,
            "command_ids": command_ids,
            "count": len(order_nos),
        }

    def advance_simulation_one_day(self, payload: dict[str, Any]) -> dict[str, Any]:
        current = self._get_simulation_state()
        baseline_date = _normalize_date_text(
            current.get("current_date") or payload.get("client_date") or _today_text()
        )
        if baseline_date is None:
            raise server_error(
                code="SIMULATION_CURRENT_DATE_INVALID",
                message="Current simulation date is invalid.",
                details={"current_date": current.get("current_date")},
            )
        next_date = (date.fromisoformat(baseline_date) + timedelta(days=1)).isoformat()
        with transaction(self.connection):
            snapshot_created = self._ensure_simulation_restore_snapshot(
                baseline_date=baseline_date
            )
            seeded_capacity_count = self._seed_simulated_morning_capacity(
                baseline_date
            )
            reporting_stats = self._simulate_same_day_reportings(baseline_date)
            rebuilt_actual_row_count, skipped_rebuild_report_count = (
                self._rebuild_line_daily_actual_capacity_rows(baseline_date)
            )
            self.connection.execute(
                """
                INSERT INTO simulation_state (
                    singleton_key,
                    current_date,
                    updated_at
                ) VALUES (?, ?, ?)
                ON CONFLICT(singleton_key) DO UPDATE SET
                    current_date = excluded.current_date,
                    updated_at = excluded.updated_at
                """,
                (RULES_SINGLETON_KEY, next_date, utc_now()),
            )
        return {
            "current_date": next_date,
            "message": f"Simulation advanced to {next_date}.",
            "simulated_date": baseline_date,
            "snapshot_created": snapshot_created,
            "seeded_morning_capacity_count": seeded_capacity_count,
            "simulated_reporting_count": reporting_stats["inserted_report_count"],
            "skipped_existing_reporting_count": reporting_stats[
                "skipped_existing_report_count"
            ],
            "skipped_zero_capacity_reporting_count": reporting_stats[
                "skipped_zero_capacity_report_count"
            ],
            "rebuild_actual_row_count": rebuilt_actual_row_count,
            "rebuild_skipped_report_count": skipped_rebuild_report_count,
        }

    def reset_manual_simulation(self) -> dict[str, Any]:
        snapshot = fetch_one(
            self.connection,
            """
            SELECT snapshot_current_date
            FROM simulation_restore_snapshot_meta
            WHERE singleton_key = ?
            """,
            (RULES_SINGLETON_KEY,),
        )
        if snapshot is None:
            raise bad_request(
                code="SIMULATION_SNAPSHOT_NOT_FOUND",
                message=(
                    "No simulation snapshot is available to restore. "
                    "Advance simulation first."
                ),
            )
        restored_date = _normalize_date_text(snapshot.get("snapshot_current_date"))
        if restored_date is None:
            raise server_error(
                code="SIMULATION_SNAPSHOT_DATE_INVALID",
                message="Snapshot simulation date is invalid.",
                details={"snapshot_current_date": snapshot.get("snapshot_current_date")},
            )
        restored_plan_count = int(
            _to_number(
                fetch_one(
                    self.connection,
                    """
                    SELECT COUNT(1) AS total
                    FROM simulation_restore_snapshot_daily_line_capacity_plan
                    """,
                )["total"],
                0,
            )
        )
        restored_actual_count = int(
            _to_number(
                fetch_one(
                    self.connection,
                    """
                    SELECT COUNT(1) AS total
                    FROM simulation_restore_snapshot_daily_line_capacity_actual
                    """,
                )["total"],
                0,
            )
        )
        restored_report_count = int(
            _to_number(
                fetch_one(
                    self.connection,
                    """
                    SELECT COUNT(1) AS total
                    FROM simulation_restore_snapshot_work_reports
                    """,
                )["total"],
                0,
            )
        )
        restored_plan_audit_count = int(
            _to_number(
                fetch_one(
                    self.connection,
                    """
                    SELECT COUNT(1) AS total
                    FROM simulation_restore_snapshot_daily_line_capacity_plan_audit
                    """,
                )["total"],
                0,
            )
        )
        with transaction(self.connection):
            self.connection.execute("DELETE FROM daily_line_capacity_plan_audit")
            self.connection.execute("DELETE FROM daily_line_capacity_actual")
            self.connection.execute("DELETE FROM daily_line_capacity_plan")
            self.connection.execute("DELETE FROM work_reports")
            self.connection.execute(
                """
                INSERT INTO daily_line_capacity_plan_audit (
                    audit_id,
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    old_planned_capacity_qty,
                    new_planned_capacity_qty,
                    old_worker_count,
                    new_worker_count,
                    old_machine_count,
                    new_machine_count,
                    operator_user_id,
                    operator_username,
                    operator_display_name,
                    changed_at
                )
                SELECT
                    audit_id,
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    old_planned_capacity_qty,
                    new_planned_capacity_qty,
                    old_worker_count,
                    new_worker_count,
                    old_machine_count,
                    new_machine_count,
                    operator_user_id,
                    operator_username,
                    operator_display_name,
                    changed_at
                FROM simulation_restore_snapshot_daily_line_capacity_plan_audit
                """
            )
            self.connection.execute(
                """
                INSERT INTO daily_line_capacity_plan (
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    planned_capacity_qty,
                    worker_count,
                    machine_count,
                    source_note,
                    updated_at
                )
                SELECT
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    planned_capacity_qty,
                    worker_count,
                    machine_count,
                    source_note,
                    updated_at
                FROM simulation_restore_snapshot_daily_line_capacity_plan
                """
            )
            self.connection.execute(
                """
                INSERT INTO daily_line_capacity_actual (
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    actual_capacity_qty,
                    report_count,
                    last_report_time,
                    updated_at
                )
                SELECT
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    actual_capacity_qty,
                    report_count,
                    last_report_time,
                    updated_at
                FROM simulation_restore_snapshot_daily_line_capacity_actual
                """
            )
            self.connection.execute(
                """
                INSERT INTO work_reports (
                    report_id,
                    production_order_no,
                    process_code,
                    process_name,
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    daily_capacity_compare_audit_id,
                    daily_capacity_compare_qty,
                    daily_capacity_compare_selected_at,
                    updated_at
                )
                SELECT
                    report_id,
                    production_order_no,
                    process_code,
                    process_name,
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    daily_capacity_compare_audit_id,
                    daily_capacity_compare_qty,
                    daily_capacity_compare_selected_at,
                    updated_at
                FROM simulation_restore_snapshot_work_reports
                """
            )
            self.connection.execute(
                """
                INSERT INTO simulation_state (
                    singleton_key,
                    current_date,
                    updated_at
                ) VALUES (?, ?, ?)
                ON CONFLICT(singleton_key) DO UPDATE SET
                    current_date = excluded.current_date,
                    updated_at = excluded.updated_at
                """,
                (RULES_SINGLETON_KEY, restored_date, utc_now()),
            )
            self._clear_simulation_restore_snapshot()
        return {
            "current_date": restored_date,
            "message": "Simulation reset.",
            "restored_plan_count": restored_plan_count,
            "restored_actual_count": restored_actual_count,
            "restored_report_count": restored_report_count,
            "restored_plan_audit_count": restored_plan_audit_count,
        }

    def _ensure_simulation_restore_snapshot(self, *, baseline_date: str) -> bool:
        existing = fetch_one(
            self.connection,
            """
            SELECT singleton_key
            FROM simulation_restore_snapshot_meta
            WHERE singleton_key = ?
            """,
            (RULES_SINGLETON_KEY,),
        )
        if existing is not None:
            return False
        normalized_date = _normalize_date_text(baseline_date)
        if normalized_date is None:
            raise bad_request(
                code="SIMULATION_BASELINE_DATE_INVALID",
                message="baseline_date must be a valid YYYY-MM-DD date.",
                details={"baseline_date": baseline_date},
            )
        self._clear_simulation_restore_snapshot()
        self.connection.execute(
            """
            INSERT INTO simulation_restore_snapshot_meta (
                singleton_key,
                snapshot_current_date,
                snapshot_created_at
            ) VALUES (?, ?, ?)
            """,
            (RULES_SINGLETON_KEY, normalized_date, utc_now()),
        )
        self.connection.execute(
            """
            INSERT INTO simulation_restore_snapshot_daily_line_capacity_plan (
                calendar_date,
                company_code,
                workshop_code,
                line_code,
                process_code,
                planned_capacity_qty,
                worker_count,
                machine_count,
                source_note,
                updated_at
            )
            SELECT
                calendar_date,
                company_code,
                workshop_code,
                line_code,
                process_code,
                planned_capacity_qty,
                worker_count,
                machine_count,
                source_note,
                updated_at
            FROM daily_line_capacity_plan
            """
        )
        self.connection.execute(
            """
            INSERT INTO simulation_restore_snapshot_daily_line_capacity_actual (
                calendar_date,
                company_code,
                workshop_code,
                line_code,
                process_code,
                actual_capacity_qty,
                report_count,
                last_report_time,
                updated_at
            )
            SELECT
                calendar_date,
                company_code,
                workshop_code,
                line_code,
                process_code,
                actual_capacity_qty,
                report_count,
                last_report_time,
                updated_at
            FROM daily_line_capacity_actual
            """
        )
        self.connection.execute(
            """
            INSERT INTO simulation_restore_snapshot_daily_line_capacity_plan_audit (
                audit_id,
                calendar_date,
                company_code,
                workshop_code,
                line_code,
                process_code,
                old_planned_capacity_qty,
                new_planned_capacity_qty,
                old_worker_count,
                new_worker_count,
                old_machine_count,
                new_machine_count,
                operator_user_id,
                operator_username,
                operator_display_name,
                changed_at
            )
            SELECT
                audit_id,
                calendar_date,
                company_code,
                workshop_code,
                line_code,
                process_code,
                old_planned_capacity_qty,
                new_planned_capacity_qty,
                old_worker_count,
                new_worker_count,
                old_machine_count,
                new_machine_count,
                operator_user_id,
                operator_username,
                operator_display_name,
                changed_at
            FROM daily_line_capacity_plan_audit
            """
        )
        self.connection.execute(
            """
            INSERT INTO simulation_restore_snapshot_work_reports (
                report_id,
                production_order_no,
                process_code,
                process_name,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                report_qty,
                report_time,
                operator_name,
                daily_capacity_compare_audit_id,
                daily_capacity_compare_qty,
                daily_capacity_compare_selected_at,
                updated_at
            )
            SELECT
                report_id,
                production_order_no,
                process_code,
                process_name,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                report_qty,
                report_time,
                operator_name,
                daily_capacity_compare_audit_id,
                daily_capacity_compare_qty,
                daily_capacity_compare_selected_at,
                updated_at
            FROM work_reports
            """
        )
        return True

    def _clear_simulation_restore_snapshot(self) -> None:
        self.connection.execute("DELETE FROM simulation_restore_snapshot_daily_line_capacity_plan")
        self.connection.execute("DELETE FROM simulation_restore_snapshot_daily_line_capacity_actual")
        self.connection.execute("DELETE FROM simulation_restore_snapshot_daily_line_capacity_plan_audit")
        self.connection.execute("DELETE FROM simulation_restore_snapshot_work_reports")
        self.connection.execute(
            "DELETE FROM simulation_restore_snapshot_meta WHERE singleton_key = ?",
            (RULES_SINGLETON_KEY,),
        )

    def _seed_simulated_morning_capacity(self, calendar_date: str) -> int:
        normalized_date = _normalize_date_text(calendar_date)
        if normalized_date is None:
            raise bad_request(
                code="CALENDAR_DATE_REQUIRED",
                message="calendar_date must be a valid YYYY-MM-DD date.",
            )
        capacity_data = self.list_line_daily_capacity(normalized_date)
        capacity_items = capacity_data.get("items") if isinstance(capacity_data, dict) else []
        if not isinstance(capacity_items, list) or len(capacity_items) == 0:
            raise server_error(
                code="SIMULATION_CAPACITY_TOPOLOGY_EMPTY",
                message="No line topology rows are available for simulation capacity seeding.",
                details={"calendar_date": normalized_date},
            )
        existing_rows = fetch_all(
            self.connection,
            """
            SELECT company_code, workshop_code, line_code, process_code
            FROM daily_line_capacity_plan
            WHERE calendar_date = ?
            """,
            (normalized_date,),
        )
        existing_keys = {
            (
                str(row.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper()
                or DEFAULT_COMPANY_CODE,
                str(row.get("workshop_code") or "").strip().upper(),
                str(row.get("line_code") or "").strip().upper(),
                str(row.get("process_code") or "").strip().upper(),
            )
            for row in existing_rows
        }
        rows_to_insert: list[tuple[Any, ...]] = []
        updated_at = utc_now()
        for item in capacity_items:
            company_code = (
                str(item.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper()
                or DEFAULT_COMPANY_CODE
            )
            workshop_code = str(item.get("workshop_code") or "").strip().upper()
            line_code = str(item.get("line_code") or "").strip().upper()
            process_code = str(item.get("process_code") or "").strip().upper()
            if not workshop_code or not line_code or not process_code:
                raise server_error(
                    code="SIMULATION_CAPACITY_KEY_INVALID",
                    message="Simulation capacity row key is invalid.",
                    details={
                        "calendar_date": normalized_date,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                        "process_code": process_code,
                    },
                )
            key = (company_code, workshop_code, line_code, process_code)
            if key in existing_keys:
                continue
            rows_to_insert.append(
                (
                    normalized_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    _to_number(item.get("planned_capacity_qty"), 0),
                    int(_to_number(item.get("worker_count"), 0)),
                    int(_to_number(item.get("machine_count"), 0)),
                    "SIMULATION_AUTO_MORNING_ESTIMATE",
                    updated_at,
                )
            )
        if rows_to_insert:
            self.connection.executemany(
                """
                INSERT INTO daily_line_capacity_plan (
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    planned_capacity_qty,
                    worker_count,
                    machine_count,
                    source_note,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                rows_to_insert,
            )
        return len(rows_to_insert)

    def _simulate_same_day_reportings(self, calendar_date: str) -> dict[str, int]:
        normalized_date = _normalize_date_text(calendar_date)
        if normalized_date is None:
            raise bad_request(
                code="CALENDAR_DATE_REQUIRED",
                message="calendar_date must be a valid YYYY-MM-DD date.",
            )
        plan_rows = fetch_all(
            self.connection,
            """
            SELECT
                plan.company_code,
                plan.workshop_code,
                COALESCE(top.workshop_name, plan.workshop_code) AS workshop_name,
                plan.line_code,
                COALESCE(top.line_name, plan.line_code) AS line_name,
                plan.process_code,
                plan.planned_capacity_qty
            FROM daily_line_capacity_plan plan
            LEFT JOIN masterdata_line_topology top
              ON top.company_code = plan.company_code
             AND top.workshop_code = plan.workshop_code
             AND top.line_code = plan.line_code
             AND top.process_code = plan.process_code
            WHERE plan.calendar_date = ?
            ORDER BY
                plan.workshop_code ASC,
                plan.line_code ASC,
                plan.process_code ASC
            """,
            (normalized_date,),
        )
        if len(plan_rows) == 0:
            raise server_error(
                code="SIMULATION_CAPACITY_PLAN_EMPTY",
                message=(
                    "No daily line capacity plan rows exist for the simulation date. "
                    "Simulation reporting cannot be generated."
                ),
                details={"calendar_date": normalized_date},
            )
        existing_reporting_rows = fetch_all(
            self.connection,
            """
            SELECT DISTINCT
                UPPER(TRIM(COALESCE(workshop_code, ''))) AS workshop_code,
                UPPER(TRIM(COALESCE(line_code, ''))) AS line_code,
                UPPER(TRIM(COALESCE(process_code, ''))) AS process_code
            FROM work_reports
            WHERE date(report_time, '+8 hours') = ?
            """,
            (normalized_date,),
        )
        existing_report_keys = {
            (
                str(row.get("workshop_code") or "").strip().upper(),
                str(row.get("line_code") or "").strip().upper(),
                str(row.get("process_code") or "").strip().upper(),
            )
            for row in existing_reporting_rows
        }
        process_name_by_code = self._process_name_by_code()
        report_base_time = (
            datetime.fromisoformat(f"{normalized_date}T10:00:00+08:00")
            .astimezone(timezone.utc)
            .replace(microsecond=0)
        )
        rows_to_insert: list[tuple[Any, ...]] = []
        skipped_existing_report_count = 0
        skipped_zero_capacity_report_count = 0
        for index, row in enumerate(plan_rows):
            company_code = (
                str(row.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper()
                or DEFAULT_COMPANY_CODE
            )
            workshop_code = str(row.get("workshop_code") or "").strip().upper()
            line_code = str(row.get("line_code") or "").strip().upper()
            process_code = str(row.get("process_code") or "").strip().upper()
            if not workshop_code or not line_code or not process_code:
                raise server_error(
                    code="SIMULATION_REPORTING_KEY_INVALID",
                    message="Simulation reporting row key is invalid.",
                    details={
                        "calendar_date": normalized_date,
                        "workshop_code": workshop_code,
                        "line_code": line_code,
                        "process_code": process_code,
                    },
                )
            if (workshop_code, line_code, process_code) in existing_report_keys:
                skipped_existing_report_count += 1
                continue
            planned_capacity_qty = _to_number(row.get("planned_capacity_qty"), 0)
            if planned_capacity_qty <= SCHEDULE_NUMBER_EPSILON:
                skipped_zero_capacity_report_count += 1
                continue
            ratio = self._simulation_reporting_ratio(
                calendar_date=normalized_date,
                company_code=company_code,
                workshop_code=workshop_code,
                line_code=line_code,
                process_code=process_code,
            )
            report_qty = round(planned_capacity_qty * ratio)
            if report_qty <= SCHEDULE_NUMBER_EPSILON:
                skipped_zero_capacity_report_count += 1
                continue
            report_time = (report_base_time + timedelta(minutes=index)).isoformat()
            rows_to_insert.append(
                (
                    f"RPT-SIM-{uuid4().hex[:10].upper()}",
                    None,
                    process_code,
                    process_name_by_code.get(process_code, process_code),
                    workshop_code,
                    str(row.get("workshop_name") or workshop_code).strip() or workshop_code,
                    line_code,
                    str(row.get("line_name") or line_code).strip() or line_code,
                    float(report_qty),
                    report_time,
                    "simulation-auto",
                    report_time,
                )
            )
        if rows_to_insert:
            self.connection.executemany(
                """
                INSERT INTO work_reports (
                    report_id,
                    production_order_no,
                    process_code,
                    process_name,
                    workshop_code,
                    workshop_name,
                    line_code,
                    line_name,
                    report_qty,
                    report_time,
                    operator_name,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                rows_to_insert,
            )
        return {
            "inserted_report_count": len(rows_to_insert),
            "skipped_existing_report_count": skipped_existing_report_count,
            "skipped_zero_capacity_report_count": skipped_zero_capacity_report_count,
        }

    def _simulation_reporting_ratio(
        self,
        *,
        calendar_date: str,
        company_code: str,
        workshop_code: str,
        line_code: str,
        process_code: str,
    ) -> float:
        seed_text = (
            f"{calendar_date}|{company_code}|{workshop_code}|{line_code}|{process_code}"
        )
        weighted_sum = sum((index + 1) * ord(ch) for index, ch in enumerate(seed_text))
        bucket = weighted_sum % 31
        return 0.65 + (bucket / 100.0)

    def test_material_issues(self, order_no: str, mode: str) -> dict[str, Any]:
        normalized_mode = str(mode or "fast").strip().lower()
        items = self.factory.material_gateway.fetch_order_materials(
            order_no,
            mode=normalized_mode,
        )
        return {
            "items": [
                {
                    "order_no": item.get("production_order_no") or order_no,
                    "child_material_code": item.get("child_material_code"),
                    "child_material_name_cn": item.get("child_material_name"),
                    "spec_model": item.get("spec_model"),
                }
                for item in items
            ]
        }

    def test_material_supply(self, material_code: str) -> dict[str, Any]:
        code = str(material_code or "").strip()
        if not code:
            raise bad_request(
                code="MATERIAL_CODE_REQUIRED",
                message="material_code is required.",
            )
        rows = self.supply_gateway.fetch_material_supply([code])
        if len(rows) == 0:
            raise not_found(
                code="MATERIAL_SUPPLY_NOT_FOUND",
                message="Material supply does not exist in ERP.",
                details={"material_code": code},
            )
        updated_at = utc_now()
        with transaction(self.connection):
            self.factory.supply_repository.upsert_many(
                [
                    {
                        "material_code": str(row["material_code"]),
                        "supply_type_code": row["supply_type_code"],
                        "supply_type_name": row["supply_type_name"],
                        "snapshot_time": updated_at,
                        "updated_at": updated_at,
                    }
                    for row in rows
                ]
            )
        row = rows[0]
        return {
            "material_code": row["material_code"],
            "supply_type": row["supply_type_code"],
            "supply_type_name_cn": row["supply_type_name"],
            "is_self_made": str(row["supply_type_code"]).strip().upper() == "SELF_MADE",
        }

    def test_material_inventory(self, material_code: str) -> dict[str, Any]:
        code = str(material_code or "").strip()
        if not code:
            raise bad_request(
                code="MATERIAL_CODE_REQUIRED",
                message="material_code is required.",
            )
        snapshots = self.inventory_gateway.fetch_inventory([code])
        snapshot = snapshots[0] if snapshots else {
            "material_code": code,
            "inventory_qty": None,
            "inventory_status": "UNKNOWN",
        }
        updated_at = utc_now()
        with transaction(self.connection):
            self.factory.inventory_repository.upsert_many(
                [
                    {
                        "material_code": code,
                        "inventory_qty": snapshot.get("inventory_qty"),
                        "inventory_status": snapshot["inventory_status"],
                        "snapshot_time": updated_at,
                        "updated_at": updated_at,
                    }
                ]
            )
        reference = self._find_material_reference(code)
        inventory_qty = snapshot.get("inventory_qty")
        items = []
        if inventory_qty is not None:
            items.append(
                {
                    "material_code": code,
                    "material_name_cn": reference.get("material_name_cn") or code,
                    "spec_model": reference.get("spec_model"),
                    "warehouse_name": None,
                    "batch_no": None,
                    "base_unit_name": reference.get("child_unit"),
                    "stock_qty": inventory_qty,
                    "stock_org_name": None,
                }
            )
        return {
            "material_code": code,
            "items": items,
            "total_stock_qty": _to_number(inventory_qty, 0),
        }

    def import_production_orders(self, payload: dict[str, Any]) -> dict[str, Any]:
        random_guidewire = payload.get("random_guidewire_protection_order") is True
        count = 1 if random_guidewire else max(1, int(_to_number(payload.get("n"), 1)))
        material_code = (
            "YXN.044.02.1020"
            if random_guidewire
            else str(payload.get("material_code") or "").strip().upper()
        )
        if not material_code:
            raise bad_request(
                code="MATERIAL_CODE_REQUIRED",
                message="material_code is required.",
            )
        completed = payload.get("completed") is True
        product_name = "导丝保护组件" if random_guidewire else material_code
        updated_at = utc_now()
        next_number = self._next_import_order_sequence()
        order_nos: list[str] = []
        order_qty_list: list[int] = []
        with transaction(self.connection):
            for _ in range(count):
                order_no = f"{next_number:04d}MO-TEST"
                next_number += 1
                order_qty = random.randint(1000, 5000) if random_guidewire else 100
                start_date = _today_text()
                end_date = (date.today() + timedelta(days=3)).isoformat()
                source_bill_no = f"TEST-{order_no}"
                self.connection.execute(
                    """
                    INSERT INTO production_orders (
                        production_order_no,
                        material_code,
                        material_name,
                        material_specification,
                        production_qty,
                        status,
                        planned_start_date,
                        planned_end_date,
                        source_bill_no,
                        material_list_no,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(production_order_no) DO UPDATE SET
                        material_code = excluded.material_code,
                        material_name = excluded.material_name,
                        material_specification = excluded.material_specification,
                        production_qty = excluded.production_qty,
                        status = excluded.status,
                        planned_start_date = excluded.planned_start_date,
                        planned_end_date = excluded.planned_end_date,
                        source_bill_no = excluded.source_bill_no,
                        material_list_no = excluded.material_list_no,
                        updated_at = excluded.updated_at
                    """,
                    (
                        order_no,
                        material_code,
                        product_name,
                        None,
                        order_qty,
                        "DONE" if completed else "OPEN",
                        start_date,
                        end_date,
                        source_bill_no,
                        f"ML-{material_code}",
                        updated_at,
                    ),
                )
                self._upsert_order_state(
                    {
                        "production_order_no": order_no,
                        "promised_due_date": end_date,
                        "expected_start_date": start_date,
                        "expected_start_time": _iso_at(start_date, "08:00:00"),
                        "expected_finish_time": _iso_at(end_date, "18:00:00"),
                        "priority_level": PRIORITY_LEVEL_MAX,
                        "urgent_flag": 0,
                        "lock_flag": 0,
                        "frozen_flag": 0,
                        "status": "DONE" if completed else "OPEN",
                        "order_status": "DONE" if completed else "OPEN",
                        "completed_qty": order_qty if completed else 0,
                        "remaining_qty": 0 if completed else order_qty,
                        "progress_rate": 100 if completed else 0,
                        "production_batch_no": f"{source_bill_no}-B1",
                    }
                )
                order_nos.append(order_no)
                order_qty_list.append(order_qty)
        return {
            "imported_count": count,
            "material_code": material_code,
            "product_name_cn": product_name,
            "completed": completed,
            "order_nos": order_nos,
            "order_qty_list": order_qty_list,
        }

    def import_production_orders_from_erp(self, payload: dict[str, Any]) -> dict[str, Any]:
        if payload.get("random_guidewire_protection_order") is True:
            return self.import_production_orders(payload)
        material_code = str(payload.get("material_code") or "").strip().upper()
        keyword = str(payload.get("keyword") or "").strip()
        if not material_code and not keyword:
            raise bad_request(
                code="IMPORT_FILTER_REQUIRED",
                message="material_code or keyword is required.",
            )
        requested_count = max(1, min(200, int(_to_number(payload.get("n"), 1))))
        completed = payload.get("completed") is True
        erp_orders = self.order_gateway.fetch_orders(
            material_code=material_code or None,
            keyword=keyword or None,
            limit=requested_count,
        )
        if len(erp_orders) == 0:
            raise bad_request(
                code="ERP_PRODUCTION_ORDERS_EMPTY",
                message="ERP production order list is empty.",
            )

        normalized_keyword = keyword.lower()
        matched_orders: list[dict[str, Any]] = []
        for row in erp_orders:
            row_material_code = str(row.get("material_code") or "").strip().upper()
            row_material_name = str(row.get("material_name") or "").strip()
            if material_code and row_material_code != material_code:
                continue
            if normalized_keyword and normalized_keyword not in row_material_name.lower():
                continue
            matched_orders.append(row)

        selected_orders = matched_orders[:requested_count]
        updated_at = utc_now()
        imported_order_nos: list[str] = []
        order_qty_list: list[float] = []
        with transaction(self.connection):
            for row in selected_orders:
                order_no = str(row.get("production_order_no") or "").strip()
                if not order_no:
                    continue
                row_material_code = str(row.get("material_code") or "").strip().upper()
                row_material_name = str(row.get("material_name") or row_material_code).strip() or row_material_code
                order_qty = _to_number(row.get("production_qty"), 0)
                if order_qty <= 0:
                    continue
                start_date = _normalize_date_text(row.get("planned_start_date")) or _today_text()
                end_date = _normalize_date_text(row.get("planned_end_date")) or start_date
                source_bill_no = str(row.get("source_bill_no") or "").strip() or None
                material_list_no = str(row.get("material_list_no") or "").strip() or None
                normalized_status = "DONE" if completed else "OPEN"
                self.connection.execute(
                    """
                    INSERT INTO production_orders (
                        production_order_no,
                        material_code,
                        material_name,
                        material_specification,
                        production_qty,
                        status,
                        planned_start_date,
                        planned_end_date,
                        source_bill_no,
                        material_list_no,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(production_order_no) DO UPDATE SET
                        material_code = excluded.material_code,
                        material_name = excluded.material_name,
                        material_specification = excluded.material_specification,
                        production_qty = excluded.production_qty,
                        status = excluded.status,
                        planned_start_date = excluded.planned_start_date,
                        planned_end_date = excluded.planned_end_date,
                        source_bill_no = excluded.source_bill_no,
                        material_list_no = excluded.material_list_no,
                        updated_at = excluded.updated_at
                    """,
                    (
                        order_no,
                        row_material_code,
                        row_material_name,
                        row.get("material_specification"),
                        order_qty,
                        normalized_status,
                        start_date,
                        end_date,
                        source_bill_no,
                        material_list_no,
                        updated_at,
                    ),
                )
                self._upsert_order_state(
                    {
                        "production_order_no": order_no,
                        "promised_due_date": end_date,
                        "expected_start_date": start_date,
                        "expected_start_time": _iso_at(start_date, "08:00:00"),
                        "expected_finish_time": _iso_at(end_date, "18:00:00"),
                        "priority_level": PRIORITY_LEVEL_MAX,
                        "urgent_flag": 0,
                        "lock_flag": 0,
                        "frozen_flag": 0,
                        "status": normalized_status,
                        "order_status": normalized_status,
                        "completed_qty": order_qty if completed else 0,
                        "remaining_qty": 0 if completed else order_qty,
                        "progress_rate": 100 if completed else 0,
                        "production_batch_no": f"{source_bill_no}-B1" if source_bill_no else None,
                    }
                )
                imported_order_nos.append(order_no)
                order_qty_list.append(order_qty)
        return {
            "material_code": material_code,
            "keyword": keyword,
            "requested_count": requested_count,
            "matched_count": len(matched_orders),
            "imported_count": len(imported_order_nos),
            "completed": completed,
            "imported_order_nos": imported_order_nos,
            "order_nos": imported_order_nos,
            "order_qty_list": order_qty_list,
        }

    def _require_order(self, order_no: str) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT
                production_order_no,
                material_code,
                material_name,
                material_specification,
                production_qty,
                status,
                planned_start_date,
                planned_end_date,
                source_bill_no,
                material_list_no,
                updated_at
            FROM production_orders
            WHERE production_order_no = ?
            """,
            (order_no,),
        )
        if row is None:
            raise not_found(
                code="ORDER_NOT_FOUND",
                message="Order does not exist.",
                details={"order_no": order_no},
            )
        return row

    def _get_order_rows_by_nos(self, order_nos: list[str]) -> dict[str, dict[str, Any]]:
        if len(order_nos) == 0:
            return {}
        placeholders = ",".join("?" for _ in order_nos)
        rows = fetch_all(
            self.connection,
            f"""
            SELECT
                production_order_no,
                material_code,
                material_name,
                material_specification,
                production_qty,
                status,
                planned_start_date,
                planned_end_date,
                source_bill_no,
                material_list_no,
                updated_at
            FROM production_orders
            WHERE production_order_no IN ({placeholders})
            """,
            tuple(order_nos),
        )
        return {str(row["production_order_no"]): row for row in rows}

    def _list_order_rows(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                production_order_no,
                material_code,
                material_name,
                material_specification,
                production_qty,
                status,
                planned_start_date,
                planned_end_date,
                source_bill_no,
                material_list_no,
                updated_at
            FROM production_orders
            ORDER BY updated_at DESC, production_order_no DESC
            """,
        )

    def _get_order_state(self, order_no: str) -> dict[str, Any] | None:
        return fetch_one(
            self.connection,
            """
            SELECT
                production_order_no,
                promised_due_date,
                expected_start_date,
                expected_start_time,
                expected_finish_time,
                priority_level,
                urgent_flag,
                lock_flag,
                frozen_flag,
                status,
                order_status,
                completed_qty,
                remaining_qty,
                progress_rate,
                production_batch_no,
                updated_at
            FROM order_pool_state
            WHERE production_order_no = ?
            """,
            (order_no,),
        )

    def _get_order_state_map(self, order_nos: list[str]) -> dict[str, dict[str, Any]]:
        if len(order_nos) == 0:
            return {}
        placeholders = ",".join("?" for _ in order_nos)
        rows = fetch_all(
            self.connection,
            f"""
            SELECT
                production_order_no,
                promised_due_date,
                expected_start_date,
                expected_start_time,
                expected_finish_time,
                priority_level,
                urgent_flag,
                lock_flag,
                frozen_flag,
                status,
                order_status,
                completed_qty,
                remaining_qty,
                progress_rate,
                production_batch_no,
                updated_at
            FROM order_pool_state
            WHERE production_order_no IN ({placeholders})
            """,
            tuple(order_nos),
        )
        return {str(row["production_order_no"]): row for row in rows}

    def _upsert_order_state(self, row: dict[str, Any]) -> None:
        updated_at = utc_now()
        self.connection.execute(
            """
            INSERT INTO order_pool_state (
                production_order_no,
                promised_due_date,
                expected_start_date,
                expected_start_time,
                expected_finish_time,
                priority_level,
                urgent_flag,
                lock_flag,
                frozen_flag,
                status,
                order_status,
                completed_qty,
                remaining_qty,
                progress_rate,
                production_batch_no,
                updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(production_order_no) DO UPDATE SET
                promised_due_date = excluded.promised_due_date,
                expected_start_date = excluded.expected_start_date,
                expected_start_time = excluded.expected_start_time,
                expected_finish_time = excluded.expected_finish_time,
                priority_level = excluded.priority_level,
                urgent_flag = excluded.urgent_flag,
                lock_flag = excluded.lock_flag,
                frozen_flag = excluded.frozen_flag,
                status = excluded.status,
                order_status = excluded.order_status,
                completed_qty = excluded.completed_qty,
                remaining_qty = excluded.remaining_qty,
                progress_rate = excluded.progress_rate,
                production_batch_no = excluded.production_batch_no,
                updated_at = excluded.updated_at
            """,
            (
                row["production_order_no"],
                row.get("promised_due_date"),
                row.get("expected_start_date"),
                row.get("expected_start_time"),
                row.get("expected_finish_time"),
                _normalize_priority_level(row.get("priority_level"), PRIORITY_LEVEL_MAX),
                _urgent_flag_from_priority_level(row.get("priority_level")),
                int(row.get("lock_flag") or 0),
                int(row.get("frozen_flag") or 0),
                row.get("status"),
                row.get("order_status"),
                row.get("completed_qty"),
                row.get("remaining_qty"),
                row.get("progress_rate"),
                row.get("production_batch_no"),
                updated_at,
            ),
        )

    def _sync_order_pool_state_schedule_window_from_tasks(
        self,
        *,
        tasks: list[tuple[Any, ...]],
        order_rows: list[dict[str, Any]],
        states: dict[str, dict[str, Any]],
    ) -> None:
        windows_by_order_no: dict[str, dict[str, str]] = {}
        for task in tasks:
            if len(task) < 9:
                continue
            order_no = str(task[2] or "").strip()
            calendar_date = _normalize_date_text(task[7])
            if not order_no or not calendar_date:
                continue
            finish_date = date.fromisoformat(calendar_date)
            if _normalize_shift_code(task[8]) == "NIGHT":
                finish_date = finish_date + timedelta(days=1)
            finish_date_text = finish_date.isoformat()
            current = windows_by_order_no.get(order_no)
            if current is None:
                windows_by_order_no[order_no] = {
                    "start_date": calendar_date,
                    "finish_date": finish_date_text,
                }
                continue
            if calendar_date < current["start_date"]:
                current["start_date"] = calendar_date
            if finish_date_text > current["finish_date"]:
                current["finish_date"] = finish_date_text

        if len(windows_by_order_no) == 0:
            return

        order_row_by_no = {
            str(row.get("production_order_no") or "").strip(): row
            for row in order_rows
        }
        for order_no, window in windows_by_order_no.items():
            state_row = states.get(order_no) or {}
            base_row = order_row_by_no.get(order_no) or {}
            status = str(
                state_row.get("status")
                or state_row.get("order_status")
                or base_row.get("status")
                or "OPEN"
            ).strip().upper() or "OPEN"
            order_status = str(
                state_row.get("order_status")
                or state_row.get("status")
                or status
            ).strip().upper() or status
            production_qty = _to_number(base_row.get("production_qty"), 0)
            completed_qty = (
                _to_number(state_row.get("completed_qty"), 0)
                if state_row.get("completed_qty") is not None
                else (production_qty if order_status == "DONE" else 0)
            )
            remaining_qty = (
                _to_number(state_row.get("remaining_qty"), 0)
                if state_row.get("remaining_qty") is not None
                else max(0.0, production_qty - completed_qty)
            )
            progress_rate = (
                _to_number(state_row.get("progress_rate"), 0)
                if state_row.get("progress_rate") is not None
                else ((completed_qty / production_qty * 100) if production_qty > 0 else 0)
            )
            source_bill_no = str(base_row.get("source_bill_no") or "").strip()
            production_batch_no = state_row.get("production_batch_no")
            if production_batch_no is None and source_bill_no:
                production_batch_no = f"{source_bill_no}-B1"

            next_row = {
                "production_order_no": order_no,
                "promised_due_date": state_row.get("promised_due_date") or base_row.get("planned_end_date"),
                "expected_start_date": window["start_date"],
                "expected_start_time": _iso_at(window["start_date"], "08:00:00"),
                "expected_finish_time": _iso_at(window["finish_date"], "18:00:00"),
                "priority_level": _normalize_priority_level(
                    state_row.get("priority_level"),
                    PRIORITY_LEVEL_MAX,
                ),
                "urgent_flag": _urgent_flag_from_priority_level(
                    state_row.get("priority_level")
                ),
                "lock_flag": int(_to_number(state_row.get("lock_flag"), 0)),
                "frozen_flag": int(_to_number(state_row.get("frozen_flag"), 0)),
                "status": status,
                "order_status": order_status,
                "completed_qty": completed_qty,
                "remaining_qty": remaining_qty,
                "progress_rate": progress_rate,
                "production_batch_no": production_batch_no,
            }
            self._upsert_order_state(next_row)

    def _get_capacity_rows(self, order_no: str) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                process_code,
                process_name,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                capacity_qty
            FROM capacity_bindings
            WHERE production_order_no = ?
            ORDER BY process_code ASC, workshop_code ASC, line_code ASC
            """,
            (order_no,),
        )

    def _get_capacity_map(self, order_nos: list[str]) -> dict[str, list[dict[str, Any]]]:
        if len(order_nos) == 0:
            return {}
        placeholders = ",".join("?" for _ in order_nos)
        rows = fetch_all(
            self.connection,
            f"""
            SELECT
                production_order_no,
                process_code,
                process_name,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                capacity_qty
            FROM capacity_bindings
            WHERE production_order_no IN ({placeholders})
            ORDER BY production_order_no ASC, process_code ASC, workshop_code ASC, line_code ASC
            """,
            tuple(order_nos),
        )
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            grouped[str(row["production_order_no"])].append(row)
        return grouped

    def _build_order_pool_row(
        self,
        *,
        base_row: dict[str, Any],
        state_row: dict[str, Any] | None,
        capacity_rows: list[dict[str, Any]],
        route_rows: list[dict[str, Any]],
        topology_by_process: dict[str, list[dict[str, Any]]],
        reference_version: dict[str, Any] | None,
        schedule_fact: dict[str, Any] | None,
        published_version: dict[str, Any] | None,
        published_schedule_fact: dict[str, Any] | None,
        shortage_summary: dict[str, Any] | None,
        final_process_metrics: dict[str, Any] | None,
    ) -> dict[str, Any]:
        production_qty = _to_number(base_row.get("production_qty"), 0)
        completed_qty = (
            _to_number((state_row or {}).get("completed_qty"), 0)
            if (state_row or {}).get("completed_qty") is not None
            else 0
        )
        remaining_qty = (state_row or {}).get("remaining_qty")
        if remaining_qty is None:
            remaining_qty = max(0, production_qty - completed_qty)
        progress_rate = (state_row or {}).get("progress_rate")
        if progress_rate is None:
            progress_rate = (completed_qty / production_qty * 100) if production_qty > 0 else 0
        promised_due_date = _normalize_date_text(
            (state_row or {}).get("promised_due_date") or base_row.get("planned_end_date")
        )
        base_expected_start_date = _normalize_date_text(base_row.get("planned_start_date"))
        expected_start_date = _normalize_date_text(
            (state_row or {}).get("expected_start_date") or base_expected_start_date
        )
        base_expected_start_time = _iso_at(base_expected_start_date, "08:00:00")
        expected_start_time = (
            (state_row or {}).get("expected_start_time")
            or _iso_at(expected_start_date, "08:00:00")
        )
        base_expected_finish_time = _iso_at(promised_due_date, "18:00:00")
        expected_finish_time = (
            (state_row or {}).get("expected_finish_time")
            or base_expected_finish_time
        )
        manual_expected_start_date = _normalize_date_text((state_row or {}).get("expected_start_date"))
        manual_expected_start_time = str((state_row or {}).get("expected_start_time") or "").strip() or None
        manual_expected_start_shift = _expected_start_shift_from_datetime_text(manual_expected_start_time)
        manual_expected_start_override = (
            manual_expected_start_date is not None
            and (
                manual_expected_start_date != base_expected_start_date
                or (
                    manual_expected_start_time is not None
                    and manual_expected_start_time != base_expected_start_time
                )
            )
        )
        base_status = str(base_row.get("status") or "").strip().upper()
        status = str(
            (state_row or {}).get("status")
            or (state_row or {}).get("order_status")
            or base_row.get("status")
            or "OPEN"
        ).strip().upper()
        order_status = str(
            (state_row or {}).get("order_status") or status
        ).strip().upper()
        erp_status = base_status if base_status and base_status not in LOCAL_ORDER_STATUS_CODES else None
        source_bill_no = str(base_row.get("source_bill_no") or "").strip()
        final_process_code = str(
            (final_process_metrics or {}).get("final_process_code") or ""
        ).strip().upper() or None
        final_process_name_cn = str(
            (final_process_metrics or {}).get("final_process_name_cn")
            or final_process_code
            or ""
        ).strip() or None
        final_process_completed_qty = _to_number(
            (final_process_metrics or {}).get("final_process_completed_qty"),
            0,
        )
        final_process_eta_date = _normalize_date_text(
            (final_process_metrics or {}).get("final_process_eta_date")
        )
        shortage_items = (
            shortage_summary.get("items")
            if isinstance(shortage_summary, dict) and isinstance(shortage_summary.get("items"), list)
            else []
        )
        reference_version_no = str((reference_version or {}).get("version_no") or "").strip() or None
        reference_version_status = str((reference_version or {}).get("status") or "").strip().upper() or None
        scheduled_in_reference_version = schedule_fact is not None
        published_in_reference_version = (
            scheduled_in_reference_version and reference_version_status == "PUBLISHED"
        )
        published_version_no = str((published_version or {}).get("version_no") or "").strip() or None
        published_version_status = str((published_version or {}).get("status") or "").strip().upper() or None
        scheduled_in_published_version = published_schedule_fact is not None
        scheduled_start_date = _normalize_date_text((schedule_fact or {}).get("scheduled_start_date"))
        scheduled_start_time = str((schedule_fact or {}).get("scheduled_start_time") or "").strip() or None
        scheduled_start_shift = str((schedule_fact or {}).get("scheduled_start_shift") or "").strip().upper() or None
        scheduled_finish_date = _normalize_date_text((schedule_fact or {}).get("scheduled_finish_date"))
        scheduled_finish_time = str((schedule_fact or {}).get("scheduled_finish_time") or "").strip() or None
        published_scheduled_start_date = _normalize_date_text((published_schedule_fact or {}).get("scheduled_start_date"))
        published_scheduled_start_time = str((published_schedule_fact or {}).get("scheduled_start_time") or "").strip() or None
        published_scheduled_start_shift = str((published_schedule_fact or {}).get("scheduled_start_shift") or "").strip().upper() or None
        published_scheduled_finish_date = _normalize_date_text((published_schedule_fact or {}).get("scheduled_finish_date"))
        published_scheduled_finish_time = str((published_schedule_fact or {}).get("scheduled_finish_time") or "").strip() or None
        actual_workshop_codes = (
            schedule_fact.get("actual_workshop_codes")
            if isinstance(schedule_fact, dict) and isinstance(schedule_fact.get("actual_workshop_codes"), list)
            else []
        )
        actual_line_codes = (
            schedule_fact.get("actual_line_codes")
            if isinstance(schedule_fact, dict) and isinstance(schedule_fact.get("actual_line_codes"), list)
            else []
        )
        actual_process_codes = (
            schedule_fact.get("actual_process_codes")
            if isinstance(schedule_fact, dict) and isinstance(schedule_fact.get("actual_process_codes"), list)
            else []
        )
        expected_start_due_gap_days = _signed_due_gap_days(
            promised_due_date,
            expected_start_time,
            event_default_clock_text="08:00:00",
        )
        is_naturally_overdue = (
            expected_start_due_gap_days is not None and expected_start_due_gap_days > 0
        )
        order_window_finish_date = _normalize_date_text(expected_finish_time)
        order_window_due_gap_days = _signed_due_gap_days(promised_due_date, expected_finish_time)
        order_window_risk_level = "UNKNOWN"
        if order_window_due_gap_days is not None:
            if order_window_due_gap_days > 0:
                order_window_risk_level = "OVERDUE"
            elif order_window_due_gap_days >= -1:
                order_window_risk_level = "TIGHT"
            else:
                order_window_risk_level = "SAFE"
        final_process_due_gap_days = (
            _signed_due_gap_days(promised_due_date, final_process_eta_date)
            if final_process_eta_date
            else None
        )
        scheduled_due_gap_days = _signed_due_gap_days(promised_due_date, scheduled_finish_time or scheduled_finish_date)
        final_process_risk_level = "UNKNOWN"
        if final_process_due_gap_days is not None:
            if final_process_due_gap_days > 0:
                final_process_risk_level = "OVERDUE"
            elif final_process_due_gap_days >= -1:
                final_process_risk_level = "TIGHT"
            else:
                final_process_risk_level = "SAFE"
        scheduled_risk_level = "UNKNOWN"
        if scheduled_due_gap_days is not None:
            if scheduled_due_gap_days > 0:
                scheduled_risk_level = "OVERDUE"
            elif scheduled_due_gap_days >= -1:
                scheduled_risk_level = "TIGHT"
            else:
                scheduled_risk_level = "SAFE"
        delay_risk_source = "UNKNOWN"
        if final_process_eta_date:
            delay_risk_source = "FINAL_PROCESS_ETA"
        elif scheduled_finish_date:
            delay_risk_source = "SCHEDULE_FACT"
        elif _normalize_date_text(expected_finish_time):
            delay_risk_source = "ORDER_WINDOW"
        process_contexts = self._merge_process_contexts_with_schedule_fact(
            process_contexts=self._build_process_contexts(
                order_no=str(base_row["production_order_no"]),
                product_code=str(base_row["material_code"]),
                capacity_rows=capacity_rows,
                route_rows=route_rows,
                topology_by_process=topology_by_process,
            ),
            schedule_fact=schedule_fact,
        )
        manual_intervention_types = self._build_manual_intervention_types(
            manual_expected_start_override=manual_expected_start_override,
            priority_level=_normalize_priority_level((state_row or {}).get("priority_level"), PRIORITY_LEVEL_MAX),
            lock_flag=int((state_row or {}).get("lock_flag") or 0),
            frozen_flag=int((state_row or {}).get("frozen_flag") or 0),
        )
        reference_schedule_version_status_label = _schedule_version_status_label(
            reference_version_status
        )
        published_schedule_version_status_label = _schedule_version_status_label(
            published_version_status
        )
        if published_in_reference_version:
            reference_schedule_version_label = f"正式发布版 {reference_version_no}"
        elif scheduled_in_reference_version and reference_version_no:
            reference_schedule_version_label = (
                f"参考版 {reference_version_no}（{reference_schedule_version_status_label}）"
            )
        else:
            reference_schedule_version_label = "未进入任何参考版本"
        viewing_schedule_version_label = (
            f"当前查看版 {reference_version_no}（{reference_schedule_version_status_label}）"
            if reference_version_no
            else "当前查看版：未选择"
        )
        published_schedule_version_label = (
            f"正式执行版 {published_version_no}"
            if published_version_no
            else "正式执行版：未发布"
        )
        return {
            "order_no": base_row["production_order_no"],
            "product_code": base_row["material_code"],
            "product_name_cn": base_row["material_name"],
            "product_name": base_row["material_name"],
            "has_process_route": len(route_rows) > 0,
            "order_qty": production_qty,
            "completed_qty": completed_qty,
            "remaining_qty": remaining_qty,
            "progress_rate": progress_rate,
            "promised_due_date": promised_due_date,
            "due_in_days": _days_from_today(promised_due_date),
            "expected_start_time": expected_start_time,
            "expected_start_shift": _expected_start_shift_from_datetime_text(expected_start_time),
            "expected_finish_time": expected_finish_time,
            "expected_start_date": expected_start_date,
            "manual_expected_start_date": manual_expected_start_date,
            "manual_expected_start_time": manual_expected_start_time,
            "manual_expected_start_shift": manual_expected_start_shift,
            "manual_expected_start_override": manual_expected_start_override,
            "scheduled_start_date": scheduled_start_date,
            "scheduled_start_time": scheduled_start_time,
            "scheduled_start_shift": scheduled_start_shift,
            "scheduled_finish_date": scheduled_finish_date,
            "scheduled_finish_time": scheduled_finish_time,
            "expected_start_due_gap_days": expected_start_due_gap_days,
            "is_naturally_overdue": is_naturally_overdue,
            "order_window_due_gap_days": order_window_due_gap_days,
            "order_window_risk_level": order_window_risk_level,
            "priority_level": _normalize_priority_level((state_row or {}).get("priority_level"), PRIORITY_LEVEL_MAX),
            "urgent_flag": _urgent_flag_from_priority_level((state_row or {}).get("priority_level")),
            "lock_flag": int((state_row or {}).get("lock_flag") or 0),
            "frozen_flag": int((state_row or {}).get("frozen_flag") or 0),
            "status": status,
            "order_status": order_status,
            "erp_status": erp_status,
            "production_batch_no": (state_row or {}).get("production_batch_no")
            or (f"{source_bill_no}-B1" if source_bill_no else "-"),
            "final_process_code": final_process_code,
            "final_process_name_cn": final_process_name_cn,
            "final_process_completed_qty": final_process_completed_qty,
            "final_process_eta_date": final_process_eta_date,
            "final_process_due_gap_days": final_process_due_gap_days,
            "scheduled_due_gap_days": scheduled_due_gap_days,
            "final_process_risk_level": final_process_risk_level,
            "scheduled_risk_level": scheduled_risk_level,
            "delay_risk_source": delay_risk_source,
            "reference_version_no": reference_version_no,
            "reference_schedule_version_no": reference_version_no,
            "reference_schedule_version_status": reference_version_status,
            "reference_schedule_version_status_label": reference_schedule_version_status_label,
            "reference_schedule_version_label": reference_schedule_version_label,
            "viewing_schedule_version_no": reference_version_no,
            "viewing_schedule_version_status": reference_version_status,
            "viewing_schedule_version_status_label": reference_schedule_version_status_label,
            "viewing_schedule_version_label": viewing_schedule_version_label,
            "published_schedule_version_no": published_version_no,
            "published_schedule_version_status": published_version_status,
            "published_schedule_version_status_label": published_schedule_version_status_label,
            "published_schedule_version_label": published_schedule_version_label,
            "scheduled_in_reference_version": scheduled_in_reference_version,
            "scheduled_in_viewing_version": scheduled_in_reference_version,
            "published_in_reference_version": published_in_reference_version,
            "scheduled_in_published_version": scheduled_in_published_version,
            "current_schedule_version_no": reference_version_no if published_in_reference_version else None,
            "published_scheduled_start_date": published_scheduled_start_date,
            "published_scheduled_start_time": published_scheduled_start_time,
            "published_scheduled_start_shift": published_scheduled_start_shift,
            "published_scheduled_finish_date": published_scheduled_finish_date,
            "published_scheduled_finish_time": published_scheduled_finish_time,
            "material_shortage_count": int((shortage_summary or {}).get("shortage_material_count") or 0),
            "material_shortage_summary": str((shortage_summary or {}).get("summary_text") or "").strip(),
            "material_shortage_start_date": _normalize_date_text(
                (shortage_summary or {}).get("shortage_start_date")
            ),
            "material_shortage_first_process_code": str(
                (shortage_summary or {}).get("first_impacted_process_code") or ""
            ).strip().upper()
            or None,
            "material_shortage_first_process_name_cn": str(
                (shortage_summary or {}).get("first_impacted_process_name_cn") or ""
            ).strip()
            or None,
            "material_shortage_first_material_code": str(
                (shortage_summary or {}).get("first_material_code") or ""
            ).strip().upper()
            or None,
            "material_shortage_first_material_name": str(
                (shortage_summary or {}).get("first_material_name") or ""
            ).strip()
            or None,
            "material_shortage_items": shortage_items,
            "actual_workshop_codes": actual_workshop_codes,
            "actual_line_codes": actual_line_codes,
            "actual_process_codes": actual_process_codes,
            "manual_intervention_types": manual_intervention_types,
            "manual_intervention_count": len(manual_intervention_types),
            "process_contexts": process_contexts,
        }

    def _build_expected_start_save_impact(self, row: dict[str, Any]) -> dict[str, Any]:
        expected_start_date = _normalize_date_text(row.get("expected_start_date"))
        expected_start_shift = _normalize_shift_code(row.get("expected_start_shift"))
        expected_start_text = (
            f"{expected_start_date} {'夜班' if expected_start_shift == 'NIGHT' else '白班'}"
            if expected_start_date
            else "未设置"
        )
        viewing_version_no = str(row.get("viewing_schedule_version_no") or "").strip() or None
        published_version_no = str(row.get("published_schedule_version_no") or "").strip() or None
        impacts_viewing_version = bool(row.get("scheduled_in_viewing_version"))
        impacts_published_version = bool(row.get("scheduled_in_published_version"))
        viewing_conflict = impacts_viewing_version and (
            _normalize_date_text(row.get("scheduled_start_date")) != expected_start_date
            or _normalize_shift_code(row.get("scheduled_start_shift")) != expected_start_shift
        )
        published_conflict = impacts_published_version and (
            _normalize_date_text(row.get("published_scheduled_start_date")) != expected_start_date
            or _normalize_shift_code(row.get("published_scheduled_start_shift")) != expected_start_shift
        )
        has_schedule_conflict = viewing_conflict or published_conflict
        causes_unavoidable_delay = bool(row.get("is_naturally_overdue"))
        requires_reschedule = has_schedule_conflict
        summary_items = [f"手工开工硬约束已设为 {expected_start_text}"]
        if impacts_published_version and published_version_no:
            summary_items.append(f"会影响正式执行版 {published_version_no}")
        else:
            summary_items.append("当前不直接改写正式执行版")
        if impacts_viewing_version and viewing_version_no:
            summary_items.append(f"会影响当前查看版 {viewing_version_no} 的判断口径")
        if has_schedule_conflict:
            summary_items.append("与现有排程事实冲突")
        if causes_unavoidable_delay:
            summary_items.append("该订单已必然延期")
        if requires_reschedule:
            summary_items.append("需要立即重排")
        return {
            "expected_start_text": expected_start_text,
            "impacts_viewing_version": impacts_viewing_version,
            "viewing_version_no": viewing_version_no,
            "impacts_published_version": impacts_published_version,
            "published_version_no": published_version_no,
            "has_schedule_conflict": has_schedule_conflict,
            "causes_unavoidable_delay": causes_unavoidable_delay,
            "requires_reschedule": requires_reschedule,
            "summary_items": summary_items,
        }

    def _build_reference_schedule_context(self, version_no: str | None) -> dict[str, Any]:
        if not version_no:
            return {"version": None, "order_map": {}}
        version = self.get_schedule_version(version_no)
        task_rows = self._list_schedule_task_detail_rows(version_no)
        order_map: dict[str, dict[str, Any]] = {}
        for row in task_rows:
            order_no = str(row.get("production_order_no") or "").strip()
            calendar_date = _normalize_date_text(row.get("calendar_date"))
            if not order_no or not calendar_date:
                continue
            shift_code = _normalize_shift_code(row.get("shift_code"))
            start_slot = _slot_index_from_text(calendar_date, shift_code)
            scheduled_finish_date = (
                (date.fromisoformat(calendar_date) + timedelta(days=1)).isoformat()
                if shift_code == "NIGHT"
                else calendar_date
            )
            scheduled_start_time = str(row.get("plan_start_time") or "").strip() or _iso_at(
                calendar_date,
                "20:00:00" if shift_code == "NIGHT" else "08:00:00",
            )
            scheduled_finish_time = _iso_at(
                scheduled_finish_date,
                "08:00:00" if shift_code == "NIGHT" else "20:00:00",
            )
            current = order_map.get(order_no)
            if current is None:
                current = {
                    "scheduled_start_slot": start_slot,
                    "scheduled_start_date": calendar_date,
                    "scheduled_start_time": scheduled_start_time,
                    "scheduled_start_shift": shift_code,
                    "scheduled_finish_slot": start_slot,
                    "scheduled_finish_date": scheduled_finish_date,
                    "scheduled_finish_time": scheduled_finish_time,
                    "actual_workshop_codes": set(),
                    "actual_line_codes": set(),
                    "actual_process_codes": set(),
                    "actual_process_map": {},
                }
                order_map[order_no] = current
            if start_slot < int(current["scheduled_start_slot"]):
                current["scheduled_start_slot"] = start_slot
                current["scheduled_start_date"] = calendar_date
                current["scheduled_start_time"] = scheduled_start_time
                current["scheduled_start_shift"] = shift_code
            if start_slot >= int(current["scheduled_finish_slot"]):
                current["scheduled_finish_slot"] = start_slot
                current["scheduled_finish_date"] = scheduled_finish_date
                current["scheduled_finish_time"] = scheduled_finish_time

            workshop_code = str(row.get("workshop_code") or "").strip().upper()
            line_code = str(row.get("line_code") or "").strip().upper()
            process_code = str(row.get("process_code") or "").strip().upper()
            process_name_cn = str(row.get("process_name_cn") or process_code).strip() or process_code
            if workshop_code:
                current["actual_workshop_codes"].add(workshop_code)
            if line_code:
                current["actual_line_codes"].add(line_code)
            if process_code:
                current["actual_process_codes"].add(process_code)
                process_item = current["actual_process_map"].get(process_code)
                if process_item is None:
                    process_item = {
                        "process_code": process_code,
                        "process_name_cn": process_name_cn,
                        "actual_workshop_codes": set(),
                        "actual_line_codes": set(),
                    }
                    current["actual_process_map"][process_code] = process_item
                if workshop_code:
                    process_item["actual_workshop_codes"].add(workshop_code)
                if line_code:
                    process_item["actual_line_codes"].add(line_code)

        normalized_order_map: dict[str, dict[str, Any]] = {}
        for order_no, item in order_map.items():
            normalized_order_map[order_no] = {
                "scheduled_start_date": item["scheduled_start_date"],
                "scheduled_start_time": item["scheduled_start_time"],
                "scheduled_start_shift": item["scheduled_start_shift"],
                "scheduled_finish_date": item["scheduled_finish_date"],
                "scheduled_finish_time": item["scheduled_finish_time"],
                "actual_workshop_codes": sorted(item["actual_workshop_codes"]),
                "actual_line_codes": sorted(item["actual_line_codes"]),
                "actual_process_codes": sorted(item["actual_process_codes"]),
                "actual_process_map": {
                    process_code: {
                        "process_code": process_code,
                        "process_name_cn": process_item["process_name_cn"],
                        "actual_workshop_codes": sorted(process_item["actual_workshop_codes"]),
                        "actual_line_codes": sorted(process_item["actual_line_codes"]),
                    }
                    for process_code, process_item in item["actual_process_map"].items()
                },
            }
        return {"version": version, "order_map": normalized_order_map}

    def _merge_process_contexts_with_schedule_fact(
        self,
        *,
        process_contexts: list[dict[str, Any]],
        schedule_fact: dict[str, Any] | None,
    ) -> list[dict[str, Any]]:
        actual_process_map = (
            schedule_fact.get("actual_process_map")
            if isinstance(schedule_fact, dict) and isinstance(schedule_fact.get("actual_process_map"), dict)
            else {}
        )
        merged: list[dict[str, Any]] = []
        for context in process_contexts:
            process_code = str(context.get("process_code") or "").strip().upper()
            actual_process_item = actual_process_map.get(process_code, {})
            merged.append(
                {
                    **context,
                    "actual_workshop_codes": actual_process_item.get("actual_workshop_codes", []),
                    "actual_line_codes": actual_process_item.get("actual_line_codes", []),
                }
            )
        return merged

    def _build_manual_intervention_types(
        self,
        *,
        manual_expected_start_override: bool,
        priority_level: int,
        lock_flag: int,
        frozen_flag: int,
    ) -> list[str]:
        items: list[str] = []
        if manual_expected_start_override:
            items.append("EXPECTED_START")
        if priority_level < PRIORITY_LEVEL_MAX:
            items.append("PRIORITY")
        if lock_flag == 1:
            items.append("LOCK")
        if frozen_flag == 1:
            items.append("FREEZE")
        return items

    def _build_process_contexts(
        self,
        *,
        order_no: str,
        product_code: str,
        capacity_rows: list[dict[str, Any]],
        route_rows: list[dict[str, Any]],
        topology_by_process: dict[str, list[dict[str, Any]]],
    ) -> list[dict[str, Any]]:
        capacity_rows_by_process = self._best_capacity_row_by_process(capacity_rows)

        if route_rows:
            contexts: list[dict[str, Any]] = []
            for index, row in enumerate(route_rows):
                process_code = str(row.get("process_code") or "").strip().upper()
                if not process_code:
                    continue
                candidate_rows = capacity_rows_by_process.get(process_code) or topology_by_process.get(process_code) or []
                candidate_contexts = self._build_candidate_process_contexts(
                    process_code=process_code,
                    process_name_cn=str(row.get("process_name_cn") or process_code).strip() or process_code,
                    candidate_rows=candidate_rows,
                )
                primary_candidate = candidate_contexts[0] if candidate_contexts else {}
                capacity_per_shift = _to_number(primary_candidate.get("capacity_per_shift"), 0)
                contexts.append(
                    {
                        "sequence_no": index + 1,
                        "process_code": process_code,
                        "company_code": str(
                            primary_candidate.get("company_code") or DEFAULT_COMPANY_CODE
                        ).strip().upper()
                        or DEFAULT_COMPANY_CODE,
                        "process_name_cn": str(row.get("process_name_cn") or process_code).strip() or process_code,
                        "workshop_code": str(primary_candidate.get("workshop_code") or "-").strip() or "-",
                        "line_code": str(primary_candidate.get("line_code") or "-").strip() or "-",
                        "candidate_workshop_codes": [item["workshop_code"] for item in candidate_contexts],
                        "candidate_line_codes": [item["line_code"] for item in candidate_contexts],
                        "candidate_contexts": candidate_contexts,
                        "dependency_type": str(row.get("dependency_type") or "FS").strip().upper() or "FS",
                        "default_capacity_per_shift": capacity_per_shift,
                        "capacity_per_shift": capacity_per_shift,
                    }
                )
            return contexts

        contexts = []
        deduplicated_rows = sorted(
            capacity_rows_by_process.values(),
            key=lambda items: str((items[0] if items else {}).get("process_code") or ""),
        )
        for index, process_rows in enumerate(deduplicated_rows):
            primary_row = process_rows[0] if process_rows else {}
            process_code = str(primary_row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            candidate_contexts = self._build_candidate_process_contexts(
                process_code=process_code,
                process_name_cn=str(primary_row.get("process_name") or process_code).strip() or process_code,
                candidate_rows=process_rows,
            )
            primary_candidate = candidate_contexts[0] if candidate_contexts else {}
            contexts.append(
                {
                    "sequence_no": index + 1,
                    "process_code": process_code,
                    "company_code": str(primary_candidate.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() or DEFAULT_COMPANY_CODE,
                    "process_name_cn": str(primary_row.get("process_name") or process_code).strip() or process_code,
                    "workshop_code": str(primary_candidate.get("workshop_code") or "-").strip() or "-",
                    "line_code": str(primary_candidate.get("line_code") or "-").strip() or "-",
                    "candidate_workshop_codes": [item["workshop_code"] for item in candidate_contexts],
                    "candidate_line_codes": [item["line_code"] for item in candidate_contexts],
                    "candidate_contexts": candidate_contexts,
                    "dependency_type": "FS",
                    "default_capacity_per_shift": _to_number(primary_candidate.get("capacity_per_shift"), 0),
                    "capacity_per_shift": _to_number(primary_candidate.get("capacity_per_shift"), 0),
                }
            )
        return contexts

    def _build_candidate_process_contexts(
        self,
        *,
        process_code: str,
        process_name_cn: str,
        candidate_rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        seen_keys: set[tuple[str, str, str]] = set()
        for row in candidate_rows:
            company_code = str(row.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() or DEFAULT_COMPANY_CODE
            workshop_code = str(row.get("workshop_code") or "").strip().upper()
            line_code = str(row.get("line_code") or "").strip().upper()
            if not workshop_code or not line_code:
                continue
            key = (company_code, workshop_code, line_code)
            if key in seen_keys:
                continue
            seen_keys.add(key)
            capacity_per_shift = _to_number(
                row.get("capacity_qty"),
                _to_number(row.get("capacity_per_shift"), 0),
            )
            out.append(
                {
                    "company_code": company_code,
                    "workshop_code": workshop_code,
                    "line_code": line_code,
                    "process_code": process_code,
                    "process_name_cn": process_name_cn,
                    "default_capacity_per_shift": capacity_per_shift,
                    "capacity_per_shift": capacity_per_shift,
                }
            )
        out.sort(
            key=lambda item: (
                -_to_number(item.get("capacity_per_shift"), 0),
                str(item.get("workshop_code") or ""),
                str(item.get("line_code") or ""),
            )
        )
        return out

    def _build_schedule_process_contexts(
        self,
        *,
        order_no: str,
        product_code: str,
        capacity_rows: list[dict[str, Any]],
        route_rows: list[dict[str, Any]],
        topology_by_process: dict[str, list[dict[str, Any]]],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> list[dict[str, Any]]:
        if len(route_rows) == 0:
            raise server_error(
                code="SCHEDULE_ROUTE_REQUIRED",
                message="Process route is required for schedule generation.",
                details={"order_no": order_no, "product_code": product_code},
            )

        contexts = self._build_process_contexts(
            order_no=order_no,
            product_code=product_code,
            capacity_rows=capacity_rows,
            route_rows=route_rows,
            topology_by_process=topology_by_process,
        )
        if len(contexts) == 0:
            raise server_error(
                code="SCHEDULE_PROCESS_CONTEXTS_EMPTY",
                message="No process contexts available for schedule generation.",
                details={"order_no": order_no, "product_code": product_code},
            )

        for context in contexts:
            process_code = str(context.get("process_code") or "").strip().upper()
            candidate_contexts = (
                context.get("candidate_contexts")
                if isinstance(context.get("candidate_contexts"), list)
                else []
            )
            company_code = str(context.get("company_code") or "").strip().upper() or "COMPANY-MAIN"
            default_capacity_per_shift = _to_number(
                context.get("default_capacity_per_shift"),
                _to_number(context.get("capacity_per_shift"), 0),
            )
            if not process_code:
                raise server_error(
                    code="SCHEDULE_PROCESS_CODE_REQUIRED",
                    message="Process code is missing in schedule process context.",
                    details={"order_no": order_no, "product_code": product_code},
                )
            if len(candidate_contexts) == 0:
                raise server_error(
                    code="SCHEDULE_PROCESS_CANDIDATE_LINES_EMPTY",
                    message="No candidate workshop/line exists for scheduled process.",
                    details={
                        "order_no": order_no,
                        "product_code": product_code,
                        "process_code": process_code,
                    },
                )
            if default_capacity_per_shift <= SCHEDULE_NUMBER_EPSILON:
                raise server_error(
                    code="SCHEDULE_CAPACITY_PER_SHIFT_INVALID",
                    message="capacity_per_shift must be greater than 0 for all scheduled processes.",
                    details={
                        "order_no": order_no,
                        "product_code": product_code,
                        "process_code": process_code,
                        "capacity_per_shift": default_capacity_per_shift,
                    },
            )
            context["company_code"] = company_code
            primary_candidate = candidate_contexts[0]
            context["workshop_code"] = str(primary_candidate.get("workshop_code") or "").strip().upper()
            context["line_code"] = str(primary_candidate.get("line_code") or "").strip().upper()
            context["process_code"] = process_code
            context["default_capacity_per_shift"] = default_capacity_per_shift
            context["capacity_per_shift"] = default_capacity_per_shift
        return contexts

    def _build_schedule_capacity_resolver(
        self,
        *,
        capacity_source_mode: str,
    ) -> dict[str, dict[tuple[str, str, str, str, str], float]]:
        planned_map: dict[tuple[str, str, str, str, str], float] = {}
        actual_map: dict[tuple[str, str, str, str, str], float] = {}
        if capacity_source_mode in {"PLANNED", "ACTUAL"}:
            for row in fetch_all(
                self.connection,
                """
                SELECT
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    planned_capacity_qty
                FROM daily_line_capacity_plan
                """
            ):
                key = (
                    str(row.get("calendar_date") or "").strip(),
                    str(row.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN",
                    str(row.get("workshop_code") or "").strip().upper(),
                    str(row.get("line_code") or "").strip().upper(),
                    str(row.get("process_code") or "").strip().upper(),
                )
                planned_map[key] = _to_number(row.get("planned_capacity_qty"), 0)
        if capacity_source_mode == "ACTUAL":
            for row in fetch_all(
                self.connection,
                """
                SELECT
                    calendar_date,
                    company_code,
                    workshop_code,
                    line_code,
                    process_code,
                    actual_capacity_qty
                FROM daily_line_capacity_actual
                """
            ):
                key = (
                    str(row.get("calendar_date") or "").strip(),
                    str(row.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN",
                    str(row.get("workshop_code") or "").strip().upper(),
                    str(row.get("line_code") or "").strip().upper(),
                    str(row.get("process_code") or "").strip().upper(),
                )
                actual_map[key] = _to_number(row.get("actual_capacity_qty"), 0)
        return {"planned": planned_map, "actual": actual_map}

    def _resolve_effective_capacity_per_shift(
        self,
        *,
        process_context: dict[str, Any],
        calendar_date: str,
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> float:
        company_code = str(process_context.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
        workshop_code = str(process_context.get("workshop_code") or "").strip().upper()
        line_code = str(process_context.get("line_code") or "").strip().upper()
        process_code = str(process_context.get("process_code") or "").strip().upper()
        lookup_key = (
            str(calendar_date or "").strip(),
            company_code,
            workshop_code,
            line_code,
            process_code,
        )
        actual_map = capacity_resolver.get("actual") or {}
        if lookup_key in actual_map:
            return _to_number(actual_map[lookup_key], 0)
        planned_map = capacity_resolver.get("planned") or {}
        if lookup_key in planned_map:
            return _to_number(planned_map[lookup_key], 0)
        return _to_number(process_context.get("default_capacity_per_shift"), 0)

    def _capacity_usage_key(
        self,
        *,
        process_context: dict[str, Any],
        slot_index: int,
        calendar_date: str,
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> tuple[Any, ...]:
        company_code = str(process_context.get("company_code") or DEFAULT_COMPANY_CODE).strip().upper() or DEFAULT_COMPANY_CODE
        workshop_code = str(process_context.get("workshop_code") or "").strip().upper()
        line_code = str(process_context.get("line_code") or "").strip().upper()
        process_code = str(process_context.get("process_code") or "").strip().upper()
        lookup_key = (
            str(calendar_date or "").strip(),
            company_code,
            workshop_code,
            line_code,
            process_code,
        )
        actual_map = capacity_resolver.get("actual") or {}
        planned_map = capacity_resolver.get("planned") or {}
        if lookup_key in actual_map or lookup_key in planned_map:
            return ("DAY_TOTAL", *lookup_key)
        return ("SHIFT", int(slot_index), workshop_code, line_code, process_code)

    def _select_candidate_context_for_slot(
        self,
        *,
        process_context: dict[str, Any],
        slot_index: int,
        calendar_date: str,
        used_capacity_by_slot: dict[tuple[Any, ...], float],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> tuple[dict[str, Any] | None, float]:
        candidate_contexts = (
            process_context.get("candidate_contexts")
            if isinstance(process_context.get("candidate_contexts"), list)
            else []
        )
        if len(candidate_contexts) == 0:
            candidate_contexts = [process_context]

        best_candidate: dict[str, Any] | None = None
        best_available_capacity = 0.0
        best_sort_key: tuple[float, float, str, str] | None = None
        for candidate in candidate_contexts:
            capacity_per_shift = self._resolve_effective_capacity_per_shift(
                process_context=candidate,
                calendar_date=calendar_date,
                capacity_resolver=capacity_resolver,
            )
            usage_key = self._capacity_usage_key(
                process_context=candidate,
                slot_index=slot_index,
                calendar_date=calendar_date,
                capacity_resolver=capacity_resolver,
            )
            used_capacity = _to_number(used_capacity_by_slot.get(usage_key), 0)
            available_capacity = max(0.0, capacity_per_shift - used_capacity)
            if available_capacity <= SCHEDULE_NUMBER_EPSILON:
                continue
            load_ratio = used_capacity / capacity_per_shift if capacity_per_shift > SCHEDULE_NUMBER_EPSILON else 1.0
            sort_key = (
                round(load_ratio, 8),
                -available_capacity,
                str(candidate.get("workshop_code") or ""),
                str(candidate.get("line_code") or ""),
            )
            if best_sort_key is None or sort_key < best_sort_key:
                best_sort_key = sort_key
                best_candidate = candidate
                best_available_capacity = available_capacity
        return best_candidate, best_available_capacity

    def _estimate_required_shifts_for_process(
        self,
        *,
        process_context: dict[str, Any],
        required_qty: float,
        start_slot: int,
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> int:
        remaining_qty = max(0.0, _to_number(required_qty, 0))
        if remaining_qty <= SCHEDULE_NUMBER_EPSILON:
            return 0
        slot_cursor = max(0, int(start_slot))
        occupied_slot_set: set[int] = set()
        simulated_used_capacity: dict[tuple[Any, ...], float] = {}
        for loop_guard in range(SCHEDULE_SLOT_SEARCH_GUARD):
            if remaining_qty <= SCHEDULE_NUMBER_EPSILON:
                break
            slot_index, calendar_date, _ = self._next_working_slot(
                start_slot=slot_cursor,
                planning_rules=planning_rules,
                day_mode_cache=day_mode_cache,
                allow_exact_start_slot=slot_cursor == max(0, int(start_slot)),
            )
            candidate_context, available_capacity = self._select_candidate_context_for_slot(
                process_context=process_context,
                slot_index=slot_index,
                calendar_date=calendar_date,
                used_capacity_by_slot=simulated_used_capacity,
                capacity_resolver=capacity_resolver,
            )
            if candidate_context is None or available_capacity <= SCHEDULE_NUMBER_EPSILON:
                slot_cursor = slot_index + 1
                continue
            usage_key = self._capacity_usage_key(
                process_context=candidate_context,
                slot_index=slot_index,
                calendar_date=calendar_date,
                capacity_resolver=capacity_resolver,
            )
            used_capacity = _to_number(simulated_used_capacity.get(usage_key), 0)
            planned_qty = min(remaining_qty, available_capacity)
            simulated_used_capacity[usage_key] = used_capacity + planned_qty
            remaining_qty -= planned_qty
            occupied_slot_set.add(slot_index)
            slot_cursor = slot_index if remaining_qty > SCHEDULE_NUMBER_EPSILON else slot_index + 1
        else:
            raise server_error(
                code="SCHEDULE_REQUIRED_SHIFTS_ESTIMATION_EXCEEDED",
                message="Unable to estimate required shifts within slot guard limit.",
                details={
                    "order_no": str(process_context.get("order_no") or ""),
                    "process_code": str(process_context.get("process_code") or ""),
                    "required_qty": required_qty,
                },
            )
        return len(occupied_slot_set)

    def _refresh_bom_children(self, parent_material_code: str) -> None:
        children = self.factory.material_gateway.fetch_bom_children(parent_material_code)
        child_codes = [str(item["child_material_code"]) for item in children]
        supply_rows = self.factory.supply_gateway.fetch_material_supply(child_codes)
        supply_map = {str(row["material_code"]): row for row in supply_rows}
        updated_at = utc_now()
        rows_to_store: list[dict[str, Any]] = []
        normalized_supply_rows: list[dict[str, Any]] = []
        for row in supply_rows:
            normalized_supply_rows.append(
                {
                    "material_code": str(row["material_code"]),
                    "supply_type_code": row["supply_type_code"],
                    "supply_type_name": row["supply_type_name"],
                    "snapshot_time": updated_at,
                    "updated_at": updated_at,
                }
            )
        for index, item in enumerate(children):
            material_code = str(item["child_material_code"])
            supply_row = supply_map.get(material_code, {})
            supply_type_code = item.get("supply_type_code") or supply_row.get("supply_type_code")
            supply_type_name = item.get("supply_type_name") or supply_row.get("supply_type_name")
            if not supply_type_code or not supply_type_name:
                raise bad_request(
                    code="SUPPLY_TYPE_MISSING",
                    message="Supply type is missing in ERP BOM refresh result.",
                    details={"material_code": material_code},
                )
            rows_to_store.append(
                {
                    "parent_material_code": parent_material_code,
                    "child_material_code": material_code,
                    "child_material_name": item["child_material_name"],
                    "child_specification": item.get("child_specification"),
                    "usage_numerator": item.get("usage_numerator"),
                    "usage_denominator": item.get("usage_denominator"),
                    "child_unit": item.get("child_unit"),
                    "supply_type_code": supply_type_code,
                    "supply_type_name": supply_type_name,
                    "inventory_qty": None,
                    "inventory_status": "UNKNOWN",
                    "expandable": supply_type_code == "SELF_MADE",
                    "display_order": index,
                    "updated_at": updated_at,
                }
            )
        with transaction(self.connection):
            self.factory.supply_repository.upsert_many(normalized_supply_rows)
            self.factory.bom_children_repository.replace_for_parent(
                parent_material_code,
                rows_to_store,
            )

    def _process_name_by_code(self) -> dict[str, str]:
        names: dict[str, str] = {}
        for row in self._list_route_rows():
            process_code = str(row.get("process_code") or "").strip().upper()
            if process_code and process_code not in names:
                names[process_code] = str(row.get("process_name_cn") or process_code)
        for row in fetch_all(
            self.connection,
            """
            SELECT DISTINCT process_code, process_name
            FROM capacity_bindings
            WHERE process_code IS NOT NULL
            """,
        ):
            process_code = str(row.get("process_code") or "").strip().upper()
            if process_code and process_code not in names:
                names[process_code] = str(row.get("process_name") or process_code)
        return names

    def _normalize_strategy_code(self, value: object) -> str:
        normalized = str(value or "KEY_ORDER_FIRST").strip().upper() or "KEY_ORDER_FIRST"
        if normalized not in SUPPORTED_STRATEGY_CODES:
            raise bad_request(
                code="SCHEDULE_STRATEGY_INVALID",
                message="strategy_code is invalid.",
                details={
                    "strategy_code": normalized,
                    "supported": sorted(SUPPORTED_STRATEGY_CODES),
                },
            )
        return normalized

    def _normalize_capacity_source_mode(self, value: object) -> str:
        normalized = str(value or "ACTUAL").strip().upper() or "ACTUAL"
        if normalized not in SUPPORTED_CAPACITY_SOURCE_MODES:
            raise bad_request(
                code="SCHEDULE_CAPACITY_SOURCE_MODE_INVALID",
                message="capacity_source_mode is invalid.",
                details={
                    "capacity_source_mode": normalized,
                    "supported": sorted(SUPPORTED_CAPACITY_SOURCE_MODES),
                },
            )
        return normalized

    def _normalize_use_order_state_window(self, value: object) -> bool:
        if value is None:
            return False
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            if value == 1:
                return True
            if value == 0:
                return False
        normalized = str(value).strip().lower()
        if normalized in {"1", "true", "yes", "y", "on"}:
            return True
        if normalized in {"0", "false", "no", "n", "off", ""}:
            return False
        raise bad_request(
            code="SCHEDULE_USE_ORDER_STATE_WINDOW_INVALID",
            message="use_order_state_window is invalid.",
            details={"use_order_state_window": value},
        )

    def _normalize_weekend_rest_mode(self, value: object) -> str:
        normalized = str(value or "DOUBLE").strip().upper() or "DOUBLE"
        if normalized not in WEEKEND_REST_MODES:
            raise bad_request(
                code="WEEKEND_REST_MODE_INVALID",
                message="weekend_rest_mode is invalid.",
                details={
                    "weekend_rest_mode": normalized,
                    "supported": sorted(WEEKEND_REST_MODES),
                },
            )
        return normalized

    def _normalize_date_shift_mode(self, value: object) -> str:
        normalized = str(value or "").strip().upper()
        if normalized not in DATE_SHIFT_MODES:
            raise bad_request(
                code="DATE_SHIFT_MODE_INVALID",
                message="date_shift_mode is invalid.",
                details={
                    "date_shift_mode": normalized,
                    "supported": sorted(DATE_SHIFT_MODES),
                },
            )
        return normalized

    def _normalize_date_shift_mode_by_date(self, value: object) -> dict[str, str]:
        if value is None:
            return {}
        if not isinstance(value, dict):
            raise bad_request(
                code="DATE_SHIFT_MODE_BY_DATE_INVALID",
                message="date_shift_mode_by_date must be an object.",
            )
        out: dict[str, str] = {}
        for key, mode_value in value.items():
            date_text = _normalize_date_text(key)
            if not date_text:
                raise bad_request(
                    code="DATE_SHIFT_MODE_DATE_INVALID",
                    message="date_shift_mode_by_date contains invalid date key.",
                    details={"date": str(key)},
                )
            out[date_text] = self._normalize_date_shift_mode(mode_value)
        return out

    def _build_planning_rules_from_current_config(self) -> dict[str, Any]:
        row = self._get_rules_row()
        return {
            "skip_statutory_holidays": bool(row.get("skip_statutory_holidays")),
            "weekend_rest_mode": self._normalize_weekend_rest_mode(
                row.get("weekend_rest_mode")
            ),
            "date_shift_mode_by_date": self._normalize_date_shift_mode_by_date(
                loads(row.get("date_shift_mode_by_date_json") or "{}")
            ),
        }

    def _best_capacity_row_by_process(
        self,
        capacity_rows: list[dict[str, Any]],
    ) -> dict[str, list[dict[str, Any]]]:
        out: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in capacity_rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            out[process_code].append(row)
        for process_code in list(out):
            out[process_code] = sorted(
                out[process_code],
                key=lambda item: (
                    -_to_number(item.get("capacity_qty"), 0),
                    str(item.get("workshop_code") or ""),
                    str(item.get("line_code") or ""),
                ),
            )
        return out

    def _enabled_topology_by_process(
        self,
        topology_rows: list[dict[str, Any]],
    ) -> dict[str, list[dict[str, Any]]]:
        out: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in topology_rows:
            if int(_to_number(row.get("enabled_flag"), 0)) != 1:
                continue
            process_code = str(row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            out[process_code].append(row)
        for process_code in list(out):
            out[process_code] = sorted(
                out[process_code],
                key=lambda item: (
                    -_to_number(item.get("capacity_per_shift"), 0),
                    str(item.get("workshop_code") or ""),
                    str(item.get("line_code") or ""),
                ),
            )
        return out

    def _build_base_schedule_hints(self, version_no: str) -> dict[str, dict[str, Any]]:
        rows = fetch_all(
            self.connection,
            """
            SELECT
                task_no,
                production_order_no,
                process_code,
                calendar_date,
                shift_code
            FROM schedule_tasks
            WHERE version_no = ?
            ORDER BY task_no ASC
            """,
            (version_no,),
        )
        hints: dict[str, dict[str, Any]] = {}
        for row in rows:
            order_no = str(row.get("production_order_no") or "").strip()
            process_code = str(row.get("process_code") or "").strip().upper()
            calendar_date = _normalize_date_text(row.get("calendar_date"))
            if not order_no or not process_code or not calendar_date:
                raise server_error(
                    code="BASE_SCHEDULE_TASK_INVALID",
                    message="Base schedule task is invalid.",
                    details={"version_no": version_no, "task_no": row.get("task_no")},
                )
            shift_code = _normalize_shift_code(row.get("shift_code"))
            slot_index = _slot_index_from_text(calendar_date, shift_code)
            task_no = int(_to_number(row.get("task_no"), 0))
            hint = hints.setdefault(
                order_no,
                {
                    "first_task_no": task_no,
                    "first_slot": slot_index,
                    "process_first_slot": {},
                },
            )
            if task_no < int(hint["first_task_no"]):
                hint["first_task_no"] = task_no
            if slot_index < int(hint["first_slot"]):
                hint["first_slot"] = slot_index
            process_first_slot = hint["process_first_slot"]
            existing_process_slot = process_first_slot.get(process_code)
            if existing_process_slot is None or slot_index < int(existing_process_slot):
                process_first_slot[process_code] = slot_index
        return hints

    def _resolve_day_shift_mode(
        self,
        *,
        date_text: str,
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
    ) -> str:
        cached = day_mode_cache.get(date_text)
        if cached:
            return cached

        date_shift_mode_by_date = planning_rules["date_shift_mode_by_date"]
        manual_mode = date_shift_mode_by_date.get(date_text)
        if manual_mode:
            day_mode_cache[date_text] = manual_mode
            return manual_mode

        mode = "DAY"
        if planning_rules["skip_statutory_holidays"] is True:
            year_value = int(date_text[:4])
            if year_value not in SUPPORTED_CN_HOLIDAY_YEARS:
                raise bad_request(
                    code="HOLIDAY_YEAR_UNSUPPORTED",
                    message="Statutory holiday data is unsupported for schedule date.",
                    details={"year": year_value, "date": date_text},
                )
            if self._is_cn_statutory_holiday(date_text):
                mode = "REST"
            else:
                weekday = date.fromisoformat(date_text).weekday()
                weekend_mode = planning_rules["weekend_rest_mode"]
                if weekend_mode == "DOUBLE" and weekday in (5, 6):
                    mode = "REST"
                elif weekend_mode == "SINGLE" and weekday == 6:
                    mode = "REST"

        day_mode_cache[date_text] = mode
        return mode

    def _allowed_shifts_for_day_mode(self, mode: str) -> tuple[str, ...]:
        if mode == "REST":
            return ()
        if mode == "BOTH":
            return ("DAY", "NIGHT")
        if mode == "NIGHT":
            return ("NIGHT",)
        return ("DAY",)

    def _next_working_slot(
        self,
        *,
        start_slot: int,
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
        allow_exact_start_slot: bool = False,
    ) -> tuple[int, str, str]:
        slot_cursor = max(0, int(start_slot))
        for _ in range(SCHEDULE_SLOT_SEARCH_GUARD):
            day_value, shift_code = _slot_to_date_shift(slot_cursor)
            date_text = day_value.isoformat()
            if allow_exact_start_slot and slot_cursor == max(0, int(start_slot)):
                return slot_cursor, date_text, shift_code
            day_mode = self._resolve_day_shift_mode(
                date_text=date_text,
                planning_rules=planning_rules,
                day_mode_cache=day_mode_cache,
            )
            if shift_code in self._allowed_shifts_for_day_mode(day_mode):
                return slot_cursor, date_text, shift_code
            slot_cursor += 1
        raise server_error(
            code="SCHEDULE_WORKING_SLOT_NOT_FOUND",
            message="Cannot find working slot under current calendar rules.",
            details={"start_slot": start_slot},
        )

    def _allocate_process_tasks(
        self,
        *,
        order_no: str,
        process_context: dict[str, Any],
        required_qty: float,
        first_slot: int,
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
        used_capacity_by_slot: dict[tuple[Any, ...], float],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> tuple[list[dict[str, Any]], int]:
        remaining_qty = max(0.0, _to_number(required_qty, 0))
        if remaining_qty <= SCHEDULE_NUMBER_EPSILON:
            return [], int(first_slot)

        process_code = str(process_context.get("process_code") or "").strip().upper()
        default_capacity_per_shift = _to_number(process_context.get("default_capacity_per_shift"), 0)
        if default_capacity_per_shift <= SCHEDULE_NUMBER_EPSILON:
            raise server_error(
                code="SCHEDULE_CAPACITY_PER_SHIFT_INVALID",
                message="capacity_per_shift must be greater than 0 for all scheduled processes.",
                details={"order_no": order_no, "process_code": process_code},
            )

        allocations: list[dict[str, Any]] = []
        slot_cursor = max(0, int(first_slot))
        last_slot = slot_cursor
        loop_guard = 0
        while remaining_qty > SCHEDULE_NUMBER_EPSILON:
            if loop_guard >= SCHEDULE_SLOT_SEARCH_GUARD:
                raise server_error(
                    code="SCHEDULE_ALLOCATION_GUARD_EXCEEDED",
                    message="Unable to allocate process workload within slot guard limit.",
                    details={
                        "order_no": order_no,
                        "process_code": process_code,
                        "remaining_qty": remaining_qty,
                    },
                )
            slot_index, calendar_date, shift_code = self._next_working_slot(
                start_slot=slot_cursor,
                planning_rules=planning_rules,
                day_mode_cache=day_mode_cache,
                allow_exact_start_slot=slot_cursor == max(0, int(first_slot)),
            )
            candidate_context, available_capacity = self._select_candidate_context_for_slot(
                process_context=process_context,
                slot_index=slot_index,
                calendar_date=calendar_date,
                used_capacity_by_slot=used_capacity_by_slot,
                capacity_resolver=capacity_resolver,
            )
            if candidate_context is None or available_capacity <= SCHEDULE_NUMBER_EPSILON:
                slot_cursor = slot_index + 1
                loop_guard += 1
                continue
            key = self._capacity_usage_key(
                process_context=candidate_context,
                slot_index=slot_index,
                calendar_date=calendar_date,
                capacity_resolver=capacity_resolver,
            )
            used_capacity = _to_number(used_capacity_by_slot.get(key), 0)

            planned_qty = min(remaining_qty, available_capacity)
            used_capacity_by_slot[key] = used_capacity + planned_qty
            allocations.append(
                {
                    "calendar_date": calendar_date,
                    "shift_code": shift_code,
                    "workshop_code": str(candidate_context.get("workshop_code") or "").strip().upper(),
                    "line_code": str(candidate_context.get("line_code") or "").strip().upper(),
                    "plan_qty": planned_qty,
                }
            )
            remaining_qty -= planned_qty
            last_slot = slot_index
            slot_cursor = slot_index if remaining_qty > SCHEDULE_NUMBER_EPSILON else slot_index + 1
            loop_guard += 1

        return allocations, last_slot

    def _build_schedule_order_summary_map_from_generated_tasks(
        self,
        *,
        tasks: list[tuple[Any, ...]],
        order_rows: list[dict[str, Any]],
    ) -> dict[str, dict[str, Any]]:
        order_row_by_no = {
            str(row.get("production_order_no") or "").strip(): row for row in order_rows
        }
        out: dict[str, dict[str, Any]] = {}
        for task in tasks:
            if len(task) < 10:
                continue
            order_no = str(task[2] or "").strip()
            if not order_no:
                continue
            base_row = order_row_by_no.get(order_no) or {}
            current = out.get(order_no)
            if current is None:
                current = {
                    "order_no": order_no,
                    "material_code": str(base_row.get("material_code") or "").strip(),
                    "material_name": str(
                        base_row.get("material_name") or base_row.get("material_code") or ""
                    ).strip(),
                    "order_qty": _to_number(base_row.get("production_qty"), 0),
                    "planned_ratio": 0.0,
                    "planned_qty": 0.0,
                }
                out[order_no] = current
            current["planned_qty"] += _to_number(task[9], 0)
        for row in out.values():
            order_qty = _to_number(row.get("order_qty"), 0)
            row["planned_ratio"] = (
                max(0.0, min(1.0, _to_number(row.get("planned_qty"), 0) / order_qty))
                if order_qty > 0
                else 0.0
            )
        return out

    def _build_generated_schedule_material_shortages(
        self,
        *,
        order_summary_map: dict[str, dict[str, Any]],
    ) -> dict[str, Any]:
        material_map = self._build_schedule_material_usage_map(
            order_summary_map,
            material_rows=self._list_expanded_schedule_material_rows(),
        )
        impacted_order_nos: set[str] = set()
        items: list[dict[str, Any]] = []
        for material_code in sorted(material_map):
            current = material_map[material_code]
            estimated_inventory_qty = _to_number(current.get("estimated_inventory_qty"), 0)
            if estimated_inventory_qty >= 0:
                continue
            order_nos = sorted(
                {
                    str(order_no).strip()
                    for order_no in current.get("order_nos", [])
                    if str(order_no).strip()
                }
            )
            impacted_order_nos.update(order_nos)
            items.append(
                {
                    "material_code": material_code,
                    "material_name": current.get("material_name"),
                    "supply_type_name": current.get("supply_type_name"),
                    "inventory_qty": _to_number(current.get("inventory_qty"), 0),
                    "planned_consume_qty": _to_number(current.get("planned_consume_qty"), 0),
                    "estimated_inventory_qty": estimated_inventory_qty,
                    "shortage_qty": abs(estimated_inventory_qty),
                    "order_nos": order_nos,
                }
            )
        items.sort(
            key=lambda item: (
                -_to_number(item.get("shortage_qty"), 0),
                str(item.get("material_code") or ""),
            )
        )
        return {
            "summary": {
                "shortage_material_count": len(items),
                "impacted_order_count": len(impacted_order_nos),
            },
            "items": items,
        }

    def _sort_schedule_candidates(
        self,
        *,
        strategy_code: str,
        candidates: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        def updated_rank(text: str) -> float:
            safe_text = str(text or "").strip()
            if not safe_text:
                return 0.0
            try:
                return datetime.fromisoformat(safe_text.replace("Z", "+00:00")).timestamp()
            except ValueError:
                return 0.0

        if strategy_code == "KEY_ORDER_FIRST":
            return sorted(
                candidates,
                key=lambda item: (
                    -int(item.get("frozen_flag") or 0),
                    -int(item.get("lock_flag") or 0),
                    _normalize_priority_level(item.get("priority_level"), PRIORITY_LEVEL_MAX),
                    int(item.get("base_first_task_no") or 10**9),
                    item.get("start_date"),
                    item.get("due_date"),
                    -updated_rank(str(item.get("updated_at") or "")),
                    str(item.get("order_no") or ""),
                ),
            )
        if strategy_code == "MAX_CAPACITY_FIRST":
            return sorted(
                candidates,
                key=lambda item: (
                    -int(item.get("frozen_flag") or 0),
                    -int(item.get("lock_flag") or 0),
                    _normalize_priority_level(item.get("priority_level"), PRIORITY_LEVEL_MAX),
                    -_to_number(item.get("total_capacity_per_shift"), 0),
                    -_to_number(item.get("min_capacity_per_shift"), 0),
                    int(item.get("base_first_task_no") or 10**9),
                    item.get("start_date"),
                    item.get("due_date"),
                    str(item.get("order_no") or ""),
                ),
            )
        if strategy_code == "MIN_DELAY_FIRST":
            return sorted(
                candidates,
                key=lambda item: (
                    -int(item.get("frozen_flag") or 0),
                    -int(item.get("lock_flag") or 0),
                    _normalize_priority_level(item.get("priority_level"), PRIORITY_LEVEL_MAX),
                    int(_to_number(item.get("slack_days"), 0)),
                    item.get("due_date"),
                    item.get("start_date"),
                    int(item.get("base_first_task_no") or 10**9),
                    str(item.get("order_no") or ""),
                ),
            )
        raise bad_request(
            code="SCHEDULE_STRATEGY_INVALID",
            message="strategy_code is invalid.",
            details={"strategy_code": strategy_code},
        )

    def _first_available_slot_for_process(
        self,
        *,
        process_context: dict[str, Any],
        start_slot: int,
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
        used_capacity_by_slot: dict[tuple[Any, ...], float],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> int:
        process_code = str(process_context.get("process_code") or "").strip().upper()
        default_capacity_per_shift = _to_number(process_context.get("default_capacity_per_shift"), 0)
        if default_capacity_per_shift <= SCHEDULE_NUMBER_EPSILON:
            raise server_error(
                code="SCHEDULE_CAPACITY_PER_SHIFT_INVALID",
                message="capacity_per_shift must be greater than 0 for all scheduled processes.",
                details={"process_code": process_code},
            )

        slot_cursor = max(0, int(start_slot))
        for _ in range(SCHEDULE_SLOT_SEARCH_GUARD):
            slot_index, calendar_date, _ = self._next_working_slot(
                start_slot=slot_cursor,
                planning_rules=planning_rules,
                day_mode_cache=day_mode_cache,
                allow_exact_start_slot=slot_cursor == max(0, int(start_slot)),
            )
            capacity_per_shift = self._resolve_effective_capacity_per_shift(
                process_context=process_context,
                calendar_date=calendar_date,
                capacity_resolver=capacity_resolver,
            )
            key = self._capacity_usage_key(
                process_context=process_context,
                slot_index=slot_index,
                calendar_date=calendar_date,
                capacity_resolver=capacity_resolver,
            )
            used_capacity = _to_number(used_capacity_by_slot.get(key), 0)
            available_capacity = max(0.0, capacity_per_shift - used_capacity)
            if available_capacity > SCHEDULE_NUMBER_EPSILON:
                return slot_index
            slot_cursor = slot_index + 1
        raise server_error(
            code="SCHEDULE_WORKING_SLOT_NOT_FOUND",
            message="Cannot find available slot for process.",
            details={"process_code": process_code, "start_slot": start_slot},
        )

    def _candidate_runtime_metrics(
        self,
        *,
        candidate: dict[str, Any],
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
        used_capacity_by_slot: dict[tuple[Any, ...], float],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> dict[str, int]:
        process_contexts = candidate["process_contexts"]
        if len(process_contexts) == 0:
            raise server_error(
                code="SCHEDULE_PROCESS_CONTEXTS_EMPTY",
                message="No process contexts available for schedule candidate.",
                details={"order_no": candidate.get("order_no")},
            )

        first_context = process_contexts[0]
        first_process_code = str(first_context.get("process_code") or "").strip().upper()
        first_slot = int(candidate.get("start_slot") or 0)
        base_process_first_slot = candidate.get("base_process_first_slot", {})
        first_process_base_slot = base_process_first_slot.get(first_process_code)
        if (
            (int(candidate.get("lock_flag") or 0) == 1 or int(candidate.get("frozen_flag") or 0) == 1)
            and first_process_base_slot is not None
        ):
            first_slot = max(first_slot, int(first_process_base_slot))

        first_ready_slot = self._first_available_slot_for_process(
            process_context=first_context,
            start_slot=first_slot,
            planning_rules=planning_rules,
            day_mode_cache=day_mode_cache,
            used_capacity_by_slot=used_capacity_by_slot,
            capacity_resolver=capacity_resolver,
        )
        required_shifts = max(1, int(_to_number(candidate.get("required_shifts"), 1)))
        projected_finish_slot = first_ready_slot + required_shifts - 1
        due_date = candidate.get("due_date")
        if not isinstance(due_date, date):
            due_date = _parse_date_or_today(due_date)
        due_slot = _slot_index_for(due_date, "NIGHT")
        projected_lateness_slots = projected_finish_slot - due_slot
        projected_slack_slots = due_slot - projected_finish_slot
        return {
            "first_ready_slot": first_ready_slot,
            "projected_finish_slot": projected_finish_slot,
            "due_slot": due_slot,
            "projected_lateness_slots": projected_lateness_slots,
            "projected_slack_slots": projected_slack_slots,
        }

    def _select_next_candidate_index(
        self,
        *,
        strategy_code: str,
        candidates: list[dict[str, Any]],
        planning_rules: dict[str, Any],
        day_mode_cache: dict[str, str],
        used_capacity_by_slot: dict[tuple[Any, ...], float],
        capacity_resolver: dict[str, dict[tuple[str, str, str, str, str], float]],
    ) -> int:
        if len(candidates) == 0:
            raise server_error(
                code="SCHEDULE_CANDIDATE_EMPTY",
                message="No schedule candidates available.",
            )

        best_index = 0
        best_key: tuple[Any, ...] | None = None
        for index, candidate in enumerate(candidates):
            metrics = self._candidate_runtime_metrics(
                candidate=candidate,
                planning_rules=planning_rules,
                day_mode_cache=day_mode_cache,
                used_capacity_by_slot=used_capacity_by_slot,
                capacity_resolver=capacity_resolver,
            )
            late_slots = int(metrics["projected_lateness_slots"])
            late_rank = 0 if late_slots > 0 else 1
            late_severity = -max(late_slots, 0)
            business_head = (
                -int(candidate.get("frozen_flag") or 0),
                -int(candidate.get("lock_flag") or 0),
                _normalize_priority_level(candidate.get("priority_level"), PRIORITY_LEVEL_MAX),
            )
            base_tail = (
                int(candidate.get("base_first_task_no") or 10**9),
                str(candidate.get("order_no") or ""),
            )

            if strategy_code == "KEY_ORDER_FIRST":
                key = (
                    *business_head,
                    late_rank,
                    late_severity,
                    int(metrics["due_slot"]),
                    int(metrics["first_ready_slot"]),
                    *base_tail,
                )
            elif strategy_code == "MAX_CAPACITY_FIRST":
                key = (
                    *business_head,
                    late_rank,
                    late_severity,
                    -_to_number(candidate.get("total_capacity_per_shift"), 0),
                    int(_to_number(candidate.get("required_shifts"), 0)),
                    int(metrics["due_slot"]),
                    int(metrics["first_ready_slot"]),
                    *base_tail,
                )
            elif strategy_code == "MIN_DELAY_FIRST":
                key = (
                    *business_head,
                    late_rank,
                    late_severity,
                    int(metrics["projected_slack_slots"]),
                    int(metrics["due_slot"]),
                    int(metrics["first_ready_slot"]),
                    *base_tail,
                )
            else:
                raise bad_request(
                    code="SCHEDULE_STRATEGY_INVALID",
                    message="strategy_code is invalid.",
                    details={"strategy_code": strategy_code},
                )

            if best_key is None or key < best_key:
                best_key = key
                best_index = index

        return best_index

    def _is_cn_statutory_holiday(self, date_text: str) -> bool:
        return date_text in CN_STATUTORY_HOLIDAY_DATE_SET

    def _get_rules_row(self) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT
                singleton_key,
                horizon_start_date,
                horizon_days,
                skip_statutory_holidays,
                weekend_rest_mode,
                date_shift_mode_by_date_json,
                updated_at
            FROM schedule_calendar_rules
            WHERE singleton_key = ?
            """,
            (RULES_SINGLETON_KEY,),
        )
        if row is not None:
            return row
        with transaction(self.connection):
            self.connection.execute(
                """
                INSERT INTO schedule_calendar_rules (
                    singleton_key,
                    horizon_start_date,
                    horizon_days,
                    skip_statutory_holidays,
                    weekend_rest_mode,
                    date_shift_mode_by_date_json,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    RULES_SINGLETON_KEY,
                    _today_text(),
                    31,
                    0,
                    "DOUBLE",
                    "{}",
                    utc_now(),
                ),
            )
        return self._get_rules_row()

    def _get_simulation_state(self) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT
                simulation_state.singleton_key,
                simulation_state.current_date AS current_date,
                simulation_state.updated_at
            FROM simulation_state
            WHERE singleton_key = ?
            """,
            (RULES_SINGLETON_KEY,),
        )
        if row is not None:
            return row
        with transaction(self.connection):
            self.connection.execute(
                """
                INSERT INTO simulation_state (
                    singleton_key,
                    current_date,
                    updated_at
                ) VALUES (?, ?, ?)
                """,
                (RULES_SINGLETON_KEY, _today_text(), utc_now()),
            )
        return self._get_simulation_state()

    def _ensure_masterdata_seeded(self) -> None:
        topology_count = fetch_one(
            self.connection,
            "SELECT COUNT(1) AS total FROM masterdata_line_topology",
        )
        route_count = fetch_one(
            self.connection,
            "SELECT COUNT(1) AS total FROM masterdata_process_routes",
        )
        skeleton_count = fetch_one(
            self.connection,
            "SELECT COUNT(1) AS total FROM masterdata_line_skeletons",
        )
        need_topology_seed = int(topology_count["total"]) == 0
        need_route_seed = int(route_count["total"]) == 0
        need_skeleton_seed = int(skeleton_count["total"]) == 0
        if not need_topology_seed and not need_route_seed and not need_skeleton_seed:
            return

        upstream_route_rows = (
            self.masterdata_gateway.fetch_process_routes()
            if need_topology_seed or need_route_seed
            else []
        )
        upstream_capability_rows = (
            self.masterdata_gateway.fetch_equipment_process_capabilities()
            if need_topology_seed
            else []
        )
        topology_rows = (
            self._build_line_topology_seed_rows(
                upstream_route_rows,
                upstream_capability_rows,
            )
            if need_topology_seed
            else []
        )
        route_rows = self._build_route_seed_rows(upstream_route_rows) if need_route_seed else self._list_route_rows()
        if len(route_rows) == 0:
            raise server_error(
                code="MASTERDATA_PROCESS_ROUTES_EMPTY",
                message="Upstream process route seed returned no rows.",
            )
        if need_topology_seed and len(topology_rows) == 0:
            raise server_error(
                code="MASTERDATA_LINE_TOPOLOGY_EMPTY",
                message="Upstream line topology seed returned no rows.",
            )
        base_topology_rows = topology_rows if need_topology_seed else self._list_line_topology_rows()
        normalized_topology_rows = (
            self._build_default_angio_line_topology_rows(route_rows, base_topology_rows)
            if need_topology_seed or need_skeleton_seed
            else base_topology_rows
        )
        skeleton_rows = (
            self._build_line_skeleton_seed_rows(base_topology_rows)
            if need_skeleton_seed
            else []
        )
        updated_at = utc_now()
        with transaction(self.connection):
            if len(route_rows) > 0 and need_route_seed:
                self.connection.executemany(
                    """
                    INSERT INTO masterdata_process_routes (
                        product_code,
                        sequence_no,
                        process_code,
                        process_name_cn,
                        dependency_type,
                        route_no,
                        route_name_cn,
                        product_name_cn,
                        is_final_process,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            row["product_code"],
                            row["sequence_no"],
                            row["process_code"],
                            row["process_name_cn"],
                            row["dependency_type"],
                            row["route_no"],
                            row["route_name_cn"],
                            row["product_name_cn"],
                            int(_to_number(row.get("is_final_process"), 0)),
                            updated_at,
                        )
                        for row in route_rows
                    ],
                )
            if need_skeleton_seed:
                self.connection.execute("DELETE FROM masterdata_line_skeletons")
                self.connection.executemany(
                    """
                    INSERT INTO masterdata_line_skeletons (
                        company_code,
                        workshop_code,
                        workshop_name,
                        line_code,
                        line_name,
                        enabled_flag,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            row["company_code"],
                            row["workshop_code"],
                            row["workshop_name"],
                            row["line_code"],
                            row["line_name"],
                            row["enabled_flag"],
                            updated_at,
                        )
                        for row in skeleton_rows
                    ],
                )
            if need_topology_seed or need_skeleton_seed:
                self.connection.execute("DELETE FROM masterdata_line_topology")
                self.connection.executemany(
                    """
                    INSERT INTO masterdata_line_topology (
                        company_code,
                        workshop_code,
                        workshop_name,
                        line_code,
                        line_name,
                        process_code,
                        capacity_per_shift,
                        required_workers,
                        required_machines,
                        enabled_flag,
                        updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            row["company_code"],
                            row["workshop_code"],
                            row["workshop_name"],
                            row["line_code"],
                            row["line_name"],
                            row["process_code"],
                            row["capacity_per_shift"],
                            row["required_workers"],
                            row["required_machines"],
                            row["enabled_flag"],
                            updated_at,
                        )
                        for row in normalized_topology_rows
                    ],
                )

    def _build_route_seed_rows(
        self,
        rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        seen: set[tuple[str, int, str]] = set()
        for row in rows:
            product_code = str(row.get("product_code") or "").strip().upper()
            process_code = str(row.get("process_code") or "").strip().upper()
            sequence_no = int(_to_number(row.get("sequence_no"), 0))
            if not product_code or not process_code or sequence_no <= 0:
                continue
            key = (product_code, sequence_no, process_code)
            if key in seen:
                continue
            seen.add(key)
            product_name_cn = str(row.get("product_name_cn") or product_code).strip() or product_code
            route_no = str(row.get("route_no") or f"ROUTE-{product_code}").strip() or f"ROUTE-{product_code}"
            route_name_cn = str(row.get("route_name_cn") or product_name_cn).strip() or product_name_cn
            out.append(
                {
                    "product_code": product_code,
                    "sequence_no": sequence_no,
                    "process_code": process_code,
                    "process_name_cn": str(row.get("process_name_cn") or process_code).strip() or process_code,
                    "dependency_type": str(row.get("dependency_type") or "FS").strip().upper() or "FS",
                    "route_no": route_no,
                    "route_name_cn": route_name_cn,
                    "product_name_cn": product_name_cn,
                }
            )
        out.sort(key=lambda item: (str(item["product_code"]), int(item["sequence_no"])))
        return self._apply_default_final_process_flags(out)

    def _apply_default_final_process_flags(
        self,
        rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        if not rows:
            return []
        max_sequence_by_product: dict[str, int] = {}
        marked_count_by_product: dict[str, int] = defaultdict(int)
        for row in rows:
            product_code = str(row.get("product_code") or "").strip().upper()
            sequence_no = int(_to_number(row.get("sequence_no"), 0))
            if not product_code or sequence_no <= 0:
                continue
            max_sequence_by_product[product_code] = max(
                max_sequence_by_product.get(product_code, 0),
                sequence_no,
            )
            if int(_to_number(row.get("is_final_process"), 0)) == 1:
                marked_count_by_product[product_code] += 1

        normalized_rows: list[dict[str, Any]] = []
        for row in rows:
            product_code = str(row.get("product_code") or "").strip().upper()
            sequence_no = int(_to_number(row.get("sequence_no"), 0))
            has_marked = marked_count_by_product.get(product_code, 0) > 0
            is_final_process = (
                int(_to_number(row.get("is_final_process"), 0)) == 1
                if has_marked
                else (
                    bool(product_code)
                    and sequence_no > 0
                    and sequence_no == max_sequence_by_product.get(product_code, -1)
                )
            )
            normalized_rows.append(
                {
                    **row,
                    "is_final_process": 1 if is_final_process else 0,
                }
            )
        return normalized_rows

    def _build_line_topology_seed_rows(
        self,
        route_rows: list[dict[str, Any]],
        capability_rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        route_meta_by_process: dict[str, dict[str, Any]] = {}
        for row in route_rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            current = route_meta_by_process.get(process_code)
            next_meta = {
                "process_name_cn": str(row.get("process_name_cn") or process_code).strip() or process_code,
                "capacity_per_shift": _to_number(row.get("capacity_per_shift"), 0),
                "required_workers": int(_to_number(row.get("required_manpower_per_group"), 0)),
                "required_machines": int(_to_number(row.get("required_equipment_count"), 0)),
            }
            if current is None:
                route_meta_by_process[process_code] = next_meta
                continue
            current["capacity_per_shift"] = max(
                _to_number(current.get("capacity_per_shift"), 0),
                next_meta["capacity_per_shift"],
            )
            current["required_workers"] = max(
                int(_to_number(current.get("required_workers"), 0)),
                next_meta["required_workers"],
            )
            current["required_machines"] = max(
                int(_to_number(current.get("required_machines"), 0)),
                next_meta["required_machines"],
            )
            if not current.get("process_name_cn"):
                current["process_name_cn"] = next_meta["process_name_cn"]

        out: list[dict[str, Any]] = []
        seen: set[tuple[str, str, str, str]] = set()
        for row in capability_rows:
            company_code = str(row.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
            workshop_code = str(row.get("workshop_code") or "").strip().upper()
            line_code = str(row.get("line_code") or "").strip().upper()
            process_code = str(row.get("process_code") or "").strip().upper()
            if not workshop_code or not line_code or not process_code:
                continue
            meta = route_meta_by_process.get(process_code)
            if meta is None:
                raise server_error(
                    code="MASTERDATA_TOPOLOGY_PROCESS_META_MISSING",
                    message="Missing route metadata for process in line topology seed.",
                    details={"process_code": process_code},
                )
            key = (company_code, workshop_code, line_code, process_code)
            if key in seen:
                continue
            seen.add(key)
            out.append(
                {
                    "company_code": company_code,
                    "workshop_code": workshop_code,
                    "workshop_name": workshop_code,
                    "line_code": line_code,
                    "line_name": str(row.get("line_name") or line_code).strip() or line_code,
                    "process_code": process_code,
                    "capacity_per_shift": _to_number(meta.get("capacity_per_shift"), 0),
                    "required_workers": int(_to_number(meta.get("required_workers"), 0)),
                    "required_machines": int(_to_number(meta.get("required_machines"), 0)),
                    "enabled_flag": int(_to_number(row.get("enabled_flag"), 0)),
                }
            )
        out.sort(
            key=lambda item: (
                str(item["workshop_code"]),
                str(item["line_code"]),
                str(item["process_code"]),
            )
        )
        return out

    def _build_line_skeleton_seed_rows(
        self,
        topology_rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        seen: set[tuple[str, str, str]] = set()
        for row in topology_rows:
            company_code = str(row.get("company_code") or "COMPANY-MAIN").strip().upper() or "COMPANY-MAIN"
            workshop_code = str(row.get("workshop_code") or "").strip().upper()
            line_code = str(row.get("line_code") or "").strip().upper()
            if not workshop_code or not line_code:
                continue
            key = (company_code, workshop_code, line_code)
            if key in seen:
                continue
            seen.add(key)
            out.append(
                {
                    "company_code": company_code,
                    "workshop_code": workshop_code,
                    "workshop_name": str(row.get("workshop_name") or workshop_code).strip() or workshop_code,
                    "line_code": line_code,
                    "line_name": str(row.get("line_name") or line_code).strip() or line_code,
                    "enabled_flag": 1,
                }
            )
        default_key = ("COMPANY-MAIN", DEFAULT_ANGIO_WORKSHOP_CODE, DEFAULT_ANGIO_LINE_CODE)
        if default_key not in seen:
            out.append(
                {
                    "company_code": "COMPANY-MAIN",
                    "workshop_code": DEFAULT_ANGIO_WORKSHOP_CODE,
                    "workshop_name": DEFAULT_ANGIO_WORKSHOP_NAME,
                    "line_code": DEFAULT_ANGIO_LINE_CODE,
                    "line_name": DEFAULT_ANGIO_LINE_NAME,
                    "enabled_flag": 1,
                }
            )
        out.sort(
            key=lambda item: (
                str(item["workshop_code"]),
                str(item["line_code"]),
            )
        )
        return out

    def _build_default_angio_line_topology_rows(
        self,
        route_rows: list[dict[str, Any]],
        topology_rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        angio_route_rows = [
            row
            for row in route_rows
            if str(row.get("product_code") or "").strip().upper() in ANGIO_CATHETER_PRODUCT_CODES
        ]
        if len(angio_route_rows) == 0:
            raise server_error(
                code="ANGIO_ROUTE_ROWS_EMPTY",
                message="No angio catheter process routes are available for default topology generation.",
            )
        ordered_process_codes: list[str] = []
        meta_by_process: dict[str, dict[str, Any]] = {}
        for row in angio_route_rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            if process_code not in meta_by_process:
                ordered_process_codes.append(process_code)
                meta_by_process[process_code] = {
                    "process_name_cn": str(row.get("process_name_cn") or process_code).strip() or process_code,
                    "capacity_per_shift": _to_number(row.get("capacity_per_shift"), 0),
                    "required_workers": int(_to_number(row.get("required_manpower_per_group"), 0)),
                    "required_machines": int(_to_number(row.get("required_equipment_count"), 0)),
                    "enabled_flag": 1 if int(_to_number(row.get("enabled_flag"), 1)) == 1 else 0,
                }
                continue
            current = meta_by_process[process_code]
            current["capacity_per_shift"] = max(
                _to_number(current.get("capacity_per_shift"), 0),
                _to_number(row.get("capacity_per_shift"), 0),
            )
            current["required_workers"] = max(
                int(_to_number(current.get("required_workers"), 0)),
                int(_to_number(row.get("required_manpower_per_group"), 0)),
            )
            current["required_machines"] = max(
                int(_to_number(current.get("required_machines"), 0)),
                int(_to_number(row.get("required_equipment_count"), 0)),
            )
            if not current.get("process_name_cn"):
                current["process_name_cn"] = (
                    str(row.get("process_name_cn") or process_code).strip() or process_code
                )
            if int(_to_number(row.get("enabled_flag"), 1)) == 1:
                current["enabled_flag"] = 1
        for row in topology_rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if process_code not in meta_by_process:
                continue
            current = meta_by_process[process_code]
            current["capacity_per_shift"] = max(
                _to_number(current.get("capacity_per_shift"), 0),
                _to_number(row.get("capacity_per_shift"), 0),
            )
            current["required_workers"] = max(
                int(_to_number(current.get("required_workers"), 0)),
                int(_to_number(row.get("required_workers"), 0)),
            )
            current["required_machines"] = max(
                int(_to_number(current.get("required_machines"), 0)),
                int(_to_number(row.get("required_machines"), 0)),
            )
            if not current.get("process_name_cn"):
                current["process_name_cn"] = (
                    str(row.get("process_name_cn") or process_code).strip() or process_code
                )
            if int(_to_number(row.get("enabled_flag"), 0)) == 1:
                current["enabled_flag"] = 1
        return [
            {
                "company_code": "COMPANY-MAIN",
                "workshop_code": DEFAULT_ANGIO_WORKSHOP_CODE,
                "workshop_name": DEFAULT_ANGIO_WORKSHOP_NAME,
                "line_code": DEFAULT_ANGIO_LINE_CODE,
                "line_name": DEFAULT_ANGIO_LINE_NAME,
                "process_code": process_code,
                "capacity_per_shift": _to_number(meta_by_process[process_code].get("capacity_per_shift"), 0),
                "required_workers": int(_to_number(meta_by_process[process_code].get("required_workers"), 0)),
                "required_machines": int(_to_number(meta_by_process[process_code].get("required_machines"), 0)),
                "enabled_flag": 1 if int(_to_number(meta_by_process[process_code].get("enabled_flag"), 0)) == 1 else 0,
            }
            for process_code in ordered_process_codes
        ]

    def _resolve_final_process_meta_by_product(
        self,
        product_codes: set[str],
    ) -> dict[str, dict[str, str]]:
        normalized_product_codes = {
            str(product_code or "").strip().upper()
            for product_code in product_codes
            if str(product_code or "").strip()
        }
        if not normalized_product_codes:
            return {}
        placeholders = ",".join("?" for _ in normalized_product_codes)
        rows = fetch_all(
            self.connection,
            f"""
            SELECT
                product_code,
                sequence_no,
                process_code,
                process_name_cn,
                COALESCE(is_final_process, 0) AS is_final_process
            FROM masterdata_process_routes
            WHERE product_code IN ({placeholders})
            ORDER BY product_code ASC, sequence_no ASC
            """,
            tuple(sorted(normalized_product_codes)),
        )
        rows_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            product_code = str(row.get("product_code") or "").strip().upper()
            process_code = str(row.get("process_code") or "").strip().upper()
            sequence_no = int(_to_number(row.get("sequence_no"), 0))
            if not product_code or not process_code or sequence_no <= 0:
                continue
            rows_by_product[product_code].append(
                {
                    "sequence_no": sequence_no,
                    "process_code": process_code,
                    "process_name_cn": str(row.get("process_name_cn") or process_code).strip() or process_code,
                    "is_final_process": 1 if int(_to_number(row.get("is_final_process"), 0)) == 1 else 0,
                }
            )

        out: dict[str, dict[str, str]] = {}
        for product_code, route_rows in rows_by_product.items():
            ordered_rows = sorted(route_rows, key=lambda item: int(item["sequence_no"]))
            marked_rows = [row for row in ordered_rows if int(row.get("is_final_process") or 0) == 1]
            if len(marked_rows) > 1:
                raise server_error(
                    code="ROUTE_FINAL_PROCESS_DUPLICATED",
                    message="Route has more than one final process step.",
                    details={"product_code": product_code},
                )
            target_row = marked_rows[0] if marked_rows else ordered_rows[-1]
            process_code = str(target_row.get("process_code") or "").strip().upper()
            if not process_code:
                continue
            out[product_code] = {
                "process_code": process_code,
                "process_name_cn": str(target_row.get("process_name_cn") or process_code).strip() or process_code,
            }
        return out

    def _is_workshop_manager(self, user: dict[str, Any] | None) -> bool:
        role_code = str((user or {}).get("role_code") or "").strip().upper()
        return role_code == ROLE_WORKSHOP_MANAGER

    def _current_user_role_code(self, user: dict[str, Any] | None) -> str:
        return str((user or {}).get("role_code") or "").strip().upper()

    def _resolve_manager_user_id(self, user: dict[str, Any] | None) -> str | None:
        if not self._is_workshop_manager(user):
            return None
        user_id = str((user or {}).get("user_id") or "").strip()
        if not user_id:
            raise forbidden(
                code="WORKSHOP_MANAGER_USER_ID_REQUIRED",
                message="Current workshop manager user_id is missing.",
            )
        return user_id

    def _resolve_order_summary_scope_manager_user_id(
        self,
        *,
        current_user: dict[str, Any] | None,
        workshop_manager_user_id: str | None,
    ) -> str | None:
        role_code = self._current_user_role_code(current_user)
        requested_user_id = str(workshop_manager_user_id or "").strip()
        if workshop_manager_user_id is not None and not requested_user_id:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_USER_ID_REQUIRED",
                message="workshop_manager_user_id must be non-empty when provided.",
            )
        if role_code == ROLE_WORKSHOP_MANAGER:
            actor_user_id = self._resolve_manager_user_id(current_user)
            assert actor_user_id is not None
            if requested_user_id and requested_user_id != actor_user_id:
                raise forbidden(
                    code="ORDER_SUMMARY_WORKSHOP_MANAGER_FILTER_FORBIDDEN",
                    message="Current workshop manager is not allowed to access another workshop manager scope.",
                    details={
                        "requested_user_id": requested_user_id,
                        "allowed_user_id": actor_user_id,
                    },
                )
            return actor_user_id
        if role_code == ROLE_SCHEDULER:
            if not requested_user_id:
                return None
            return self._validate_order_summary_scheduler_filter_target(requested_user_id)
        raise forbidden(
            code="ORDER_SUMMARY_ROLE_FORBIDDEN",
            message="Current role is not allowed to access order summary.",
            details={"role_code": role_code},
        )

    def _validate_order_summary_scheduler_filter_target(
        self,
        workshop_manager_user_id: str,
    ) -> str:
        normalized_user_id = str(workshop_manager_user_id or "").strip()
        if not normalized_user_id:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_USER_ID_REQUIRED",
                message="workshop_manager_user_id must be non-empty when provided.",
            )
        row = fetch_one(
            self.connection,
            """
            SELECT
                users.user_id,
                users.role_code,
                users.enabled_flag,
                COALESCE(visibility.visible_flag, 1) AS visible_flag,
                COUNT(scope.line_code) AS line_scope_count
            FROM app_users users
            LEFT JOIN masterdata_workshop_manager_visibility visibility
              ON visibility.user_id = users.user_id
            LEFT JOIN app_user_line_scopes scope
              ON scope.user_id = users.user_id
            WHERE users.user_id = ?
            GROUP BY users.user_id, users.role_code, users.enabled_flag, COALESCE(visibility.visible_flag, 1)
            LIMIT 1
            """,
            (normalized_user_id,),
        )
        if row is None:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_NOT_FOUND",
                message="workshop_manager_user_id does not exist.",
                details={"workshop_manager_user_id": normalized_user_id},
            )
        role_code = str(row.get("role_code") or "").strip().upper()
        if role_code != ROLE_WORKSHOP_MANAGER:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_ROLE_INVALID",
                message="workshop_manager_user_id must reference an enabled workshop manager user.",
                details={"workshop_manager_user_id": normalized_user_id},
            )
        if int(row.get("enabled_flag") or 0) != 1:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_DISABLED",
                message="workshop_manager_user_id references a disabled workshop manager user.",
                details={"workshop_manager_user_id": normalized_user_id},
            )
        if int(row.get("visible_flag") or 0) != 1:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_HIDDEN",
                message="workshop_manager_user_id references a hidden workshop manager user.",
                details={"workshop_manager_user_id": normalized_user_id},
            )
        if int(_to_number(row.get("line_scope_count"), 0)) <= 0:
            raise bad_request(
                code="ORDER_SUMMARY_WORKSHOP_MANAGER_LINE_SCOPE_EMPTY",
                message="workshop_manager_user_id must have at least one assigned line scope.",
                details={"workshop_manager_user_id": normalized_user_id},
            )
        return normalized_user_id

    def _list_user_line_scope_rows(self, user_id: str) -> list[dict[str, Any]]:
        normalized_user_id = str(user_id or "").strip()
        if not normalized_user_id:
            return []
        return fetch_all(
            self.connection,
            """
            SELECT
                user_id,
                company_code,
                workshop_code,
                line_code
            FROM app_user_line_scopes
            WHERE user_id = ?
            """,
            (normalized_user_id,),
        )

    def _build_line_scope_key(
        self,
        *,
        company_code: str | None,
        workshop_code: str | None,
        line_code: str | None,
    ) -> tuple[str, str, str]:
        return (
            str(company_code or DEFAULT_COMPANY_CODE).strip().upper() or DEFAULT_COMPANY_CODE,
            str(workshop_code or "").strip().upper(),
            str(line_code or "").strip().upper(),
        )

    def _user_line_scope_key_set(self, user_id: str) -> set[tuple[str, str, str]]:
        rows = self._list_user_line_scope_rows(user_id)
        return {
            self._build_line_scope_key(
                company_code=row.get("company_code"),
                workshop_code=row.get("workshop_code"),
                line_code=row.get("line_code"),
            )
            for row in rows
        }

    def _assert_actor_can_access_line(
        self,
        actor: dict[str, Any] | None,
        *,
        company_code: str,
        workshop_code: str,
        line_code: str,
        missing_user_error_code: str,
        forbidden_error_code: str,
        forbidden_message: str,
    ) -> None:
        if not self._is_workshop_manager(actor):
            return
        actor_user_id = str((actor or {}).get("user_id") or "").strip()
        if not actor_user_id:
            raise bad_request(
                code=missing_user_error_code,
                message="workshop manager actor.user_id is required.",
            )
        scope_key = self._build_line_scope_key(
            company_code=company_code,
            workshop_code=workshop_code,
            line_code=line_code,
        )
        allowed_scope_keys = self._user_line_scope_key_set(actor_user_id)
        if scope_key not in allowed_scope_keys:
            raise forbidden(
                code=forbidden_error_code,
                message=forbidden_message,
                details={
                    "user_id": actor_user_id,
                    "company_code": scope_key[0],
                    "workshop_code": scope_key[1],
                    "line_code": scope_key[2],
                },
            )

    def _list_enabled_workshop_manager_users(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                user_id,
                username,
                display_name,
                role_code,
                enabled_flag
            FROM app_users
            WHERE role_code = ?
              AND enabled_flag = 1
            ORDER BY username ASC, user_id ASC
            """,
            (ROLE_WORKSHOP_MANAGER,),
        )

    def _list_workshop_manager_visibility_rows(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                user_id,
                visible_flag
            FROM masterdata_workshop_manager_visibility
            ORDER BY user_id ASC
            """,
        )

    def _list_workshop_manager_line_scope_rows(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                scope.user_id,
                scope.company_code,
                scope.workshop_code,
                scope.line_code
            FROM app_user_line_scopes scope
            JOIN app_users users
              ON users.user_id = scope.user_id
            WHERE users.role_code = ?
              AND users.enabled_flag = 1
            ORDER BY scope.user_id ASC, scope.company_code ASC, scope.workshop_code ASC, scope.line_code ASC
            """,
            (ROLE_WORKSHOP_MANAGER,),
        )

    def _list_line_skeleton_rows(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                company_code,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                enabled_flag
            FROM masterdata_line_skeletons
            ORDER BY workshop_code ASC, line_code ASC
            """,
        )

    def _list_line_topology_rows(self) -> list[dict[str, Any]]:
        return fetch_all(
            self.connection,
            """
            SELECT
                company_code,
                workshop_code,
                workshop_name,
                line_code,
                line_name,
                process_code,
                capacity_per_shift,
                required_workers,
                required_machines,
                enabled_flag
            FROM masterdata_line_topology
            ORDER BY workshop_code ASC, line_code ASC, process_code ASC
            """,
        )

    def _derive_line_topology_rows(self) -> list[dict[str, Any]]:
        rows = fetch_all(
            self.connection,
            """
            SELECT
                workshop_code,
                COALESCE(workshop_name, workshop_code) AS workshop_name,
                line_code,
                COALESCE(line_name, line_code) AS line_name,
                process_code,
                MAX(capacity_qty) AS capacity_per_shift
            FROM capacity_bindings
            WHERE TRIM(COALESCE(workshop_code, '')) <> ''
              AND TRIM(COALESCE(line_code, '')) <> ''
              AND TRIM(COALESCE(process_code, '')) <> ''
            GROUP BY workshop_code, workshop_name, line_code, line_name, process_code
            ORDER BY workshop_code ASC, line_code ASC, process_code ASC
            """,
        )
        return [
            {
                "company_code": "COMPANY-MAIN",
                "workshop_code": str(row["workshop_code"]).strip().upper(),
                "workshop_name": str(row.get("workshop_name") or row["workshop_code"]).strip()
                or str(row["workshop_code"]).strip().upper(),
                "line_code": str(row["line_code"]).strip().upper(),
                "line_name": str(row.get("line_name") or row["line_code"]).strip()
                or str(row["line_code"]).strip().upper(),
                "process_code": str(row["process_code"]).strip().upper(),
                "capacity_per_shift": _to_number(row.get("capacity_per_shift"), 800),
                "required_workers": 4,
                "required_machines": 1,
                "enabled_flag": 1,
            }
            for row in rows
        ]

    def _first_topology_by_process(
        self,
        rows: list[dict[str, Any]],
    ) -> dict[str, dict[str, Any]]:
        out: dict[str, dict[str, Any]] = {}
        for row in rows:
            process_code = str(row.get("process_code") or "").strip().upper()
            if process_code and process_code not in out:
                out[process_code] = row
        return out

    def _list_route_rows(self) -> list[dict[str, Any]]:
        rows = fetch_all(
            self.connection,
            """
            SELECT
                product_code,
                sequence_no,
                process_code,
                process_name_cn,
                dependency_type,
                route_no,
                route_name_cn,
                product_name_cn,
                COALESCE(is_final_process, 0) AS is_final_process
            FROM masterdata_process_routes
            ORDER BY product_code ASC, sequence_no ASC
            """,
        )
        return [
            {
                "id": _route_row_id(str(row["product_code"]), int(row["sequence_no"])),
                **row,
            }
            for row in rows
        ]

    def _derive_route_rows(self) -> list[dict[str, Any]]:
        rows = fetch_all(
            self.connection,
            """
            SELECT
                po.material_code AS product_code,
                COALESCE(po.material_name, po.material_code) AS product_name_cn,
                cb.process_code,
                COALESCE(cb.process_name, cb.process_code) AS process_name_cn,
                cb.workshop_code,
                cb.line_code
            FROM production_orders po
            JOIN capacity_bindings cb
              ON cb.production_order_no = po.production_order_no
            WHERE TRIM(COALESCE(po.material_code, '')) <> ''
              AND TRIM(COALESCE(cb.process_code, '')) <> ''
            ORDER BY po.material_code ASC, cb.process_code ASC, cb.workshop_code ASC, cb.line_code ASC
            """,
        )
        seen: set[tuple[str, str]] = set()
        sequence_by_product: dict[str, int] = defaultdict(int)
        out: list[dict[str, Any]] = []
        for row in rows:
            product_code = str(row["product_code"]).strip().upper()
            process_code = str(row["process_code"]).strip().upper()
            key = (product_code, process_code)
            if key in seen:
                continue
            seen.add(key)
            sequence_by_product[product_code] += 1
            product_name = str(row.get("product_name_cn") or product_code).strip() or product_code
            out.append(
                {
                    "id": _route_row_id(product_code, sequence_by_product[product_code]),
                    "product_code": product_code,
                    "sequence_no": sequence_by_product[product_code],
                    "process_code": process_code,
                    "process_name_cn": str(row.get("process_name_cn") or process_code).strip()
                    or process_code,
                    "dependency_type": "FS",
                    "route_no": f"ROUTE-{product_code}",
                    "route_name_cn": product_name,
                    "product_name_cn": product_name,
                }
            )
        return self._apply_default_final_process_flags(out)

    def _group_routes_by_product(
        self,
        rows: list[dict[str, Any]],
    ) -> dict[str, list[dict[str, Any]]]:
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            grouped[str(row["product_code"])].append(row)
        for product_code in grouped:
            grouped[product_code].sort(key=lambda item: int(item.get("sequence_no") or 0))
        return grouped

    def _replace_process_routes(
        self,
        product_code: str,
        steps: Any,
        *,
        source_product_code: str | None,
    ) -> None:
        normalized_steps = steps if isinstance(steps, list) else []
        if len(normalized_steps) == 0:
            raise bad_request(
                code="ROUTE_STEPS_REQUIRED",
                message="steps must be a non-empty array.",
            )
        product_name = (
            self._lookup_product_name(product_code)
            or self._lookup_product_name(source_product_code)
            or product_code
        )
        process_name_by_code = self._process_name_by_code()
        updated_at = utc_now()
        rows_to_insert: list[dict[str, Any]] = []
        marked_final_count = 0
        for index, step in enumerate(normalized_steps):
            process_code = str(step.get("process_code") or "").strip().upper()
            if not process_code:
                raise bad_request(
                    code="ROUTE_STEP_INVALID",
                    message="Each route step requires process_code.",
                )
            dependency_type = str(step.get("dependency_type") or "FS").strip().upper() or "FS"
            raw_final_process_flag = step.get("is_final_process")
            if raw_final_process_flag is None:
                is_final_process = 0
            elif isinstance(raw_final_process_flag, bool):
                is_final_process = 1 if raw_final_process_flag else 0
            else:
                normalized_final_process_flag = (
                    str(raw_final_process_flag).strip().lower()
                )
                if normalized_final_process_flag in {"1", "true"}:
                    is_final_process = 1
                elif normalized_final_process_flag in {"0", "false", ""}:
                    is_final_process = 0
                else:
                    raise bad_request(
                        code="ROUTE_STEP_INVALID",
                        message="is_final_process must be 0 or 1.",
                    )
            if is_final_process == 1:
                marked_final_count += 1
            rows_to_insert.append(
                {
                    "product_code": product_code,
                    "sequence_no": index + 1,
                    "process_code": process_code,
                    "process_name_cn": process_name_by_code.get(
                        process_code,
                        process_code,
                    ),
                    "dependency_type": dependency_type,
                    "route_no": f"ROUTE-{product_code}",
                    "route_name_cn": product_name,
                    "product_name_cn": product_name,
                    "is_final_process": is_final_process,
                    "updated_at": updated_at,
                }
            )
        if marked_final_count > 1:
            raise bad_request(
                code="ROUTE_FINAL_PROCESS_INVALID",
                message="Only one step can be marked as final process.",
            )
        if marked_final_count == 0 and rows_to_insert:
            rows_to_insert[-1]["is_final_process"] = 1
        with transaction(self.connection):
            self.connection.execute(
                "DELETE FROM masterdata_process_routes WHERE product_code = ?",
                (product_code,),
            )
            self.connection.executemany(
                """
                INSERT INTO masterdata_process_routes (
                    product_code,
                    sequence_no,
                    process_code,
                    process_name_cn,
                    dependency_type,
                    route_no,
                    route_name_cn,
                    product_name_cn,
                    is_final_process,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        row["product_code"],
                        row["sequence_no"],
                        row["process_code"],
                        row["process_name_cn"],
                        row["dependency_type"],
                        row["route_no"],
                        row["route_name_cn"],
                        row["product_name_cn"],
                        row["is_final_process"],
                        row["updated_at"],
                    )
                    for row in rows_to_insert
                ],
            )

    def _lookup_product_name(self, product_code: str | None) -> str | None:
        if not product_code:
            return None
        row = fetch_one(
            self.connection,
            """
            SELECT material_name
            FROM production_orders
            WHERE material_code = ?
            ORDER BY updated_at DESC
            LIMIT 1
            """,
            (product_code,),
        )
        if row is None:
            return None
        return str(row.get("material_name") or "").strip() or None

    def _next_schedule_version_no(self) -> str:
        today = datetime.now().strftime("%Y.%m.%d")
        prefix = f"V{today}-D"
        row = fetch_one(
            self.connection,
            """
            SELECT COUNT(1) AS total
            FROM schedule_versions
            WHERE version_no LIKE ?
            """,
            (f"{prefix}%",),
        )
        next_index = int(row["total"]) + 1 if row else 1
        return f"{prefix}{next_index}"

    def _normalize_batch_dispatch_order_nos(self, value: object) -> list[str]:
        if not isinstance(value, list):
            raise bad_request(
                code="ORDER_BATCH_DISPATCH_ORDER_NOS_INVALID",
                message="order_nos must be an array.",
            )
        normalized_order_nos: list[str] = []
        seen_order_nos: set[str] = set()
        for item in value:
            order_no = str(item or "").strip()
            if not order_no or order_no in seen_order_nos:
                continue
            seen_order_nos.add(order_no)
            normalized_order_nos.append(order_no)
        return normalized_order_nos

    def _is_order_completed_for_dispatch(
        self,
        base_row: dict[str, Any],
        state_row: dict[str, Any] | None,
    ) -> bool:
        explicit_status = str(
            (state_row or {}).get("order_status") or (state_row or {}).get("status") or ""
        ).strip().upper()
        if explicit_status in {"OPEN", "IN_PROGRESS"}:
            return False
        if explicit_status in {"DONE", "COMPLETED", "CLOSED"}:
            return True

        production_qty = _to_number(base_row.get("production_qty"), 0)
        completed_qty = _to_number((state_row or {}).get("completed_qty"), 0)
        remaining_qty = (state_row or {}).get("remaining_qty")
        if remaining_qty is not None and _to_number(remaining_qty, 0) <= SCHEDULE_NUMBER_EPSILON:
            return True
        if (
            production_qty > SCHEDULE_NUMBER_EPSILON
            and completed_qty + SCHEDULE_NUMBER_EPSILON >= production_qty
        ):
            return True

        progress_rate = (state_row or {}).get("progress_rate")
        if progress_rate is not None and _to_number(progress_rate, 0) >= 99.999:
            return True
        return False

    def _resolve_dispatch_actor_name(self, actor: object) -> str:
        if isinstance(actor, dict):
            for field in ("username", "display_name", "user_id"):
                value = str(actor.get(field) or "").strip()
                if value:
                    return value
        return "system"

    def _insert_dispatch_command_record(
        self,
        *,
        target_order_no: str,
        command_type: str,
        payload: dict[str, Any],
    ) -> str:
        command_id = f"CMD-{uuid4().hex[:10].upper()}"
        now = utc_now()
        self.connection.execute(
            """
            INSERT INTO dispatch_commands (
                command_id,
                target_order_no,
                command_type,
                status,
                effective_time,
                reason,
                created_by,
                approver,
                decision,
                decision_reason,
                decision_time,
                created_at,
                updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                command_id,
                target_order_no,
                command_type,
                "PENDING",
                payload.get("effective_time"),
                payload.get("reason"),
                payload.get("created_by"),
                None,
                None,
                None,
                None,
                now,
                now,
            ),
        )
        return command_id

    def _approve_dispatch_command_record(
        self,
        *,
        command_id: str,
        target_order_no: str,
        command_type: str,
        previous_status: str,
        payload: dict[str, Any],
    ) -> None:
        decision = str(payload.get("decision") or "").strip().upper() or "APPROVED"
        now = utc_now()
        self.connection.execute(
            """
            UPDATE dispatch_commands
            SET status = ?,
                approver = ?,
                decision = ?,
                decision_reason = ?,
                decision_time = ?,
                updated_at = ?
            WHERE command_id = ?
            """,
            (
                decision,
                payload.get("approver"),
                decision,
                payload.get("decision_reason"),
                payload.get("decision_time") or now,
                now,
                command_id,
            ),
        )
        if decision == "APPROVED" and previous_status.strip().upper() != "APPROVED":
            self._apply_dispatch_command(
                target_order_no=target_order_no,
                command_type=command_type,
            )

    def _apply_dispatch_command(self, *, target_order_no: str, command_type: str) -> None:
        current = self._get_order_state(target_order_no) or {
            "production_order_no": target_order_no,
        }
        next_row = {
            "production_order_no": target_order_no,
            "promised_due_date": current.get("promised_due_date"),
            "expected_start_date": current.get("expected_start_date"),
            "expected_start_time": current.get("expected_start_time"),
            "expected_finish_time": current.get("expected_finish_time"),
            "priority_level": _normalize_priority_level(current.get("priority_level"), PRIORITY_LEVEL_MAX),
            "urgent_flag": _urgent_flag_from_priority_level(current.get("priority_level")),
            "lock_flag": int(current.get("lock_flag") or 0),
            "frozen_flag": int(current.get("frozen_flag") or 0),
            "status": current.get("status"),
            "order_status": current.get("order_status"),
            "completed_qty": current.get("completed_qty"),
            "remaining_qty": current.get("remaining_qty"),
            "progress_rate": current.get("progress_rate"),
            "production_batch_no": current.get("production_batch_no"),
        }
        normalized = str(command_type or "").strip().upper()
        if normalized == "LOCK":
            next_row["lock_flag"] = 1
        elif normalized == "UNLOCK":
            next_row["lock_flag"] = 0
        elif normalized == "PRIORITY_UP":
            next_row["priority_level"] = max(PRIORITY_LEVEL_MIN, next_row["priority_level"] - 1)
            next_row["urgent_flag"] = _urgent_flag_from_priority_level(next_row["priority_level"])
        elif normalized == "PRIORITY_DOWN":
            next_row["priority_level"] = min(PRIORITY_LEVEL_MAX, next_row["priority_level"] + 1)
            next_row["urgent_flag"] = _urgent_flag_from_priority_level(next_row["priority_level"])
        elif normalized == "PRIORITY":
            next_row["priority_level"] = PRIORITY_LEVEL_MIN
            next_row["urgent_flag"] = 1
        elif normalized == "UNPRIORITY":
            next_row["priority_level"] = PRIORITY_LEVEL_MAX
            next_row["urgent_flag"] = 0
        self._upsert_order_state(next_row)

    def _find_material_reference(self, material_code: str) -> dict[str, Any]:
        row = fetch_one(
            self.connection,
            """
            SELECT
                child_material_name AS material_name_cn,
                spec_model,
                child_unit
            FROM material_issue_items
            WHERE child_material_code = ?
            LIMIT 1
            """,
            (material_code,),
        )
        if row is not None:
            return row
        row = fetch_one(
            self.connection,
            """
            SELECT
                child_material_name AS material_name_cn,
                child_specification AS spec_model,
                child_unit
            FROM bom_children
            WHERE child_material_code = ?
            LIMIT 1
            """,
            (material_code,),
        )
        return row or {}

    def _next_import_order_sequence(self) -> int:
        rows = fetch_all(
            self.connection,
            """
            SELECT production_order_no
            FROM production_orders
            WHERE production_order_no LIKE '%MO-TEST'
            """
        )
        numbers = []
        for row in rows:
            text = str(row["production_order_no"])
            digits = "".join(ch for ch in text.split("MO-TEST")[0] if ch.isdigit())
            if digits:
                numbers.append(int(digits))
        return max(numbers, default=9000) + 1
